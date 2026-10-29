"""Финансовая модель в Excel с формулами и документ «Экономика проекта».

    python docs/ekonomika/build_finmodel.py   → docs/ekonomika/Finmodel.xlsx, Ekonomika_proekta.docx, ekonomika.md

Excel — для экономистов команды: допущения на одном листе (оранжевые ячейки — ввод), сценарий — выпадающий
список, остальное считается формулами: план на 2027–2029, экономика клиента, эффект для заказчика. Формулы
повторяют docs/ekonomika/finmodel.py; автотест пересчитывает книгу и сверяет с ним. Документ и Markdown —
для наставника, жюри и партнёра: цифры из той же модели, у каждого допущения — статус и источник.
"""
import sys
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))

import finmodel as fm  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

OUT_XLSX = HERE / "Finmodel.xlsx"
INPUT = PatternFill("solid", fgColor="FFF4E5")
HEAD = PatternFill("solid", fgColor="E9E4F3")
TOTAL = PatternFill("solid", fgColor="F4F1FA")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="4B3A75")
MONEY = '#,##0'
SHARE = '0%'
COLS = "DEF"                       # 2027, 2028, 2029
D = "'Допущения'!"


class Refs:
    """Адреса ячеек допущений: ключ → ячейка; по годам — ключ[год]; по сценариям — ключ|сценарий."""
    def __init__(self):
        self.cell: dict[str, str] = {}

    def __getitem__(self, key: str) -> str:
        return D + self.cell[key]


def build_assumptions(wb) -> Refs:
    ws = wb.active
    ws.title = "Допущения"
    ref = Refs()
    ws["A1"], ws["A1"].font = "Финансовая модель: допущения", TITLE
    ws["A3"], ws["A3"].font = "Сценарий", BOLD
    ws["D3"] = "базовый"
    ws["D3"].fill = INPUT
    dv = DataValidation(type="list", formula1='"' + ",".join(fm.SCENARIOS) + '"')
    dv.add("D3")
    ws.add_data_validation(dv)
    ws["E3"] = "← выберите: консервативный, базовый или оптимистичный"
    ws["A4"], ws["D4"] = "Номер сценария", '=IF(D3="консервативный",1,IF(D3="оптимистичный",3,2))'
    ref.cell["scenario"] = "$D$4"
    r = 6
    for c, h in zip("ABCDEFGH", ["Ключ", "Допущение", "Единица", "Значение / 2027", "2028", "2029", "Статус",
                                  "Источник, как проверить"]):
        ws[f"{c}{r}"], ws[f"{c}{r}"].font, ws[f"{c}{r}"].fill = h, BOLD, HEAD
    for group, items in fm.GROUPS.items():
        r += 1
        ws[f"B{r}"], ws[f"B{r}"].font = group, BOLD
        for it in items:
            rows = ([(f"{it.label} — {s}", it.value[s], f"{it.key}|{s}") for s in fm.SCENARIOS]
                    if isinstance(it.value, dict) else [(it.label, it.value, it.key)])
            for label, value, key in rows:
                r += 1
                ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"G{r}"], ws[f"H{r}"] = key, label, it.unit, it.status, it.source
                vals = value if isinstance(value, list) else [value]
                for col, v in zip(COLS, vals):
                    cell = ws[f"{col}{r}"]
                    cell.value, cell.fill = v, INPUT
                    cell.number_format = SHARE if isinstance(v, float) and v < 1.5 and it.unit in ("доля", "доля прочих") \
                        else ('0.00' if isinstance(v, float) else MONEY)
                if isinstance(value, list):
                    for y, col in enumerate(COLS):
                        ref.cell[f"{key}[{y}]"] = f"${col}${r}"
                else:
                    ref.cell[key] = f"$D${r}"
    for col, width in zip("ABCDEFGH", (16, 62, 12, 15, 13, 13, 10, 70)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=7):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ref


def pick(ref: Refs, key: str, y: int | None = None) -> str:
    """Значение выбранного сценария: CHOOSE(номер, консервативный, базовый, оптимистичный)."""
    cells = [ref[f"{key}|{s}" + (f"[{y}]" if y is not None else "")] for s in fm.SCENARIOS]
    return f"CHOOSE({ref['scenario']},{','.join(cells)})"


PLAN_ROWS = ["Новых клиентов", "Отток подписок", "Клиентов сохранилось с прошлого года", "Выручка с клиента в год",
             "Выручка: подписка и модули", "Выручка: внедрение", "Выручка: услуги", "Выручка", "Комиссия интеграторам",
             "Выезды", "Штат (масштаб)", "ФОТ: разработка", "ФОТ: внедрение и поддержка", "ФОТ: продажи",
             "ФОТ: экономика и договоры", "Себестоимость", "Валовая прибыль", "Прочие расходы", "Операционные расходы",
             "EBITDA", "Страховые взносы в ФОТ", "Налог (УСН 6 % с вычетом взносов)", "Гранты", "Денежный поток",
             "Деньги накопленно", "Активных клиентов на конец года", "Загрузка внедрения"]


def build_plan(wb, ref: Refs) -> dict[str, int]:
    ws = wb.create_sheet("План")
    ws["A1"], ws["A1"].font = "План на 2027–2029 — сценарий выбирается на листе «Допущения»", TITLE
    ws["A2"] = f'="Сценарий: "&{D}$D$3'
    ws["A4"], ws["A4"].font = "Показатель", BOLD
    for y, col in enumerate(COLS):
        ws[f"{col}4"], ws[f"{col}4"].font, ws[f"{col}4"].fill = fm.YEARS[y], BOLD, HEAD
    row = {name: 5 + i for i, name in enumerate(PLAN_ROWS)}
    for name, r in row.items():
        ws[f"A{r}"] = name
    for y, col in enumerate(COLS):
        prev = COLS[y - 1] if y else None

        def c(name, col=col):
            return f"{col}{row[name]}"
        idx = f"(1+{ref['price_index']})^{y}"
        f = {
            "Новых клиентов": f"={pick(ref, 'new_clients', y)}",
            "Отток подписок": f"={pick(ref, 'churn')}",
            "Клиентов сохранилось с прошлого года": "=0" if not y else f"={prev}{row['Активных клиентов на конец года']}*(1-{c('Отток подписок')})",
            "Выручка с клиента в год": (f"=(({ref['sub_dept']}*(1-{ref['ent_share']})+{ref['sub_ent']}*{ref['ent_share']})"
                                        f"+({ref['module_tkp']}*{ref['module_tkp_share']}+{ref['module_1c']}*{ref['module_1c_share']}))*{idx}"),
            "Выручка: подписка и модули": (f"={c('Клиентов сохранилось с прошлого года')}*{c('Выручка с клиента в год')}"
                                            f"+{c('Новых клиентов')}*{c('Выручка с клиента в год')}*{ref['first_year_share']}"),
            "Выручка: внедрение": f"={c('Новых клиентов')}*{ref['impl_price']}*{idx}",
            "Выручка: услуги": f"={c('Новых клиентов')}*{ref['services_share']}*{ref['services_price']}*{idx}",
            "Выручка": f"={c('Выручка: подписка и модули')}+{c('Выручка: внедрение')}+{c('Выручка: услуги')}",
            "Комиссия интеграторам": f"={c('Выручка: подписка и модули')}*{ref['channel_share']}*{ref['commission']}",
            "Выезды": f"={c('Новых клиентов')}*{ref['travel']}",
            "Штат (масштаб)": "=1" if not y else f"={pick(ref, 'team_scale')}",
        }
        for role, name in (("dev", "ФОТ: разработка"), ("impl", "ФОТ: внедрение и поддержка"),
                           ("sales", "ФОТ: продажи"), ("admin", "ФОТ: экономика и договоры")):
            f[name] = (f"={ref[f'fte_{role}[{y}]']}*{c('Штат (масштаб)')}*{ref[f'sal_{role}[{y}]']}*12"
                       f"*(1+{ref[f'contrib[{y}]']})")
        other = "+".join(ref[f"{k}[{y}]"] for k in ("opex_accounting", "opex_legal", "opex_hardware", "opex_marketing",
                                                    "opex_office", "opex_agpl"))
        payroll = "+".join(c(n) for n in ("ФОТ: разработка", "ФОТ: внедрение и поддержка", "ФОТ: продажи",
                                          "ФОТ: экономика и договоры"))
        f.update({
            "Себестоимость": f"={c('Комиссия интеграторам')}+{c('Выезды')}+{c('ФОТ: внедрение и поддержка')}",
            "Валовая прибыль": f"={c('Выручка')}-{c('Себестоимость')}",
            "Прочие расходы": f"=({other})*(1+{ref['opex_reserve']})",
            "Операционные расходы": f"={c('ФОТ: разработка')}+{c('ФОТ: продажи')}+{c('ФОТ: экономика и договоры')}+{c('Прочие расходы')}",
            "EBITDA": f"={c('Валовая прибыль')}-{c('Операционные расходы')}",
            "Страховые взносы в ФОТ": f"=({payroll})*{ref[f'contrib[{y}]']}/(1+{ref[f'contrib[{y}]']})",
            "Налог (УСН 6 % с вычетом взносов)": (f"=MAX({c('Выручка')}*{ref['usn_rate']}-{c('Страховые взносы в ФОТ')},"
                                                  f"{c('Выручка')}*{ref['usn_rate']}/2)"),
            "Гранты": f"={pick(ref, 'grant', y)}",
            "Денежный поток": f"={c('EBITDA')}-{c('Налог (УСН 6 % с вычетом взносов)')}+{c('Гранты')}",
            "Деньги накопленно": f"={c('Денежный поток')}" + (f"+{prev}{row['Деньги накопленно']}" if y else ""),
            "Активных клиентов на конец года": f"={c('Клиентов сохранилось с прошлого года')}+{c('Новых клиентов')}",
            "Загрузка внедрения": (f"=({c('Новых клиентов')}*{ref['impl_days']}*{ref['first_factor'] if not y else 1}"
                                   f"+{c('Новых клиентов')}*{ref['services_share']}*{ref['services_days']}"
                                   f"+{c('Активных клиентов на конец года')}*{ref['support_hours']}*12/8)"
                                   f"/({ref[f'fte_impl[{y}]']}*{c('Штат (масштаб)')}*{ref['work_days']}*0.8)"),
        })
        for name, formula in f.items():
            cell = ws[c(name)]
            cell.value = formula
            cell.number_format = SHARE if name in ("Отток подписок", "Загрузка внедрения") else (
                "0.0" if name in ("Новых клиентов", "Клиентов сохранилось с прошлого года", "Штат (масштаб)",
                                  "Активных клиентов на конец года") else MONEY)
    for name in ("Выручка", "Валовая прибыль", "EBITDA", "Деньги накопленно"):
        for col in "A" + COLS:
            ws[f"{col}{row[name]}"].font, ws[f"{col}{row[name]}"].fill = BOLD, TOTAL
    ws.column_dimensions["A"].width = 44
    for col in COLS:
        ws.column_dimensions[col].width = 16
    return row


UNIT_ROWS = ["Выручка с клиента в год (2028)", "Стоимость чел.-дня внедрения", "Выручка первого года",
             "Прямые затраты первого года", "Маржа первого года", "Маржа первого года, %", "Маржа со второго года",
             "Маржа со второго года, %", "Стоимость привлечения (CAC)", "Ценность клиента за 5 лет (LTV)", "LTV / CAC",
             "Окупаемость привлечения, мес."]


def build_unit(wb, ref: Refs, plan_row: dict) -> dict[str, int]:
    ws = wb.create_sheet("Клиент")
    ws["A1"], ws["A1"].font = "Экономика одного клиента — маржинальная, цены и ставки 2028 г.", TITLE
    row = {name: 3 + i for i, name in enumerate(UNIT_ROWS)}
    y = 1
    idx = f"(1+{ref['price_index']})^{y}"
    churn = pick(ref, "churn")
    commission = f"{ref['channel_share']}*{ref['commission']}"
    support = f"{ref['support_hours']}*12/8*B{row['Стоимость чел.-дня внедрения']}"
    f = {
        "Выручка с клиента в год (2028)": f"='План'!E{plan_row['Выручка с клиента в год']}",
        "Стоимость чел.-дня внедрения": f"={ref['sal_impl[1]']}*12*(1+{ref['contrib[1]']})/{ref['work_days']}",
        "Выручка первого года": (f"={ref['impl_price']}*{idx}+B{row['Выручка с клиента в год (2028)']}*{ref['first_year_share']}"
                                 f"+{ref['services_share']}*{ref['services_price']}*{idx}"),
        "Прямые затраты первого года": (f"=({ref['impl_days']}+{ref['services_share']}*{ref['services_days']})"
                                        f"*B{row['Стоимость чел.-дня внедрения']}+{ref['travel']}+{support}*{ref['first_year_share']}"
                                        f"+B{row['Выручка с клиента в год (2028)']}*{ref['first_year_share']}*{commission}"),
        "Маржа первого года": f"=B{row['Выручка первого года']}-B{row['Прямые затраты первого года']}",
        "Маржа первого года, %": f"=B{row['Маржа первого года']}/B{row['Выручка первого года']}",
        "Маржа со второго года": f"=B{row['Выручка с клиента в год (2028)']}*(1-{commission})-{support}",
        "Маржа со второго года, %": f"=B{row['Маржа со второго года']}/B{row['Выручка с клиента в год (2028)']}",
        "Стоимость привлечения (CAC)": (f"=({ref['presale_days']}+{ref['pilot_days']})*B{row['Стоимость чел.-дня внедрения']}"
                                        f"+2*{ref['travel']}+{ref['opex_marketing[1]']}/MAX({pick(ref, 'new_clients', 1)},1)"),
        "Ценность клиента за 5 лет (LTV)": (f"=B{row['Маржа первого года']}+B{row['Маржа со второго года']}*((1-{churn})"
                                            f"+(1-{churn})^2+(1-{churn})^3+(1-{churn})^4)"),
        "LTV / CAC": f"=B{row['Ценность клиента за 5 лет (LTV)']}/B{row['Стоимость привлечения (CAC)']}",
        "Окупаемость привлечения, мес.": f"=B{row['Стоимость привлечения (CAC)']}/(B{row['Маржа со второго года']}/12)",
    }
    for name, r in row.items():
        ws[f"A{r}"], ws[f"B{r}"] = name, f[name]
        ws[f"B{r}"].number_format = SHARE if name.endswith("%") else ("0.0" if name in ("LTV / CAC",
                                                                                      "Окупаемость привлечения, мес.") else MONEY)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    return row


CUST_ROWS = ["Сотрудников", "Ответов в день на сотрудника", "Минут на ответ вручную", "Экономия времени на ответе",
             "ТКП в месяц", "Минут на ТКП вручную", "Экономия времени на ТКП", "Минут поиска в день на сотрудника",
             "Экономия времени поиска", "Зарплата, ₽ в месяц", "Подписка, ₽ в год", "Модули, ₽ в год",
             "Внедрение, ₽", "Оборудование, ₽", "Администратор базы знаний, ч в месяц",
             "Час сотрудника с взносами", "Высвобождается часов в год", "Эффект, ₽ в год", "Затраты первого года",
             "Затраты со второго года", "Окупаемость, мес.", "Возврат за 3 года (ROI)", "Цена — доля от эффекта",
             "Запросов в год", "Полная стоимость запроса", "Экономия на запросе"]
CUST_INPUT = {"Сотрудников": "people", "Ответов в день на сотрудника": "answers_per_day",
              "Минут на ответ вручную": "minutes_answer", "Экономия времени на ответе": "saved_answer",
              "ТКП в месяц": "tkp_per_month", "Минут на ТКП вручную": "minutes_tkp", "Экономия времени на ТКП": "saved_tkp",
              "Минут поиска в день на сотрудника": "search_minutes_day", "Экономия времени поиска": "saved_search",
              "Зарплата, ₽ в месяц": "salary", "Подписка, ₽ в год": "subscription", "Модули, ₽ в год": "modules",
              "Внедрение, ₽": "implementation", "Оборудование, ₽": "hardware",
              "Администратор базы знаний, ч в месяц": "admin_hours"}


def build_customer(wb) -> dict[str, int]:
    ws = wb.create_sheet("Эффект для заказчика")
    ws["A1"], ws["A1"].font = "Эффект для заказчика — ввод в оранжевых ячейках; на пилоте — замер вкладки «Эффект»", TITLE
    row = {name: 4 + i for i, name in enumerate(CUST_ROWS)}
    for k, (name, c) in enumerate(fm.CUSTOMER.items()):
        col = "BC"[k]
        ws[f"{col}3"], ws[f"{col}3"].font, ws[f"{col}3"].fill = c["label"], BOLD, HEAD
        ws[f"{col}3"].alignment = Alignment(wrap_text=True)

        def x(name, col=col):
            return f"{col}{row[name]}"
        for name, key in CUST_INPUT.items():
            cell = ws[x(name)]
            cell.value, cell.fill = c.get(key, 0), INPUT
            cell.number_format = SHARE if key.startswith("saved") else MONEY
        f = {
            "Час сотрудника с взносами": f"={x('Зарплата, ₽ в месяц')}*1.3/{fm.HOURS_MONTH}",
            "Высвобождается часов в год": (f"=({x('Сотрудников')}*{x('Ответов в день на сотрудника')}*247*{x('Минут на ответ вручную')}"
                                           f"*{x('Экономия времени на ответе')}+{x('ТКП в месяц')}*12*{x('Минут на ТКП вручную')}"
                                           f"*{x('Экономия времени на ТКП')}+{x('Сотрудников')}*{x('Минут поиска в день на сотрудника')}"
                                           f"*{x('Экономия времени поиска')}*247)/60"),
            "Эффект, ₽ в год": f"={x('Высвобождается часов в год')}*{x('Час сотрудника с взносами')}",
            "Затраты со второго года": (f"={x('Подписка, ₽ в год')}+{x('Модули, ₽ в год')}+{x('Администратор базы знаний, ч в месяц')}"
                                        f"*12*{x('Час сотрудника с взносами')}+{fm.POWER_MONTH}*12"),
            "Затраты первого года": f"={x('Затраты со второго года')}+{x('Внедрение, ₽')}+{x('Оборудование, ₽')}",
            "Окупаемость, мес.": f"={x('Затраты первого года')}/({x('Эффект, ₽ в год')}/12)",
            "Возврат за 3 года (ROI)": (f"=(3*{x('Эффект, ₽ в год')}-{x('Затраты первого года')}-2*{x('Затраты со второго года')})"
                                        f"/({x('Затраты первого года')}+2*{x('Затраты со второго года')})"),
            "Цена — доля от эффекта": f"=({x('Подписка, ₽ в год')}+{x('Модули, ₽ в год')})/{x('Эффект, ₽ в год')}",
            "Запросов в год": f"={x('Сотрудников')}*{x('Ответов в день на сотрудника')}*247+{x('ТКП в месяц')}*12",
            "Полная стоимость запроса": f"=IF({x('Запросов в год')}>0,{x('Затраты со второго года')}/{x('Запросов в год')},0)",
            "Экономия на запросе": (f"=IF({x('Запросов в год')}>0,({x('Минут на ответ вручную')}*{x('Экономия времени на ответе')}"
                                    f"*{x('Сотрудников')}*{x('Ответов в день на сотрудника')}*247+{x('Минут на ТКП вручную')}"
                                    f"*{x('Экономия времени на ТКП')}*{x('ТКП в месяц')}*12)/{x('Запросов в год')}/60"
                                    f"*{x('Час сотрудника с взносами')},0)"),
        }
        for name, formula in f.items():
            ws[x(name)] = formula
            ws[x(name)].number_format = SHARE if name in ("Возврат за 3 года (ROI)", "Цена — доля от эффекта") else (
                "0.0" if name == "Окупаемость, мес." else MONEY)
    for name, r in row.items():
        ws[f"A{r}"] = name
        ws[f"D{r}"] = fm.CUSTOMER_NOTES.get(CUST_INPUT.get(name, ""), "")
    for col, width in zip("ABCD", (40, 26, 26, 70)):
        ws.column_dimensions[col].width = width
    return row


def build_static(wb) -> None:
    ws = wb.create_sheet("Чувствительность")
    ws["A1"], ws["A1"].font = "Деньги на конец 2029 г. (накопленно), базовый сценарий — рассчитано скриптом", TITLE
    for c, h in zip("ABCD", ["Допущение", "Хуже", "Базовый", "Лучше"]):
        ws[f"{c}3"], ws[f"{c}3"].font, ws[f"{c}3"].fill = h, BOLD, HEAD
    for i, (label, low, base, high) in enumerate(fm.sensitivity(), start=4):
        worse, better = sorted((low, high))
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"] = label, worse, base, better
        for c in "BCD":
            ws[f"{c}{i}"].number_format = MONEY
    ws.column_dimensions["A"].width = 42
    for c in "BCD":
        ws.column_dimensions[c].width = 16
    bs = wb.create_sheet("Созданный продукт")
    bs["A1"], bs["A1"].font = "Стоимость созданного продукта — затратный подход (нормативная трудоёмкость)", TITLE
    bs["A3"], bs["B3"] = "Работа", "чел.-дн."
    for c in "AB":
        bs[f"{c}3"].font, bs[f"{c}3"].fill = BOLD, HEAD
    for i, (work, days) in enumerate(fm.BUILT, start=4):
        bs[f"A{i}"], bs[f"B{i}"] = work, days
    last = 3 + len(fm.BUILT)
    bs[f"A{last + 1}"], bs[f"B{last + 1}"] = "Итого", f"=SUM(B4:B{last})"
    bs[f"A{last + 1}"].font = bs[f"B{last + 1}"].font = BOLD
    for j, (label, rate) in enumerate(fm.RATES.items(), start=last + 3):
        bs[f"A{j}"], bs[f"B{j}"] = f"Стоимость по ставке: {label}", f"=B{last + 1}*{rate:.2f}"
        bs[f"B{j}"].number_format = MONEY
    bs.column_dimensions["A"].width = 100
    bs.column_dimensions["B"].width = 16


def build_xlsx(path: Path = OUT_XLSX) -> dict:
    wb = Workbook()
    ref = build_assumptions(wb)
    plan_row = build_plan(wb, ref)
    unit_row = build_unit(wb, ref, plan_row)
    cust_row = build_customer(wb)
    build_static(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return {"plan": plan_row, "unit": unit_row, "customer": cust_row}


if __name__ == "__main__":
    build_xlsx()
    from build_ekonomika import build_docs        # документ и Markdown — из тех же расчётов
    build_docs()
    print(OUT_XLSX)
