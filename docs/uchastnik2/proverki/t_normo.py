"""Нормоконтроль (kbagent/normocontrol.py) на трудных документах: что он пропускает и где ошибается."""
import io
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))   # корень репозитория
from kbagent.normocontrol import check_docx  # noqa: E402

LONG = ("В целях проверки состояния и правильности применения средств измерений в подразделениях "
        "предприятия приказываю создать комиссию в составе трёх человек и провести проверку до конца месяца.")
SHORT = "Прошу выдать пропуск на территорию цеха № 3 для проведения работ по ремонту."  # < 100 знаков


def base(paras, size=13, font="Times New Roman", indent=1.25, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left=25):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(left), Mm(10), Mm(20), Mm(20)
    for text in paras:
        p = d.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.font.name = Pt(size), font
        p.paragraph_format.first_line_indent = Cm(indent)
        p.paragraph_format.alignment = align
        p.paragraph_format.line_spacing = 1.0
    return d


def save(d):
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


cases = {}
cases["контроль: правильный приказ"] = save(base(["ПРИКАЗ", "05.06.2026 № 17", "О создании комиссии", LONG]))
cases["пустой документ"] = save(Document())
cases["только таблица, без текста"] = (lambda d: (d.add_table(rows=2, cols=2), save(d))[1])(Document())
cases["служебка из коротких абзацев: 10 пт, по левому краю, без отступа"] = save(
    base(["СЛУЖЕБНАЯ ЗАПИСКА", "О выдаче пропуска", SHORT, SHORT], size=10, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT))
cases["шрифт Comic Sans 13 пт"] = save(base(["ПРИКАЗ", "О создании комиссии", LONG], font="Comic Sans MS"))
cases["двойной интервал"] = (lambda d: ([setattr(p.paragraph_format, "line_spacing", 2.0) for p in d.paragraphs],
                                        save(d))[1])(base(["ПРИКАЗ", "О создании комиссии", LONG]))
d = base(["ПРИКАЗ", "О создании комиссии", LONG])
sec = d.add_section()
sec.orientation, sec.page_width, sec.page_height = WD_ORIENT.LANDSCAPE, Mm(297), Mm(210)
sec.left_margin = Mm(5)
d.add_paragraph(LONG)
cases["второй раздел с левым полем 5 мм"] = save(d)
cases["дата 5.06.2026 (без нуля)"] = save(base(["ПРИКАЗ", "5.06.2026 № 17", "О создании комиссии", LONG]))
cases["дата 05/06/2026"] = save(base(["ПРИКАЗ", "05/06/2026 № 17", "О создании комиссии", LONG]))
cases["PDF под видом DOCX"] = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
cases["случайные байты"] = b"PK\x03\x04" + b"\x00" * 200

for name, data in cases.items():
    try:
        r = check_docx(data)
        issues = "; ".join(f"[{i.level[:4]} п.{i.clause}] {i.text[:70]}" for i in r.issues) or "ЗАМЕЧАНИЙ НЕТ"
        print(f"— {name}\n    абзацев: {r.paragraphs}; соблюдено: {', '.join(r.passed)}\n    {issues}")
    except Exception as e:  # noqa: BLE001
        print(f"— {name}\n    ИСКЛЮЧЕНИЕ {type(e).__name__}: {str(e)[:80]} (агент покажет «не удалось прочитать»)")
