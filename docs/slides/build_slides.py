"""Черновик презентации и слайд архитектуры в шаблоне акселератора.

    python docs/slides/build_slides.py

Создаёт:
  docs/slides/prezentatsiya_chernovik.pptx — все слайды 1-го этапа, заполненные данными проекта;
  docs/slides/arhitektura.pptx             — отдельный слайд «Технологическое ядро проекта».
Оранжевым в квадратных скобках выделено то, что команда вписывает сама (интервью, ФИО).
Цифры качества берутся из этого файла — после нового прогона eval/run_eval.py обновите RESULTS.
"""
import copy
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "shablon.pptx"

RESULTS = {
    "correct": "47 из 47", "refusals": "11 из 11", "db": "12 из 12", "wrong": "0",
    "questions": "74", "time_gpu": "1,3 с", "time_cpu": "около 7 с",
}

PURPLE = RGBColor(0x80, 0x64, 0xA2)
DEEP = RGBColor(0x4B, 0x3A, 0x66)
LAV = RGBColor(0xED, 0xEA, 0xF0)
LAV2 = RGBColor(0xD8, 0xD3, 0xE0)
ORANGE = RGBColor(0xEF, 0x7D, 0x20)
ORANGE_TINT = RGBColor(0xFD, 0xEF, 0xE3)
TEXT = RGBColor(0x1F, 0x1A, 0x2E)
MUTED = RGBColor(0x6E, 0x66, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


# ---------------------------------------------------------------- текст

def _add_runs(paragraph, text, size, color, bold=False):
    """Текст в [квадратных скобках] — пропуск, который заполняет команда: выделяется оранжевым."""
    import re
    for part in re.split(r"(\[[^\]]+\])", text):
        if not part:
            continue
        r = paragraph.add_run()
        r.text = part
        is_gap = part.startswith("[")
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold or is_gap
        r.font.color.rgb = ORANGE if is_gap else color


def _bullet(paragraph, on=True):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if on:
        pPr.set("marL", str(Inches(0.22)))
        pPr.set("indent", str(-Inches(0.2)))
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")
    else:
        etree.SubElement(pPr, qn("a:buNone"))


def fill(shape, items, size=14, color=TEXT, space=4, anchor=MSO_ANCHOR.TOP):
    """items: строки; «## текст» — подзаголовок, «- текст» — пункт списка, «**Метка:** текст» — жирная метка."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        p.alignment = PP_ALIGN.LEFT
        if item.startswith("## "):
            _bullet(p, False)
            _add_runs(p, item[3:], size + 1, PURPLE, bold=True)
            p.space_before = Pt(6)
            continue
        bullet = item.startswith("- ")
        _bullet(p, bullet)
        body = item[2:] if bullet else item
        if body.startswith("**") and "**" in body[2:]:
            label, rest = body[2:].split("**", 1)
            _add_runs(p, label, size, color, bold=True)
            body = rest
        _add_runs(p, body, size, color)


def textbox(slide, x, y, w, h, items, size=14, color=TEXT, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill(tb, items, size, color)
    return tb


def card(slide, x, y, w, h, big, small, fill_color=PURPLE, fg=WHITE, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.adjustments[0] = 0.14
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = big
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(26), True, fg
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = small
    r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(12), fg
    return s


def shape_by_text(slide, needle):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise KeyError(needle)


def remove(shape):
    slide_part = shape.part
    el = shape._element
    if shape.shape_type == 13:   # картинка — заодно отвязываем файл изображения
        rid = el.blipFill.blip.rEmbed
        el.getparent().remove(el)
        slide_part.drop_rel(rid)
    else:
        el.getparent().remove(el)


def big_pictures(slide, min_width=4.0):
    return [sh for sh in slide.shapes if sh.shape_type == 13 and sh.width > Inches(min_width)]


def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def set_cell(cell, text, size=12, bold=False, color=TEXT, fill_color=None):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _add_runs(p, text, size, color, bold)
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    if fill_color is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill_color


def ensure_rows(table, n):
    """Добавляет строки копированием последней (python-pptx не умеет добавлять строки)."""
    tbl = table._tbl
    while len(table.rows) < n:
        tbl.append(copy.deepcopy(tbl.tr_lst[-1]))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- архитектура

def draw_architecture(slide):
    """Схема технологического ядра (слайд 4)."""
    sh = slide.shapes

    def box(x, y, w, h, title, sub=None, fill_color=WHITE, border=PURPLE, fg=TEXT,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, title_size=12, sub_size=10.5, name=None):
        s = sh.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            s.name = name
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.18
        s.fill.solid()
        s.fill.fore_color.rgb = fill_color
        if border is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = border
            s.line.width = Pt(1.25)
        s.shadow.inherit = False
        tf = s.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(title_size), True, fg
        if sub:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = sub
            r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(sub_size), fg
        return s

    def arrow(x1, y1, x2, y2):
        c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = PURPLE
        c.line.width = Pt(1.5)
        tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")

    def badge(x, y, n):
        s = sh.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.3), Inches(0.3))
        s.name = f"Защита {n}"
        s.fill.solid()
        s.fill.fore_color.rgb = ORANGE
        s.line.color.rgb = WHITE
        s.line.width = Pt(1)
        s.shadow.inherit = False
        tf = s.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(n)
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(11), True, WHITE

    def caption(x, y, w, text, bold=True):
        textbox(slide, x, y, w, 0.22, [text], size=10, color=MUTED).text_frame.paragraphs[0].runs[0].font.bold = bold

    textbox(slide, 0.62, 1.45, 10.9, 0.35,
            ["Отвечает только по документам и учётной базе предприятия — со ссылкой на источник и без интернета"],
            size=15, name="Подзаголовок")

    Y2, H2 = 3.42, 1.0
    widths = [1.2, 1.6, 1.75, 1.6, 2.0, 1.45, 1.85]
    gap = (12.8 - 0.5 - sum(widths)) / (len(widths) - 1)
    xs = [0.5]
    for w in widths[:-1]:
        xs.append(xs[-1] + w + gap)
    lane = [
        ("Вопрос", "сотрудника", WHITE, PURPLE, TEXT),
        ("Где искать?", "GigaChat выбирает: документы или база", WHITE, PURPLE, TEXT),
        ("Гибридный поиск", "по словам (BM25) и по смыслу (bge-m3)", WHITE, PURPLE, TEXT),
        ("Оценка релевантности", "переранжировщик bge-reranker-v2-m3", WHITE, PURPLE, TEXT),
        ("GigaChat 3.1 Lightning", "отвечает только по найденным фрагментам", PURPLE, None, WHITE),
        ("Проверка ответа", "ссылки верны, ответ о том же предмете", WHITE, PURPLE, TEXT),
        ("Ответ со ссылкой", "документ · пункт · страница", DEEP, None, WHITE),
    ]
    for (title, sub, fc, border, fg), x, w in zip(lane, xs, widths):
        box(x, Y2, w, H2, title, sub, fill_color=fc, border=border, fg=fg, name=f"Шаг: {title}")
    for i in range(len(xs) - 1):
        arrow(xs[i] + widths[i], Y2 + H2 / 2, xs[i + 1], Y2 + H2 / 2)
    badge(xs[3] + widths[3] - 0.2, Y2 - 0.24, 1)
    badge(xs[4] + widths[4] - 0.2, Y2 - 0.24, 2)
    badge(xs[5] + widths[5] - 0.52, Y2 - 0.24, 3)
    badge(xs[5] + widths[5] - 0.2, Y2 - 0.24, 4)
    caption(xs[0], Y2 - 0.27, 4.0, "ВОПРОС ПО ДОКУМЕНТАМ")

    Y1, H1 = 2.22, 0.78
    sx, sw = xs[2], widths[2]
    for x, w, t, s in [(0.5, 1.2, "PDF, DOCX", "регламенты, ГОСТы"),
                       (1.88, sx - 0.18 - 1.88, "Разбор на пункты", "номер пункта и страницы"),
                       (sx, sw, "Индекс", "слова и смысловые векторы")]:
        box(x, Y1, w, H1, t, s, fill_color=LAV, border=None, name=f"Индексация: {t}")
    arrow(1.7, Y1 + H1 / 2, 1.88, Y1 + H1 / 2)
    arrow(sx - 0.18, Y1 + H1 / 2, sx, Y1 + H1 / 2)
    arrow(sx + sw / 2, Y1 + H1, sx + sw / 2, Y2)
    caption(0.5, Y1 - 0.27, 4.0, "ИНДЕКСАЦИЯ — при добавлении документов")

    lx = xs[3]
    legend = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(lx), Inches(Y1), Inches(12.8 - lx), Inches(H1))
    legend.name = "Четыре уровня защиты"
    legend.adjustments[0] = 0.12
    legend.fill.solid()
    legend.fill.fore_color.rgb = ORANGE_TINT
    legend.line.color.rgb = ORANGE
    legend.line.width = Pt(1)
    legend.shadow.inherit = False
    tf = legend.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (t, b) in enumerate([
            ("Четыре уровня защиты от выдумок", True),
            ("1 — нерелевантно: «ответа нет», модель не вызывается   2 — модель видит только найденное", False),
            ("3 — ответ без верной ссылки не показывается   4 — ответ о том же предмете, а не о похожем", False)]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = t
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(11 if b else 10.5), b, TEXT

    Y3, H3 = 4.88, 0.8
    rx, rw = xs[1], widths[1]
    box(rx, Y3, 2.0, H3, "GigaChat → SQL", "только SELECT, только чтение", name="Учёт: SQL")
    box(rx + 2.2, Y3, 1.7, H3, "Учётная база", "1С · ERP · MES", fill_color=LAV2, border=None,
        shape=MSO_SHAPE.CAN, name="Учёт: база")
    box(rx + 4.1, Y3, 2.3, H3, "Ответ + запрос + таблица", "запрос виден пользователю",
        fill_color=DEEP, border=None, fg=WHITE, name="Учёт: ответ")
    arrow(rx + rw / 2, Y2 + H2, rx + rw / 2, Y3)
    arrow(rx + 2.0, Y3 + H3 / 2, rx + 2.2, Y3 + H3 / 2)
    arrow(rx + 3.9, Y3 + H3 / 2, rx + 4.1, Y3 + H3 / 2)
    caption(rx + rw / 2 + 0.12, Y2 + H2 + 0.03, 1.6, "вопрос по учёту", bold=False)
    jx, jw = xs[6], widths[6]
    box(jx, Y3, jw, H3, "Журнал запросов", "кто, что спросил, что получил", fill_color=LAV, border=None,
        name="Журнал")
    arrow(jx + jw / 2, Y2 + H2, jx + jw / 2, Y3)
    arrow(rx + 6.4, Y3 + H3 / 2, jx, Y3 + H3 / 2)

    YF, HF = 6.12, 0.82
    facts = [("Отечественная модель", "GigaChat 3.1 Lightning, MIT, без иностранных весов"),
             ("Без интернета", "одна видеокарта 12 ГБ или только процессор"),
             ("Скорость", f"{RESULTS['time_gpu']} на ответ (RTX 5070), {RESULTS['time_cpu']} без видеокарты"),
             ("Указ № 529", "критическая технология № 13, сквозная № 25")]
    fw = (12.3 - 3 * 0.2) / 4
    for i, (t, s) in enumerate(facts):
        box(0.5 + i * (fw + 0.2), YF, fw, HF, t, s, fill_color=LAV2, border=None, title_size=12, sub_size=11,
            name=f"Факт: {t}")

    notes(slide, "Технологическое ядро — генерация с опорой на найденные фрагменты документов. Модель не обучается "
                 "на документах: документы разбираются на пункты и индексируются, на каждый вопрос модель получает "
                 "только найденные фрагменты. Поэтому у каждого ответа есть источник, новые документы подключаются "
                 "за секунды, а права доступа соблюдаются. Четыре уровня защиты от выдумок отмечены оранжевыми "
                 "номерами; четвёртый — проверка, что ответ относится к тому же предмету, а не к похожему. "
                 "Вопросы по учёту агент переводит в SQL-запрос только на чтение.")


# ---------------------------------------------------------------- слайды

def build_deck(out: Path):
    prs = Presentation(TEMPLATE)
    S = prs.slides

    # 1. Титул
    s = S[0]
    title = shape_by_text(s, "Название проекта")
    place(title, 0.55, 1.5, 9.8, 1.5)
    fill(title, ["AI-агент для поиска ответов по внутренней документации и базам данных предприятия"],
         size=30, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    for p in title.text_frame.paragraphs:
        p.line_spacing = 1.0
        for r in p.runs:
            r.font.bold = True
    sub = shape_by_text(s, "ФИО, контакты")
    place(sub, 0.55, 6.2, 10.5, 0.75)
    fill(sub, ["[ФИО участников, телефон, почта]",
               "РГРТУ им. В.Ф. Уткина · задача индустриального партнёра № 28 — ООО «НПВП «НаноТехноПром»"],
         size=12, color=WHITE, space=2)
    notes(s, "Одна фраза о проекте: агент отвечает сотрудникам по внутренним документам и учётной базе со ссылкой на "
             "пункт, честно говорит «не нашёл», работает без интернета на отечественной модели.")

    # 2. Проблема
    s = S[1]
    problem = shape_by_text(s, "Опишите, чью проблему")
    place(problem, 0.75, 2.5, 11.4, 4.2)
    fill(problem, [
        "**Чья проблема.** Сотрудники производственных и научно-производственных предприятий — инженеры, технологи, "
        "новые сотрудники. Платит руководитель: он теряет время специалистов и получает ошибки.",
        "**Как решается сейчас.** Поиск по сетевым папкам, порталу и бумажным регламентам, вопросы опытным коллегам. "
        "Облачные ИИ-сервисы использовать нельзя: коммерческая тайна и защита данных.",
        "**Почему это плохо.** Долго; ответ зависит от того, кто «помнит»; ошибки из-за устаревших версий документов; "
        "универсальные чат-боты выдумывают ответы, и им нельзя доверять в рабочих вопросах.",
        "**Масштаб.** По оценке McKinsey Global Institute, офисные сотрудники тратят на поиск и сбор информации "
        "около 1,8 часа в день — порядка 20 % рабочего времени.",
        "**Кто подтвердил.** Задача поставлена индустриальным партнёром — ООО «НПВП «НаноТехноПром» (задача № 28). "
        "[Интервью: N сотрудников партнёра и M предприятий Рязани — цифры и цитаты по итогам интервью]",
    ], size=15, space=9)
    notes(s, "Не говорить «аналогов нет». Главное — цифра потерь времени и подтверждение интервью. "
             "Цифру McKinsey называть с источником; свои цифры из интервью сильнее — вставьте их.")

    # 3. Решение
    s = S[2]
    hint = shape_by_text(s, "Опишите суть")
    place(hint, 0.58, 1.55, 8.7, 2.85)
    fill(hint, [
        "**Суть.** Сотрудник задаёт вопрос обычным языком — агент отвечает по внутренним документам и учётной базе "
        "со ссылкой на документ, пункт и страницу, понимает уточнения в диалоге. Если ответа нет — честно говорит "
        "об этом. Всё работает внутри предприятия, без интернета.",
        "**Целевые параметры пилота:** не менее 85 % верных ответов на вопросах сотрудников партнёра, не более 3 % "
        "уверенно неверных, ответ до 10 с на сервере предприятия.",
        "**Сценарии:** нештатная ситуация в ночную смену, входной контроль, метрология, охрана труда, "
        "адаптация новичков.",
        "**Эффект:** [часы поиска в день × число сотрудников — по итогам интервью]",
    ], size=13, space=6)
    y = 4.6
    for i, (big, small) in enumerate([(RESULTS["correct"], "верных ответов со ссылкой на пункт"),
                                      (RESULTS["refusals"], "честных отказов, когда ответа нет"),
                                      (RESULTS["time_gpu"], "на ответ, без интернета")]):
        card(s, 0.58 + i * 2.95, y, 2.75, 1.35, big, small, name=f"Показатель {i + 1}")
    textbox(s, 0.58, 6.1, 8.7, 0.5, [f"Прототип, эталонный набор из {RESULTS['questions']} вопросов; "
                                     f"неверных ответов, показанных пользователю, — {RESULTS['wrong']}."],
            size=11, color=MUTED)
    notes(s, "Показать три числа и сказать, как они получены: эталонный набор, скрипт eval/run_eval.py. "
             "Честно добавить: набор использовался и для отладки, независимая проверка — на вопросах партнёра.")

    # 4. Технологическое ядро
    s = S[3]
    remove(shape_by_text(s, "Опишите технологическое ядро"))
    for pic in big_pictures(s):
        remove(pic)
    draw_architecture(s)

    # 5. Конкуренты
    s = S[4]
    hint = shape_by_text(s, "Сравните свой продукт")
    place(hint, 0.5, 1.5, 12.2, 0.8)
    fill(hint, ["Отличие не в модели, а в гарантиях: работа в контуре, ссылка на пункт, измеренный честный отказ, "
                "документы и учётная база в одном окне. Решение полностью на отечественной и открытой основе, "
                "без зарубежных сервисов."], size=13)
    table = next(sh for sh in s.shapes if sh.has_table).table
    tshape = next(sh for sh in s.shapes if sh.has_table)
    ensure_rows(table, 7)
    for i, row in enumerate(table.rows):
        row.height = Inches(0.62 if i == 0 else 0.56)
    place(tshape, 0.56, 2.45, 12.06, 0.62 + 6 * 0.56)
    header = ["Характеристика", "Наша разработка", "Яндекс «Нейроэксперт» (Россия)", "Just AI (Россия)",
              "Открытые RAG-платформы: RAGFlow, Dify (open source)"]
    rows = [
        ["Работа без интернета, в контуре предприятия", "Да", "Нет — облачный сервис", "Да — on-premise", "Да"],
        ["Ссылка на документ, пункт и страницу", "Да, с выделением на странице", "Ссылки на источники",
         "Зависит от внедрения", "Ссылки на фрагменты"],
        ["Честный отказ, измеренный на эталоне", f"Да — {RESULTS['refusals']}", "Нет открытых данных",
         "Нет открытых данных", "Нет — нужно настраивать"],
        ["Документы и учётная база в одном окне", "Да, запрос к базе только на чтение", "Нет",
         "Зависит от внедрения", "Нет"],
        ["Языковая модель", "GigaChat 3.1 Lightning — отечественная, открытая (MIT)", "YandexGPT — облачная",
         "Разные модели", "Любая открытая"],
        ["Доступность для среднего бизнеса", "Одна видеокарта или процессор", "Подписка на облако",
         "Решение уровня крупного бизнеса", "Нужна своя ИТ-команда"],
    ]
    for c, t in enumerate(header):
        set_cell(table.cell(0, c), t, size=11, bold=True, color=WHITE)
    for r, row in enumerate(rows, 1):
        for c, t in enumerate(row):
            set_cell(table.cell(r, c), t, size=11, bold=(c == 0 or c == 1), color=TEXT)
    notes(s, "Перед защитой проверьте характеристики конкурентов по их сайтам — ячейки заполнены по общедоступным "
             "сведениям на октябрь 2026. Если спросят «это же просто RAG»: RAG — техника, продукт — гарантии.")

    # 6. Монетизация (первый этап — гипотеза)
    s = S[5]
    hint = shape_by_text(s, "Опишите вашего потенциального")
    place(hint, 0.5, 1.6, 7.9, 4.9)
    fill(hint, [
        "## Потребитель",
        "- Производственные и научно-производственные предприятия малого и среднего бизнеса (50–1000 сотрудников) "
        "с большим объёмом технической и нормативной документации",
        "- Организации с режимом коммерческой тайны и закрытым контуром, которым облако недоступно",
        "- Первый пилот — индустриальный партнёр ООО «НПВП «НаноТехноПром»",
        "## Путь коммерциализации (гипотеза)",
        "- Пилот у партнёра → лицензия на установку в контур и годовое сопровождение",
        "- Внедрение как проект: подключение источников, эталонный набор вопросов, права доступа",
        "- Тиражирование через интеграторов 1С и разработчиков отраслевого ПО",
        "## Рынок",
        "- [Оценка TAM / SAM / SOM и расчёт юнит-экономики — на втором этапе]",
    ], size=13, space=3)
    notes(s, "На первом этапе монетизацию показываем как гипотезу. Подробный расчёт — второй этап.")

    # 7. Команда и прогресс
    s = S[6]
    hint = shape_by_text(s, "Укажите – ФИО")
    place(hint, 0.5, 1.6, 7.5, 5.2)
    fill(hint, [
        "## Команда",
        "- **[ФИО], [курс, направление]** — лидер, разработка прототипа: поиск, языковая модель, база данных",
        "- **[ФИО], [курс, направление]** — презентация, паспорт проекта, видео",
        "- **[ФИО], [курс, направление]** — интервью с потребителями, тестирование",
        "- **Наставник:** [ФИО]",
        "## Что сделано за первый этап",
        "- Выбрана задача индустриального партнёра № 28 и сформулирована гипотеза продукта",
        f"- Собран работающий прототип с локальным развёртыванием: {RESULTS['correct']} верных ответов, "
        f"{RESULTS['refusals']} честных отказов на эталоне из {RESULTS['questions']} вопросов",
        "- Подготовлен опросный лист; [проведено N интервью: кто и что подтвердили]",
        "- Отправлен запрос партнёру на документы, вопросы сотрудников и пилот",
    ], size=13, space=3)
    notes(s, "Показать путь команды: что было на входе (идея) и что на выходе (прототип, цифры, интервью).")

    # 8. Партнёры
    s = S[7]
    hint = shape_by_text(s, "Опишите задачу партнера")
    place(hint, 0.5, 1.6, 7.5, 5.2)
    fill(hint, [
        "## Задача партнёра",
        "ООО «НПВП «НаноТехноПром», Рязань — задача № 28: «Эффективное использование AI-агентов для качественных "
        "и точных ответов, используя внутренние базы знаний и данных предприятия».",
        "## Роль партнёра",
        "- Постановщик задачи и источник требований",
        "- Обезличенные документы и вопросы сотрудников для замеров на реальных данных",
        "- Площадка пилотного внедрения",
        "## Ближайшие шаги",
        "- Получить форму описания задачи и контакты представителя",
        "- Интервью с постановщиком задачи, сотрудниками и службой ИТ",
        "- Пилот на документах партнёра: [дата]",
        "## Другие заинтересованные",
        "- [Предприятия Рязани по итогам интервью]",
    ], size=13, space=3)
    notes(s, "Задача из списка индустриальных партнёров — это плюс к оценке, назвать явно.")

    # 9. Уровень готовности технологии
    s = S[8]
    hint = shape_by_text(s, "Определите уровень")
    place(hint, 0.5, 1.5, 8.6, 1.5)
    fill(hint, [
        "**Текущий уровень: TRL 4.** Прототип работает, характеристики подтверждены в лабораторных условиях "
        f"на эталонном наборе из {RESULTS['questions']} вопросов.",
        "**Артефакты:** видео работы прототипа; протокол испытаний (eval/results/latest.md); код в репозитории; "
        "замеры скорости на ПК с видеокартой и без неё.",
    ], size=13, space=4)
    table = next(sh for sh in s.shapes if sh.has_table).table
    for r in range(len(table.rows)):
        if table.cell(r, 0).text_frame.text.strip().startswith("4"):
            for c in range(len(table.columns)):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = ORANGE_TINT
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
    notes(s, "TRL 4 — прототип, подтверждение характеристик в лабораторных условиях. Артефакты показать: видео и "
             "таблицу результатов.")

    # 10. Прогресс за время акселерации
    s = S[9]
    hint = shape_by_text(s, "Опишите чего удалось достигнуть")
    place(hint, 0.66, 1.7, 10.5, 0.5)
    fill(hint, ["За первый этап: от идеи по задаче партнёра — к работающему прототипу с измеренным качеством."],
         size=14)
    tshape = next(sh for sh in s.shapes if sh.has_table)
    tshape.top = Inches(2.4)
    table = tshape.table
    progress = [
        ("Уровень готовности технологии — TRL 1: идея по задаче партнёра",
         f"Уровень готовности технологии — TRL 4: прототип, эталон из {RESULTS['questions']} вопросов"),
        ("Производственная готовность — нет",
         "Производственная готовность — установка одной командой, работа на ПК и ноутбуке"),
        ("Рыночная готовность — гипотезы о проблеме",
         "Рыночная готовность — [N интервью, проблема подтверждена или уточнена]"),
        ("Партнёрства — задача из списка партнёров",
         "Партнёрства — запрос партнёру отправлен, [договорённость о пилоте]"),
    ]
    for r, (before, after) in enumerate(progress, 1):
        set_cell(table.cell(r, 0), before, size=12)
        set_cell(table.cell(r, 1), after, size=12, bold=True)
    notes(s, "Строки «до» и «после» — главное доказательство динамики. Обновить после интервью.")

    # 11. Планы развития
    s = S[10]
    hint = shape_by_text(s, "Опишите ваши планы")
    place(hint, 0.5, 1.6, 8.0, 5.3)
    fill(hint, [
        "## Исследования и разработка",
        "- Пилот на документах и вопросах сотрудников НаноТехноПром",
        "- Распознавание сканов, подключение к 1С и другим учётным системам",
        "## Интеллектуальная собственность",
        "- Регистрация программы для ЭВМ в Роспатенте",
        "## Внедрение и продвижение",
        "- Кейс пилота → 3–5 предприятий Рязанской области",
        "- Партнёрство с интеграторами 1С и разработчиками отраслевого ПО",
        "## Привлечение инвестиций",
        "- Грант «Студенческий стартап» Фонда содействия инновациям",
    ], size=13, space=3)

    # 12. Запрос
    s = S[11]
    hint = shape_by_text(s, "Опишите чего сейчас не хватает")
    place(hint, 0.5, 1.6, 8.2, 5.0)
    fill(hint, [
        "- **Данные партнёра.** Обезличенные документы и реальные вопросы сотрудников НаноТехноПром — "
        "чтобы измерить качество на настоящих задачах, а не на открытых документах",
        "- **Нетворкинг.** Контакты 2–3 производственных предприятий Рязани — для интервью и подтверждения рынка",
        "- **Экспертиза.** Требования информационной безопасности и порядок включения в реестр отечественного ПО",
        "- **Финансирование.** Помощь с подготовкой заявки на грант «Студенческий стартап»",
    ], size=14, space=10)

    prs.save(out)
    _fix_content_types(out)
    return out


def build_architecture(out: Path):
    """Отдельный файл с одним слайдом «Технологическое ядро проекта»."""
    tmp = out.with_suffix(".tmp.pptx")
    shutil.copy(TEMPLATE, tmp)
    prs = Presentation(tmp)
    keep = prs.slides[3]
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.slides.part.related_part(sld_id.rId) is not keep.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
    remove(shape_by_text(keep, "Опишите технологическое ядро"))
    for pic in big_pictures(keep):
        remove(pic)
    draw_architecture(keep)
    prs.save(out)
    tmp.unlink()
    _fix_content_types(out)
    return out


def _fix_content_types(path: Path):
    """В шаблоне не объявлен тип для JPG-картинок фона — PowerPoint это прощает, валидатор нет."""
    tmp = path.with_suffix(".ct.pptx")
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        names = set(zi.namelist())
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "[Content_Types].xml" and b'Extension="jpg"' not in data \
                    and any(n.lower().endswith(".jpg") for n in names):
                data = data.replace(b"<Default ", b'<Default Extension="jpg" ContentType="image/jpeg"/><Default ', 1)
            zo.writestr(item, data)
    shutil.move(tmp, path)


if __name__ == "__main__":
    print(build_deck(HERE / "prezentatsiya_chernovik.pptx"))
    print(build_architecture(HERE / "arhitektura.pptx"))
