"""Вход, сеансы, диалоги, демо-база, каталог редакций, маршрутизация, оценка, API."""
from datetime import date

from kbagent import auth, demo_db
from kbagent.agent import Agent
from kbagent.ingest import ARCHIVE, access_of, is_cancelled


def test_password_hash_roundtrip():
    stored = auth.hash_password("тест-пароль")
    assert auth.verify_password("тест-пароль", stored) and not auth.verify_password("другой", stored)


def test_sessions_and_conversations(tmp_settings, monkeypatch):
    monkeypatch.setattr(auth, "load_users", lambda: {"u": {"login": "u", "name": "У", "role": "Метролог"}})
    token = auth.create_session({"login": "u"})
    assert auth.session_user(token)["role"] == "Метролог"
    auth.end_session(token)
    assert auth.session_user(token) is None

    conv = auth.new_conversation()
    conv["items"].append({"question": "Какой срок входного контроля?", "text": "5 дней [1]"})
    auth.save_conversation("u", conv)
    listed = auth.list_conversations("u")
    assert listed[0]["title"] == "Какой срок входного контроля?"
    auth.delete_conversation("u", conv["id"])
    assert auth.list_conversations("u") == []


def test_demo_dates_shift_with_today():
    delta = demo_db.offset(date(2026, 12, 4))
    assert demo_db.shift("2026-09-25", delta) == "2026-11-26"
    assert demo_db.shift("2026-09-30 02:45", delta) == "2026-12-01 02:45"
    assert demo_db.shift("ОБ-0101", delta) == "ОБ-0101"


def test_cancelled_edition_is_archived():
    assert is_cancelled({"status": "отменён"}) and access_of({"status": "отменён", "access": "all"}) == ARCHIVE
    assert access_of({"access": "metrology"}) == "metrology"


def test_routing_rules_without_model():
    agent = Agent.__new__(Agent)          # без загрузки индекса: проверяются правила до обращения к модели
    assert agent.route("Можно ли выдать в производство партию П-2609-04?") == "both"
    assert agent.route("Что делать при искрении оборудования?") == "documents"
    assert agent.route("В какой срок приступит ремонтная служба?") == "documents"
    assert agent.route("Когда поверка СИ-0018?") == "database"
    assert agent.route("Сколько стоит поверка одного осциллографа?") == "documents"
    assert agent.route("скока дней длится входной контроль эри") == "documents"


def test_eval_helpers():
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location("run_eval", Path(__file__).parents[1] / "eval" / "run_eval.py")
    ev = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ev)
    assert ev.clause_covers("4.1–4.3", "4.2") and ev.clause_covers("3.1–3.4", "3") and not ev.clause_covers("5.1", "4.2")
    assert "12.11.2026" in ev.expand_dates("date:2026-11-12", demo_db.offset(date(2026, 10, 3)))
    assert ev.has_facts("Срок — 5 рабочих дней", ["5 рабочих дней"]) and ev.has_none("2211", ["2200"])


def test_api_requires_key(tmp_path, monkeypatch, tmp_settings):
    from fastapi.testclient import TestClient

    from api import server
    monkeypatch.setattr(server, "KEYS_PATH", tmp_path / "keys.yaml")
    (tmp_path / "keys.yaml").write_text(
        f"keys:\n- name: t\n  role: Метролог\n  hash: {server.key_hash('kb_good')}\n", encoding="utf-8")
    api = TestClient(server.app)
    assert api.post("/ask", json={"question": "тест вопроса"}).status_code == 401
    assert api.post("/ask", json={"question": "тест вопроса"}, headers={"X-API-Key": "bad"}).status_code == 401
    assert server.client("kb_good")["role"] == "Метролог"


def test_named_documents_and_rule_routing():
    from kbagent.agent import named_docs
    cards = {"g": {"title": "ГОСТ Р 7.0.97-2016 Организационно-распорядительная документация"},
             "s": {"title": "СТП ДП-01-2026 «Порядок входного контроля»"},
             "o": {"title": "Осциллографы цифровые АКИП-4115/1В, АКИП-4115/2В"},
             "m": {"title": "Мультиметр цифровой MASTECH MAS830L"}}
    assert named_docs("какие поля у документа по госту 7.0.97", cards) == {"g"}
    assert named_docs("Что в пункте 3.2 СТП ДП-01-2026?", cards) == {"s"}
    assert named_docs("Предохранитель в MAS830L?", cards) == {"m"}
    assert named_docs("Какой срок входного контроля?", cards) == set()
    agent = Agent.__new__(Agent)
    assert agent.route("Пришла партия 450 резисторов. Сколько штук нужно осмотреть?") == "documents"


def test_source_preview_marks_whole_fragment_across_lines(tmp_path):
    import pymupdf
    from kbagent import preview
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "2.1 Equipment failures are divi-")
    page.insert_text((72, 90), "ded into three categories. Other text below.")
    path = tmp_path / "d.pdf"
    doc.save(path)
    png, marked = preview.render_page(str(path), 1, "2.1 Equipment failures are divided into three categories.",
                                      crop=True)
    assert marked == 9 and png[:4] == b"\x89PNG"     # перенос «divi-/ded» склеен, «Other text below» не выделен


def test_continuation_of_long_clause_gets_its_own_pages():
    from kbagent import ingest
    sentence = "Работник выполняет требования инструкции по охране труда при работе на оборудовании участка. "
    pages = [(1, ["1 Общие положения", "1.1 " + sentence * 6]), (2, [sentence * 6]), (3, [sentence * 6])]
    meta = {"id": "doc", "title": "Документ", "kind": "Инструкция", "access": "all"}
    chunks = ingest._chunks(ingest._units(pages), meta)
    assert len(chunks) >= 3 and chunks[0].page_start == 1
    assert chunks[-1].page_start >= 2 and chunks[-1].page_end == 3     # не «стр. 1–3» у последнего куска


def test_uploaded_document_title_is_taken_from_its_text(tmp_path):
    from docx import Document
    from kbagent.uploads import guess_title
    doc = Document()
    for line in ["ООО «Демо-Прибор»", "РИ ДП-05-2026", "ОБСЛУЖИВАНИЕ УСТАНОВКИ СЕЛЕКТИВНОЙ ПАЙКИ", "1.1 Текст пункта."]:
        doc.add_paragraph(line)
    path = tmp_path / "RI-DP-05.docx"
    doc.save(path)
    assert guess_title(path) == "РИ ДП-05-2026 «Обслуживание установки селективной пайки»"


def test_unanswered_questions_are_grouped_into_topics():
    from kbagent.gaps import find_gaps
    rec = lambda q, status, user="petrov", t="2026-10-05 10:00:00": {  # noqa: E731
        "time": t, "user": user, "role": "Инженер производства", "question": q, "status": status, "route": "documents"}
    records = [rec("Сколько стоит поверка осциллографа?", "no_data"),
               rec("Какая стоимость поверки осциллографа?", "no_data", user="ivanov", t="2026-10-05 11:00:00"),
               rec("Где взять пропуск на склад?", "model_refused"),
               rec("Что делать при искрении?", "answer"),
               rec("Сколько стоит поверка осциллографа?", "no_data", user="eval")]       # прогон эталона — не в счёт
    topics, total = find_gaps(records)
    assert total == {"questions": 4, "unanswered": 3, "topics": 2, "closed": 0}
    assert topics[0]["count"] == 2 and topics[0]["users"] == 2 and topics[0]["last"] == "2026-10-05 11:00:00"


def test_gap_closes_when_the_same_question_is_answered_later():
    from kbagent.gaps import find_gaps
    rec = lambda q, status, t, user="petrov", role="Инженер производства": {  # noqa: E731
        "time": t, "user": user, "role": role, "question": q, "status": status, "route": "documents"}
    records = [rec("Какая температура припоя в ванне селективной пайки?", "no_data", "2026-10-05 10:00:00"),
               rec("Где взять пропуск на склад?", "no_data", "2026-10-05 10:05:00"),
               # загрузили регламент — тот же вопрос получил ответ (в том числе при проверке администратором)
               rec("Какая температура припоя в ванне селективной пайки?", "answer", "2026-10-05 12:00:00", user="eval"),
               # ответ на более общий вопрос или под другой ролью пробел не закрывает
               rec("Где склад?", "answer", "2026-10-05 12:00:00"),
               rec("Где взять пропуск на склад?", "answer", "2026-10-05 12:00:00", role="Метролог"),
               rec("Где взять пропуск на склад?", "answer", "2026-10-05 12:00:00", role="Администратор"),
               # случайный ответ, после которого снова нет ответа, — пробел открыт
               rec("Где взять пропуск на склад?", "answer", "2026-10-05 13:00:00"),
               rec("Где взять пропуск на склад?", "no_data", "2026-10-05 14:00:00", user="eval")]
    topics, total = find_gaps(records)
    assert total["topics"] == 1 and total["closed"] == 1
    assert topics[0]["question"] == "Где взять пропуск на склад?" and not topics[0]["closed"]
    # ответ есть, но для другой роли — права доступа; администратор видит всё и подсказкой не служит
    assert topics[0]["other_roles"] == ["Метролог"]
    assert topics[1]["closed"] == "2026-10-05 12:00:00"


def test_actions_journal_records_who_did_what(tmp_settings):
    from kbagent.actions import load_actions, log_action
    log_action("login_failed", "ivanov")
    log_action("upload", "admin", "РИ ДП-05-2026", "фрагментов: 7")
    records = load_actions()
    assert [r["action"] for r in records] == ["login_failed", "upload"]
    assert records[1]["user"] == "admin" and records[1]["target"] == "РИ ДП-05-2026"


def test_personal_data_is_masked_but_work_details_stay(tmp_settings):
    from kbagent.privacy import mask
    text = ("Иванов Пётр Сергеевич, тел. +7 (912) 345-67-89, почта p.ivanov@zavod.ru, паспорт 4510 123456, "
            "СНИЛС 112-233-445 95. Ответственный — Белов И. Н., звонить диспетчеру 2211.")
    masked = mask(text)
    for secret in ["Пётр Сергеевич", "345-67-89", "p.ivanov", "123456", "112-233-445"]:
        assert secret not in masked
    assert "Белов И. Н." in masked and "2211" in masked


def test_names_are_masked_in_any_case_and_order(tmp_settings):
    from kbagent.privacy import mask
    # примеры участника 2 (docs/uchastnik2/proverki/t_privacy.py): косвенные падежи, порядок «И О Ф», регистр
    leaks = {"Передайте Белову Игорю Николаевичу, тел. 8 (4912) 12-34-56: где флюс?": ["Николаевичу", "12-34-56"],
             "Передайте Сидоровой Анне Петровне акт": ["Сидоровой", "Петровне"],
             "Иван Петрович Сидоров уволен?": ["Сидоров"], "Кузьмин Фома Ильич заболел": ["Кузьмин", "Ильич"],
             "ИВАНОВ ИВАН ИВАНОВИЧ": ["ИВАНОВИЧ"], "Римская-Корсакова Мария Андреевна": ["Римская"],
             "СНИЛС 12345678901, паспорт серия 6123 номер 456789": ["12345678901", "456789"]}
    for text, secrets in leaks.items():
        assert not any(s in mask(text) for s in secrets), text
    # рабочие сведения и похожие на отчество прилагательные не трогаются
    for text in ["Кому сообщает начальник ночной смены?", "Это основной документ", "Ответственный Белов И. Н.",
                 "партия П-2609-04, 120 шт, поступила 02.10.2026", "телефон диспетчера 2211"]:
        assert mask(text) == text, text


def test_login_is_locked_after_failed_attempts(tmp_settings, tmp_path, monkeypatch):
    import yaml
    from kbagent import auth
    users = tmp_path / "users.yaml"
    users.write_text(yaml.safe_dump({"users": [{"login": "ivanov", "name": "Иванов", "role": "Инженер производства",
                                                "password_hash": auth.hash_password("верный")}]}), encoding="utf-8")
    monkeypatch.setattr(tmp_settings, "users_path", users)
    assert auth.authenticate("ivanov", "верный")
    for _ in range(auth.LOCK_AFTER):
        assert auth.authenticate("ivanov", "неверный") is None
    assert auth.locked_for("ivanov") > 0
    assert auth.authenticate("ivanov", "верный") is None          # подбор остановлен и с верным паролем
    assert auth.locked_for("petrov") == 0                         # другие логины не затронуты


def test_old_journal_records_are_pruned(tmp_settings, tmp_path):
    import json
    from kbagent.privacy import prune
    log = tmp_path / "audit.jsonl"
    log.write_text("".join(json.dumps({"time": t, "question": "q"}) + "\n"
                           for t in ["2020-01-01 10:00:00", "2099-01-01 10:00:00"]), encoding="utf-8")
    assert prune(log, days=180) == 1
    assert "2099" in log.read_text(encoding="utf-8") and "2020" not in log.read_text(encoding="utf-8")


def test_normocontrol_finds_formatting_errors(tmp_path):
    import sys
    from conftest import ROOT
    sys.path.insert(0, str(ROOT / "scripts"))
    from make_show_docs import order_docx
    from kbagent.normocontrol import check_docx
    bad, good = tmp_path / "bad.docx", tmp_path / "good.docx"
    order_docx(bad, correct=False)
    order_docx(good, correct=True)
    issues = check_docx(bad.read_bytes()).issues
    assert {i.clause for i in issues} == {"6.1", "3.3", "3.4", "3.5", "5.17", "5.10"}
    assert check_docx(good.read_bytes()).issues == []


def test_draft_edition_does_not_cancel_current_until_approved(tmp_path):
    import yaml
    from kbagent.ingest import ARCHIVE, DRAFT, access_of, load_catalog
    raw = tmp_path / "raw"
    (raw / "uploaded").mkdir(parents=True)
    (tmp_path / "sources.yaml").write_text(yaml.safe_dump({"documents": [{"id": "stp_02", "title": "СТП ДП-02"}]}),
                                           encoding="utf-8")
    draft = {"id": "upl_stp_02_new", "title": "СТП ДП-02 (проект)", "replaces": "stp_02", "approved": False}
    (raw / "uploaded" / "catalog.yaml").write_text(yaml.safe_dump({"documents": [draft]}), encoding="utf-8")
    cat = load_catalog(raw)
    # проект с ошибкой не подменяет действующий регламент: сотрудники по-прежнему видят старую редакцию
    assert access_of(cat["stp_02"]) == "all" and access_of(cat["upl_stp_02_new"]) == DRAFT
    (raw / "uploaded" / "catalog.yaml").write_text(yaml.safe_dump({"documents": [{**draft, "approved": True}]}),
                                                   encoding="utf-8")
    cat = load_catalog(raw)
    assert access_of(cat["stp_02"]) == ARCHIVE and cat["stp_02"]["replaced_by"] == "upl_stp_02_new"
    assert access_of(cat["upl_stp_02_new"]) == "all"


def test_gap_topics_are_not_glued_and_drafts_do_not_close_them():
    from kbagent.gaps import find_gaps
    rec = lambda q, status, t, **kw: {"time": t, "user": "petrov", "role": "Инженер производства",  # noqa: E731
                                      "question": q, "status": status, "route": "documents", **kw}
    topics, total = find_gaps([rec("Сколько стоит поверка осциллографа?", "no_data", "2026-10-05 10:00:00"),
                               rec("Сколько стоит ремонт осциллографа?", "no_data", "2026-10-05 10:01:00"),
                               rec("Какая цена поверки осциллографа?", "no_data", "2026-10-05 10:02:00")])
    assert total["topics"] == 2           # «цена поверки» — к «стоит поверка», «ремонт» — отдельно
    q = "Какая температура пайки в селективной ванне?"
    topics, total = find_gaps([rec(q, "no_data", "2026-10-05 10:00:00"),
                               rec(q, "answer", "2026-10-05 11:00:00", unapproved=True)])
    assert total == {"questions": 2, "unanswered": 1, "topics": 1, "closed": 0}


def test_registry_dates_control_access_and_labels(tmp_path):
    import time
    import yaml
    from kbagent.ingest import DRAFT, access_of, expired_on, load_catalog, registry_suffix
    raw = tmp_path / "raw"
    (raw / "internal").mkdir(parents=True)
    future = time.strftime("%Y") + "9-01-01"          # заведомо будущая дата
    docs = [{"id": "old", "title": "Приказ 903н", "effective_from": "2021-01-01", "effective_to": "2025-12-31"},
            {"id": "soon", "title": "СТП ДП-09", "version": "2027", "effective_from": future, "access": "all"},
            {"id": "now", "title": "СТП ДП-01", "version": "2026", "effective_from": "2026-03-01"}]
    (tmp_path / "sources.yaml").write_text(yaml.safe_dump({"documents": docs}, allow_unicode=True), encoding="utf-8")
    cat = load_catalog(raw)
    assert expired_on(cat["old"]) == "2025-12-31" and access_of(cat["old"]) == "all"   # не прячем, а предупреждаем
    assert access_of(cat["soon"]) == DRAFT                       # не вступил в силу — только администратору
    assert registry_suffix("now", raw) == "ред. 2026, действует с 01.03.2026"
    assert registry_suffix("old", raw) == "срок действия истёк 31.12.2025"


def test_api_rejects_oversized_history_and_hides_internal_errors(tmp_path, monkeypatch, tmp_settings):
    from fastapi.testclient import TestClient

    from api import server
    monkeypatch.setattr(server, "KEYS_PATH", tmp_path / "keys.yaml")
    (tmp_path / "keys.yaml").write_text(
        f"keys:\n- name: t\n  role: Метролог\n  hash: {server.key_hash('kb_good')}\n", encoding="utf-8")
    api, key = TestClient(server.app), {"X-API-Key": "kb_good"}
    long = [{"question": "вопрос", "answer": "о" * 9000}]
    assert api.post("/ask", json={"question": "тест вопроса", "history": long}, headers=key).status_code == 422
    many = [{"question": "вопрос", "answer": "ответ"}] * 60
    assert api.post("/ask", json={"question": "тест вопроса", "history": many}, headers=key).status_code == 422

    class Broken:
        def ask(self, *a, **k):
            raise RuntimeError("сервер векторов: слишком длинный текст")
    monkeypatch.setattr(server, "agent", lambda: Broken())
    r = api.post("/ask", json={"question": "тест вопроса"}, headers=key)
    assert r.status_code == 503 and "недоступен" in r.json()["detail"]


def test_employee_can_send_any_answer_to_expert_and_see_the_decision(tmp_settings):
    from kbagent import experts
    first = experts.add("передал сотрудник", "Какая полоса у АКИП-4115/2В?", "Ответ агента: 200 МГц [1]",
                        user="u1", role="Техпродажи и сервис", comment="в каталоге 150 МГц")
    again = experts.add("передал сотрудник", "Какая полоса у АКИП-4115/2В?", "Ответ агента: 200 МГц [1]",
                        user="u2", comment="клиент тоже спрашивает")
    assert again["id"] == first["id"] and again["asked"] == 2 and "клиент" in again["comment"]
    experts.resolve(first["id"], "Эксперт", "верно 200 МГц, каталог устарел")
    done = experts.get(first["id"])
    assert done["status"] == "решён" and done["note"].startswith("верно")
    assert experts.add("передал сотрудник", "Какая полоса у АКИП-4115/2В?", "…")["id"] != first["id"]


def test_upload_converts_images_and_excel_tables(monkeypatch):
    import datetime
    import io

    import pymupdf
    import pytest
    from docx import Document
    from openpyxl import Workbook

    from kbagent import uploads
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 800, 1100), 0)
    pix.clear_with(255)
    monkeypatch.setattr(uploads, "ocr_available", lambda: True)
    ext, data, notes = uploads.convert_upload("Скан приказа.JPG", pix.tobytes("jpg"))
    page = pymupdf.open("pdf", data)[0]
    assert ext == ".pdf" and notes == {"image": True} and page.get_images() and round(page.rect.width) == 595

    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Модель", "Цена, руб", "Дата"])
    ws.append(["АКИП-4115/1В", 45000.5, datetime.date(2026, 10, 1)])
    ws.append([None, None, None])
    ws.append(["MAS830L", 1200, None])
    buf = io.BytesIO()
    wb.save(buf)
    ext, data, notes = uploads.convert_upload("Прайс_лист.xlsx", buf.getvalue())
    lines = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert ext == ".docx" and notes == {"table_rows": 2}
    assert "Лист «Прайс», строка 2. Модель: АКИП-4115/1В; Цена, руб: 45000,5; Дата: 01.10.2026" in lines
    assert "Лист «Прайс», строка 4. Модель: MAS830L; Цена, руб: 1200" in lines

    monkeypatch.setattr(uploads, "ocr_available", lambda: False)
    with pytest.raises(ValueError, match="распознавание"):
        uploads.convert_upload("фото.png", pix.tobytes("png"))
    with pytest.raises(ValueError, match="Excel"):
        uploads.convert_upload("заметки.txt", b"text")


def test_every_demo_user_has_a_known_role():
    from kbagent.config import settings
    from scripts.make_users import DEMO_USERS
    assert {role for _, _, role, _ in DEMO_USERS} <= set(settings.roles)
    assert "Техпродажи и сервис" in settings.roles and list(settings.roles)[-1] == "Администратор"


def test_effect_counts_time_to_approval_and_flags_stale_answers(tmp_settings, tmp_path, monkeypatch):
    from kbagent import effect, experts
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    monkeypatch.setattr(effect, "PATH", tmp_path / "effect.jsonl")
    src = [{"doc_id": "osc_akip_4115v", "page": 6}]
    effect.record("ответ", "u", "Техпродажи и сервис", "Какая частота?", 90, src)
    effect.record("ответ", "u", "Техпродажи и сервис", "Полоса?", 150, [{"doc_id": "osc_akip_4115v", "page": 7}])
    effect.record("ТКП", "u", "Техпродажи и сервис", "ТКП: …", 300, [], positions=3)
    s = effect.summary()
    assert s["kinds"]["ответ"]["count"] == 2 and s["kinds"]["ответ"]["median_min"] == 2.0
    assert s["kinds"]["ТКП"]["share"] == 5 / 60 and round(s["saved_hours"], 2) == round((13.5 + 12.5 + 55) / 60, 2)
    assert not [r for r in effect.affected() if r["question"] == "Какая частота?"]
    it = experts.add("противоречие в справочнике", "АКИП-4115/1В — Частота дискретизации на канал",
                     [{"value": "1 Гвыб/с", "doc": "osc_akip_4115v", "page": 63},
                      {"value": "1 Гц/с", "doc": "osc_akip_4115v", "page": 6}])
    experts.resolve(it["id"], "Эксперт", "на с. 6 опечатка", {"value": "1 Гвыб/с", "doc": "osc_akip_4115v",
                                                              "page": 63})
    stale = [r for r in effect.affected() if r["question"] == "Какая частота?"]
    assert stale and "эксперт исправил значение" in stale[0]["reasons"][0]
    assert not [r for r in effect.affected() if r["question"] == "Полоса?"]


def test_model_requests_carry_api_key(tmp_path, monkeypatch):
    from kbagent import llm
    from kbagent.config import settings

    sent = []

    class Reply:
        status_code = 200

        @staticmethod
        def json():
            return {"ok": True}

    def fake_post(url, json, headers, timeout):
        sent.append(headers)
        return Reply()

    monkeypatch.setattr(llm.requests, "post", fake_post)
    monkeypatch.delenv("KB_MODEL_KEY", raising=False)
    monkeypatch.setattr(settings, "model_key_path", tmp_path / "model_api_key.txt")
    llm._post("http://127.0.0.1:8080/v1/x", {})
    assert sent[-1] == {}                                   # ключа нет — серверы запущены без него
    (tmp_path / "model_api_key.txt").write_text("abc123\n", encoding="utf-8")
    llm._post("http://127.0.0.1:8080/v1/x", {})
    assert sent[-1] == {"Authorization": "Bearer abc123"}   # ключ из файла, без перевода строки
    monkeypatch.setenv("KB_MODEL_KEY", "enterprise")
    llm._post("http://127.0.0.1:8080/v1/x", {})
    assert sent[-1] == {"Authorization": "Bearer enterprise"}


def test_rejected_model_key_is_explained(monkeypatch):
    import pytest
    from kbagent import llm

    class Reply:
        status_code = 401
        text = "Invalid API Key"

    monkeypatch.setattr(llm.requests, "post", lambda *a, **k: Reply())
    with pytest.raises(llm.ModelError, match="ключ доступа"):
        llm._post("http://127.0.0.1:8080/v1/x", {})


def test_api_builds_tkp_from_letter_without_models(tmp_path, monkeypatch, tmp_settings):
    import base64
    import io

    from docx import Document
    from fastapi.testclient import TestClient

    from api import server
    from kbagent import experts, llm
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    monkeypatch.setattr(server, "KEYS_PATH", tmp_path / "keys.yaml")
    monkeypatch.setattr(llm, "_post", lambda *a, **k: (_ for _ in ()).throw(AssertionError("модель не нужна")))
    (tmp_path / "keys.yaml").write_text(
        f"keys:\n- name: crm\n  role: Техпродажи и сервис\n  hash: {server.key_hash('kb_crm')}\n", encoding="utf-8")
    api, key = TestClient(server.app), {"X-API-Key": "kb_crm"}
    letter = "Просим КП: 2 осциллографа от 150 МГц на 4 канала, мультиметр на ток до 15 А — 3 шт."
    assert api.post("/tkp", json={"letter": letter}).status_code == 401
    r = api.post("/tkp", json={"letter": letter, "signer": "Иванов И. И."}, headers=key).json()
    assert r["status"] == "answer" and [p["model"] for p in r["positions"]] == ["АКИП-4143/1", "DT9205A"]
    osc = r["positions"][0]
    assert osc["fit_models"][0] == "АКИП-4143/1" and all(c["ok"] for c in osc["checks"]) and osc["checks"][0]["source"]
    assert r["total"] == sum(p["sum"] for p in r["positions"]) and r["vat"] and "Иванов И. И." in r["letter_draft"]
    doc = Document(io.BytesIO(base64.b64decode(r["docx_base64"])))
    assert doc.tables[0].rows[1].cells[1].text == "Осциллограф АКИП-4143/1"
    # выбор менеджера и письмо файлом
    again = api.post("/tkp", headers=key, json={"choices": {"1": "АКИП-4143/2"}, "include_docx": False, "file": {
        "name": "zapros.txt", "content_base64": base64.b64encode(letter.encode("utf-8")).decode()}}).json()
    assert again["positions"][0]["model"] == "АКИП-4143/2" and again["docx_base64"] is None
    assert api.post("/tkp", json={"letter": "Добрый день, как ваши дела?"}, headers=key).json()["status"] == "not_request"
    assert api.post("/tkp", json={"file": {"name": "a.txt", "content_base64": "не base64"}}, headers=key).status_code == 422
