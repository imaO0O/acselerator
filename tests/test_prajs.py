"""Прайс для ТКП: проверка при загрузке, цена и сумма в разборе письма, в письме клиенту и в таблице Word."""
import io
from datetime import date, timedelta

import pytest

from kbagent import experts, prajs, tkp

LETTER = "ТКП: 2 осциллографа от 150 МГц на 4 канала, мультиметр на ток до 15 А — 3 шт., паяльная станция на 450 °C"


def write(path, rows, head="Модель;Цена;Срок поставки;Действует с;Примечание"):
    path.write_text("\n".join([head] + rows) + "\n", encoding="utf-8-sig")
    return path


def test_price_list_is_checked_on_import(tmp_path):
    f = write(tmp_path / "p.csv", [
        "акип-4143/1;189 000 ₽;в наличии;01.10.2026;",          # имя приводится к справочнику
        "DT9205A;1450,50;;2026-10-01;",                        # копейки и дата из Excel
        "DT9205A;1450,50;;2026-10-01;",                        # повтор — не дублируется
        "MAS830L;бесплатно;;01.10.2026;",                      # не число
        "SBK936B;;;01.10.2026;",                               # пропуск
        "CXG DES H95;7800;;32.10.2026;",                       # не дата
        "Lukey 852D+;6300;;01.10.2026;",
        "Lukey 852D+;6900;;01.10.2026;",                       # две цены одной модели
        "Неизвестный 1;100;;01.10.2026;",                      # нет в справочнике — с предупреждением
    ])
    known = {"АКИП-4143/1", "DT9205A", "MAS830L", "SBK936B", "CXG DES H95", "Lukey 852D+"}
    prices, rep = prajs.parse(f, known)
    by = {p.model: p for p in prices}
    assert set(by) == {"АКИП-4143/1", "DT9205A", "Неизвестный 1"}
    assert by["АКИП-4143/1"].price == 189000 and by["DT9205A"].price == 1450.5 and by["DT9205A"].since == "2026-10-01"
    assert rep.duplicates == 1 and rep.unknown == ["Неизвестный 1"]
    assert [e["line"] for e in rep.errors] == [5, 6, 7]
    assert rep.conflicts[0]["model"] == "Lukey 852D+"
    assert prajs.money(1450.5) == "1\u00a0450,50\u00a0₽" and prajs.money(378000) == "378\u00a0000\u00a0₽"


def test_replace_keeps_old_price_list_when_nothing_passes(tmp_path, monkeypatch):
    monkeypatch.setattr(prajs, "PATH", tmp_path / "prajs.csv")
    rep = prajs.replace_file("new.csv", "Модель;Цена\nАКИП-4143/1;189000\n".encode("utf-8"))
    assert rep.prices == 1
    stored = prajs.PATH.read_text(encoding="utf-8-sig")
    assert date.today().strftime("%d.%m.%Y") in stored          # без «Действует с» — дата загрузки
    with pytest.raises(ValueError):
        prajs.replace_file("bad.csv", "Модель;Цена\nАКИП-4143/1;—\n".encode("utf-8"))
    assert prajs.PATH.read_text(encoding="utf-8-sig") == stored  # действующий прайс остался


def test_tkp_takes_price_and_sum_from_price_list(tmp_path, monkeypatch):
    from docx import Document
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    today = date.today().strftime("%d.%m.%Y")
    monkeypatch.setattr(prajs, "PATH", write(tmp_path / "prajs.csv", [
        f"АКИП-4143/1;189000;в наличии;{today};ДЕМО",
        f"DT9205A;1450;10 рабочих дней;{today};ДЕМО",          # паяльной станции в прайсе нет
    ]))
    t = tkp.parse(LETTER, {"all"})
    text = tkp.answer_text(t)
    assert f"Цена по прайсу от {today} (ДЕМО): {prajs.money(189000)} × 2 шт. = {prajs.money(378000)}" in text
    assert "В прайсе цены нет — впишет менеджер" in text            # CXG DES H95
    assert f"**Итого по прайсу:** {prajs.money(378000 + 4350)} — позиций с ценой 2 из 3" in text
    d = t.to_dict()
    draft = tkp.draft_text(d, "Сотрудник")
    assert f"АКИП-4143/1 — 2 шт. Цена: {prajs.money(189000)} за шт., сумма {prajs.money(378000)}." in draft
    assert "CXG DES H95 — 1 шт. Цена: [заполняет менеджер]." in draft
    assert "поз. 1 — в наличии; поз. 2 — 10 рабочих дней; поз. 3 — [заполняет менеджер]" in draft
    doc = Document(io.BytesIO(tkp.tkp_docx(draft, d, "Сотрудник", f"{today} 12:00", "Агент")))
    rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert rows[1][4:] == [prajs.money(189000), prajs.money(378000)] and rows[3][4:] == ["", ""]
    assert rows[-1][1] == "Итого по прайсу" and rows[-1][5] == prajs.money(382350)
    note = doc.paragraphs[-1].text
    assert "по прайсу от" in note and "ДЕМО" in note and "чего нет в прайсе, вписывает менеджер" in note


def test_old_price_list_is_flagged(tmp_path, monkeypatch):
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    old = (date.today() - timedelta(days=prajs.STALE_DAYS + 10)).strftime("%d.%m.%Y")
    monkeypatch.setattr(prajs, "PATH", write(tmp_path / "prajs.csv", [f"АКИП-4143/1;189000;в наличии;{old};"]))
    text = tkp.answer_text(tkp.parse(LETTER, {"all"}))
    assert f"прайс старше {prajs.STALE_DAYS} дней — проверьте цену" in text


def test_demo_price_list_loads_clean():
    prices, rep = prajs.load()
    assert len(prices) >= 20 and not rep.errors and not rep.conflicts and not rep.unknown
    assert all(p.demo for p in prices.values())
