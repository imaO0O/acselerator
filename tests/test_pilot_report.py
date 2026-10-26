"""Отчёт о пилоте: цифры за период из журналов, цель «минус 50 %», без персональных данных."""
import io
import json
from datetime import date

from docx import Document

from kbagent import effect, experts, pilot_report, tkp_registry


def _write(path, records):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in records) + "\n", encoding="utf-8")


def test_report_counts_period_and_hides_people(tmp_settings, tmp_path, monkeypatch):
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    monkeypatch.setattr(tkp_registry, "DIR", tmp_path / "tkp")
    logs = tmp_settings.log_dir
    q = lambda t, user, role, status, route="documents", cited=("СТП ДП-02, п. 3.1",), s=1.2: {  # noqa: E731
        "time": t, "user": user, "role": role, "question": "Сколько действует наряд-допуск?", "route": route,
        "status": status, "seconds": s, "cited": list(cited) if status == "answer" else []}
    _write(logs / "audit.jsonl", [
        q("2026-11-02 10:00:00", "ivanov", "Инженер производства", "answer"),
        q("2026-11-02 10:05:00", "ivanov", "Инженер производства", "answer", cited=()),
        q("2026-11-03 11:00:00", "petrova", "Метролог", "no_data"),
        q("2026-11-03 11:01:00", "petrova", "Метролог", "answer", route="spravochnik", s=0.1),
        q("2026-11-04 12:00:00", "eval", "Администратор", "answer"),                 # проверочный прогон
        q("2026-10-20 12:00:00", "ivanov", "Инженер производства", "answer"),        # до периода
    ])
    _write(logs / "feedback.jsonl", [{"time": "2026-11-02 10:01:00", "user": "ivanov", "score": "👍"},
                                     {"time": "2026-11-03 11:02:00", "user": "petrova", "score": "👎"}])
    _write(logs / "effect.jsonl", [
        {"time": "2026-11-02 10:03:00", "kind": "ответ", "user": "ivanov", "seconds": 180, "question": "…"},
        {"time": "2026-11-05 09:00:00", "kind": "ТКП", "user": "sidorov", "seconds": 2400, "question": "…"}])
    data = pilot_report.collect(date(2026, 11, 1), date(2026, 11, 30))
    assert (data["users"], data["questions"], data["answered"], data["cited"], data["refused"]) == (2, 4, 3, 2, 1)
    assert (data["up"], data["down"]) == (1, 1)
    assert dict(data["roles"]) == {"Инженер производства": [2, 2], "Метролог": [2, 1]}
    goal = pilot_report.goal_lines(data["effect"])
    assert "3,0 мин против 15 мин вручную (−80 %)" in goal[0] and "цель достигнута" in goal[0]
    assert "40,0 мин против 60 мин вручную (−33 %)" in goal[1] and "пока не достигнута" in goal[1]
    doc = Document(io.BytesIO(pilot_report.report_docx(data, "ООО «Заказчик»", "AI-агент")))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows
                                                                  for c in r.cells)
    assert "Заказчик: ООО «Заказчик»" in text and "Период: 01.11.2026 — 30.11.2026" in text
    assert "3 (75 %)" in text and "2 (67 %)" in text and "Заключение заказчика" in text
    assert "ivanov" not in text and "petrova" not in text                       # логины в отчёт не попадают
    assert effect.BASELINE_MIN["ответ"] == 15


def test_eval_counts_survive_pipes_in_letters(tmp_path):
    md, csv_path = tmp_path / "tkp.md", tmp_path / "tkp.csv"
    md.write_text("# Результаты (2026-10-26 19:30)\n\nвопросов: 2\n"
                  "| K14 | письма | Наименование | Количество | верно | answer |\n", encoding="utf-8")
    csv_path.write_text("id,correct,question\nK14,True,a | b\nK15,True,c\n", encoding="utf-8-sig")
    assert pilot_report.eval_counts(md) == (2, 2)
