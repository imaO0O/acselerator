"""Реестр ТКП: номер при утверждении, бланк компании, срок действия и условия, Excel, статусы и сводка."""
import datetime as dt
import io

import pytest
from docx import Document
from openpyxl import load_workbook

from kbagent import experts, prajs, tkp, tkp_registry

LETTER = "Добрый день! ООО «Пример» просит КП: 2 осциллографа от 150 МГц на 4 канала, мультиметр на ток до 15 А — 3 шт."


@pytest.fixture
def reg(tmp_path, monkeypatch, tmp_settings):
    d = tmp_path / "tkp"
    for name, value in {"DIR": d, "SETTINGS": d / "settings.yaml", "BLANK": d / "blank.docx",
                        "REGISTRY": d / "registry.json", "FILES": d / "files"}.items():
        monkeypatch.setattr(tkp_registry, name, value)
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    prices = tmp_path / "prajs.csv"
    prices.write_text("Модель;Цена;Срок поставки;Действует с\nАКИП-4143/1;189000;в наличии;01.10.2026\n"
                      "DT9205A;1450;в наличии;01.10.2026\n", encoding="utf-8-sig")
    monkeypatch.setattr(prajs, "PATH", prices)
    return tkp_registry


def blank_bytes() -> bytes:
    doc = Document()
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].header.paragraphs[0].text = "ООО «Ромашка» · ТКП № {{номер}} от {{дата}}"
    for text in ("Шапка бланка: реквизиты", "{{ТКП}}", "Подвал бланка: банковские реквизиты"):
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def approve(reg, now):
    t = tkp.parse(LETTER, {"all"}).to_dict()
    text = tkp.draft_text(t, "Менеджер")
    return reg.approve(t, text, "m", "Менеджер", "Техпродажи и сервис", LETTER, "Агент", now=now)


def test_approved_tkp_gets_number_and_company_blank(reg):
    assert "место метки" in reg.save_blank(blank_bytes())
    with pytest.raises(ValueError):
        reg.save_blank(b"not a docx")
    reg.save_settings({"company": "ООО «Ромашка»", "validity_days": 30, "payment": "100 % предоплата",
                       "delivery": "самовывоз, Рязань"})
    first = approve(reg, dt.datetime(2026, 10, 26, 15, 30))
    second = approve(reg, dt.datetime(2026, 10, 27, 9, 0))
    third = approve(reg, dt.datetime(2027, 1, 10, 9, 0))
    assert [first["number"], second["number"], third["number"]] == ["ТКП-2026-0001", "ТКП-2026-0002", "ТКП-2027-0001"]
    assert first["client"] == "ООО «Пример»" and first["models"] == ["АКИП-4143/1", "DT9205A"]
    assert first["total"] == 189000 * 2 + 1450 * 3 and first["valid_until"] == "25.11.2026"

    doc = Document(reg.files("ТКП-2026-0001")[0])
    assert doc.sections[0].header.paragraphs[0].text == "ООО «Ромашка» · ТКП № ТКП-2026-0001 от 26.10.2026"
    body = [p.text for p in doc.paragraphs if p.text.strip()]
    assert body[0] == "Шапка бланка: реквизиты" and body[-1] == "Подвал бланка: банковские реквизиты"
    assert body[1] == "ТЕХНИКО-КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ" and body[2] == "№ ТКП-2026-0001 от 26.10.2026"
    assert "Предложение действительно до 25.11.2026" in body and "Условия оплаты: 100 % предоплата" in body
    assert not any("{{" in text for text in body)
    assert doc.tables[0].rows[1].cells[4].text == prajs.money(189000)

    wb = load_workbook(reg.files("ТКП-2026-0001")[1])
    rows = [[c for c in r if c is not None] for r in wb["ТКП"].iter_rows(values_only=True)]
    assert rows[0] == ["ООО «Ромашка»"] and "№ ТКП-2026-0001 от 26.10.2026" in rows[1][0]
    assert ["Итого по прайсу", 382350] in rows and any(r[:2] == [1, "Осциллограф АКИП-4143/1"] for r in rows)
    assert wb["Источники"].max_row > 1


def test_status_history_summary_and_export(reg):
    a = approve(reg, dt.datetime(2026, 10, 26, 10, 0))
    b = approve(reg, dt.datetime(2026, 10, 26, 11, 0))
    approve(reg, dt.datetime(2026, 10, 26, 12, 0))
    reg.set_status(a["number"], "отправлен клиенту", "m", "Менеджер")
    reg.set_status(a["number"], "выигран", "m", "Менеджер", "счёт оплачен")
    reg.set_status(b["number"], "проигран", "m", "Менеджер", "выбрали конкурента — цена")
    with pytest.raises(ValueError):
        reg.set_status(a["number"], "непонятно", "m", "Менеджер")
    rec = next(r for r in reg.load() if r["number"] == a["number"])
    assert [h["status"] for h in rec["history"]] == ["утверждён", "отправлен клиенту", "выигран"]
    s = reg.summary(reg.load())
    assert (s["count"], s["won"], s["open"], s["win_rate"]) == (3, 1, 1, 50) and s["won_total"] == a["total"]
    ws = load_workbook(io.BytesIO(reg.export(reg.load()))).active
    assert ws.cell(row=1, column=1).value == "Номер" and ws.max_row == 4


def test_without_blank_tkp_is_built_as_before(reg):
    rec = approve(reg, dt.datetime(2026, 10, 26, 10, 0))
    body = [p.text for p in Document(reg.files(rec["number"])[0]).paragraphs if p.text.strip()]
    assert body[0] == "ТЕХНИКО-КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ" and body[1] == f"№ {rec['number']} от 26.10.2026"
    assert "Предложение действительно до 25.11.2026" in body          # срок действия по умолчанию — 30 дней


def test_revision_and_reminders(reg):
    first = approve(reg, dt.datetime(2026, 10, 1, 10, 0))
    t = tkp.parse(LETTER, {"all"}).to_dict()
    rev = reg.approve(t, tkp.draft_text(t, "М"), "m", "Менеджер", "Техпродажи и сервис", LETTER, "Агент",
                      now=dt.datetime(2026, 10, 2, 10, 0), revision_of=first["number"])
    assert rev["number"] == "ТКП-2026-0001/1" and rev["revision_of"] == first["number"]
    old = next(r for r in reg.load() if r["number"] == first["number"])
    assert old["status"] == "заменён" and reg.summary(reg.load())["count"] == 1
    assert approve(reg, dt.datetime(2026, 10, 3, 10, 0))["number"] == "ТКП-2026-0002"   # редакция не тратит номер
    reg.set_status(rev["number"], "отправлен клиенту", "m", "Менеджер")
    found = {(a["record"]["number"], a["kind"]) for a in reg.attention(reg.load(), today=dt.date.today()
                                                                         + dt.timedelta(days=40))}
    assert (rev["number"], "expired") in found and (rev["number"], "silent") in found
    text = reg.reminder_text(next(r for r in reg.load() if r["number"] == rev["number"]), "Менеджер")
    assert rev["number"] in text and "АКИП-4143/1" in text
    reg.note(rev["number"], "напомнили клиенту", "m", "Менеджер")
    assert next(r for r in reg.load() if r["number"] == rev["number"])["history"][-1]["comment"] == "напомнили клиенту"
