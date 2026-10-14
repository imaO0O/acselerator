"""Справочник параметров: проверка данных при загрузке, понимание вопроса, вывод «подходит / не подходит» программой."""
from kbagent import spravochnik as S

ALL = {"all", "production", "metrology", "admin"}
HEAD = "Модель;Вид;Параметр;Значение;Единица;Условие;Документ;Страница;Цитата\n"


def test_import_checks_quality_and_does_not_duplicate(tmp_path):
    f = tmp_path / "s.csv"
    f.write_text(HEAD +
                 "X-100;Мультиметр;Постоянный ток, верхний предел;10;А;;doc1;2;10A\n"
                 "X-100;Мультиметр;Постоянный ток, верхний предел;10;А;;doc1;2;10A\n"        # повтор — не дублируется
                 "X-100;Мультиметр;Масса;;г;;doc1;2;\n"                                    # пусто
                 "X-100;Мультиметр;Масса;160;фунт;;doc1;2;\n"                              # неизвестная единица
                 "X-100;Мультиметр;Масса;160;г;;doc1;две;\n"                               # страница не число
                 "X-100;Мультиметр;Масса;160;г;;nodoc;2;\n"                                # документа нет в базе
                 "X-200;Осциллограф;Полоса пропускания;100;МГц;;doc1;6;\n"
                 "X-200;Осциллограф;Полоса пропускания;0,12;ГГц;;doc1;9;\n",               # 100 МГц ≠ 120 МГц
                 encoding="utf-8")
    values, rep = S.parse(f, known_docs={"doc1"})
    assert rep.rows == 8 and len(values) == 3 and rep.duplicates == 1
    assert len(rep.errors) == 4
    assert [c["model"] for c in rep.contradictions] == ["X-200"]


def test_same_value_in_other_units_is_not_a_contradiction(tmp_path):
    f = tmp_path / "s.csv"
    f.write_text(HEAD + "X-200;Осциллограф;Полоса пропускания;100;МГц;;doc1;6;\n"
                        "X-200;Осциллограф;Полоса пропускания;0,1;ГГц;;doc1;9;\n", encoding="utf-8")
    _, rep = S.parse(f, known_docs={"doc1"})
    assert rep.contradictions == []


def test_questions_are_understood_and_checked_by_program():
    values, rep = S.load()
    assert rep.errors == [] and rep.values >= 100
    q = S.understand("Можно ли мультиметром MAS830L измерить ток 15 А?", values)
    assert q.models == ["MAS830L"] and q.reqs[0].mode == "min" and q.reqs[0].value == 15
    a = S.answer(q, ALL, "")
    assert "не подходит" in a.text and "10 А" in a.text and a.sources[0]["doc_id"] == "multimeter_mas830l"
    # подбор по нескольким параметрам: «не менее» читается раньше «менее»
    q = S.understand("Нужен осциллограф с полосой не менее 150 МГц и 4 каналами", values)
    assert [r.mode for r in q.reqs] == ["min", "min"]
    assert "Подходят: 6 из" in S.answer(q, ALL, "").text
    # обозначение модели не принимается за требование («4115/1В» — не «1 В»)
    q = S.understand("Какая полоса пропускания у АКИП-4115/2В?", values)
    assert q.models == ["АКИП-4115/2В"] and q.reqs == []
    # температура — внутри диапазона
    q = S.understand("Подойдёт ли SBK936B для пайки при 450 °C?", values)
    assert q.reqs[0].mode == "inside" and "подходит" in S.answer(q, ALL, "").text


def test_contradiction_is_shown_and_sent_to_expert(tmp_settings):
    from kbagent.experts import load
    values, _ = S.load()
    q = S.understand("Какая частота дискретизации у АКИП-4115/1В?", values)
    a = S.answer(q, ALL, "Какая частота дискретизации у АКИП-4115/1В?")
    assert "Источники расходятся" in a.text and "1 Гвыб/с" in a.text and "1 Гц/с" in a.text
    assert load()[0]["subject"].startswith("АКИП-4115/1В")
    S.answer(q, ALL, "ещё раз")                                    # открытый вопрос повторно не заводится
    assert len(load()) == 1 and load()[0]["asked"] == 2


def test_document_questions_are_not_taken_by_reference():
    values, _ = S.load()
    for question in ["Как заменить батарею в MAS830L?", "Какой плавкий предохранитель установлен в мультиметре MAS830L?",
                     "При какой температуре окружающего воздуха допускается эксплуатация паяльной станции «Магистр Ц20-ДВ»?",
                     "Как часто нужно включать осциллограф АКИП-4115 при длительном хранении?",
                     "Сколько стоит мультиметр TDM?", "Какое сопротивление должно быть у антистатического браслета?"]:
        assert S.understand(question, values) is None, question


def test_client_letter_becomes_tkp_positions_without_model(tmp_path, monkeypatch):
    from kbagent import experts, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    letter = ("Добрый день! Просим направить коммерческое предложение: 2 осциллографа с полосой не менее 150 МГц "
              "и 4 каналами; мультиметр для измерения постоянного тока до 15 А — 3 шт.; паяльная станция для пайки "
              "при 450 °C; источник питания 30 В 5 А.")
    t = tkp.parse(letter, {"all"})
    assert tkp.is_request(letter, t)
    osc, mm, st, ps = t.positions
    assert (osc.kind, osc.qty, osc.status, osc.model, osc.fit_count) == ("Осциллограф", 2, "offer", "АКИП-4143/1", 6)
    assert (mm.qty, mm.model, mm.status) == (3, "DT9205A", "offer") and all(c.ok for c in mm.checks)
    assert st.status == "offer" and st.reqs == ["температура жала паяльника 450 °C"]
    assert ps.status == "not_in_catalog" and ps.kind == "Источник питания"
    text = tkp.answer_text(t)
    assert "Предлагаем **АКИП-4143/1**" in text and "нет в справочнике" in text
    assert not experts.load()                      # разбор ничего не записывает сам


def test_tkp_is_not_triggered_by_ordinary_questions(tmp_path, monkeypatch):
    from kbagent import experts, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    for q in ["Клиенту нужен осциллограф с полосой не менее 150 МГц и 4 каналами — что предложить?",
              "Сколько осциллографов и мультиметров на поверке?", "Где лежат паяльные станции и мультиметры?",
              "Чем отличаются АКИП-4122/1(V) и АКИП-4122/10(V)?", "Что делать при искрении оборудования?"]:
        assert not tkp.is_request(q, tkp.parse(q, {"all"})), q
    t = tkp.parse("Подготовь ТКП: нужен мультиметр", {"all"})
    assert tkp.is_request("Подготовь ТКП: нужен мультиметр", t)
    assert t.positions[0].status == "no_reqs" and "уточните" in t.positions[0].ask_client


def test_tkp_word_document_has_positions_table(tmp_path, monkeypatch):
    import io

    from docx import Document

    from kbagent import experts, prais, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    monkeypatch.setattr(prais, "PATH", tmp_path / "net_praisa.csv")       # без прайса цены вписывает менеджер
    t = tkp.parse("ТКП: 2 осциллографа от 150 МГц на 4 канала, мультиметр на ток до 15 А", {"all"}).to_dict()
    text = tkp.draft_text(t, "Сотрудник")
    assert "АКИП-4143/1 — 2 шт." in text and "[заполняет менеджер]" in text
    doc = Document(io.BytesIO(tkp.tkp_docx(text, t, "Сотрудник", "13.10.2026 12:00", "Агент")))
    cells = [c.text for c in doc.tables[0].rows[1].cells]
    assert cells[1] == "Осциллограф АКИП-4143/1" and cells[2] == "2" and "соответствует" in cells[3]


def test_expert_decision_is_used_in_later_answers(tmp_path, monkeypatch):
    from kbagent import experts
    from kbagent import spravochnik as spr
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    q = spr.understand("Какая частота дискретизации у АКИП-4115/1В?")
    first = spr.answer(q, {"all"}, "вопрос")
    assert "Источники расходятся" in first.text and len(experts.load()) == 1
    it = experts.load()[0]
    right = next(d for d in it["details"] if d["value"].startswith("1 Гвыб"))
    experts.resolve(it["id"], "Эксперт", "на с. 6 опечатка", right)
    again = spr.answer(q, {"all"}, "вопрос")
    assert "подтверждено экспертом" in again.text and "Источники расходятся" not in again.text
    assert len(experts.load()) == 1                # вопрос эксперту повторно не заводится
    pick = spr.answer(spr.understand("Осциллограф с частотой дискретизации от 800 Мвыб/с"), {"all"}, "вопрос")
    assert "АКИП-4115/1В" in pick.text.split("Не подходят")[0]


def test_model_comparison_lists_differences():
    from kbagent import spravochnik as spr
    a = spr.answer(spr.understand("Чем отличаются АКИП-4122/1(V) и АКИП-4122/10(V)?"), {"all"}, "вопрос")
    assert "Различаются: частота дискретизации на канал" in a.text


def test_letter_text_is_read_from_word_text_and_email():
    import io

    from docx import Document

    from kbagent import tkp
    doc = Document()
    doc.add_paragraph("Просим КП: 2 осциллографа от 150 МГц на 4 канала")
    buf = io.BytesIO()
    doc.save(buf)
    assert "2 осциллографа" in tkp.letter_text("Письмо.DOCX", buf.getvalue())
    assert "мультиметр" in tkp.letter_text("p.txt", "Нужен мультиметр до 15 А".encode("cp1251"))
    eml = ("Subject: =?utf-8?b?" + __import__("base64").b64encode("Запрос КП".encode()).decode() + "?=\n"
           "Content-Type: text/plain; charset=utf-8\nContent-Transfer-Encoding: 8bit\n\n"
           "Нужна паяльная станция до 450 °C").encode("utf-8")
    text = tkp.letter_text("m.eml", eml)
    assert text.startswith("Запрос КП") and "паяльная станция" in text
    t = tkp.parse(text, {"all"})
    assert tkp.is_request(text, t) and t.positions[0].kind == "Паяльная станция"


def test_letter_only_with_unknown_products_is_an_honest_refusal(tmp_path, monkeypatch):
    from kbagent import experts, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    text = "Подготовьте ТКП на лабораторный источник питания 30 В 5 А и генератор сигналов"
    t = tkp.parse(text, {"all"})
    assert tkp.is_request(text, t) and not tkp.has_offer(t)
    assert "нет в справочнике" in tkp.answer_text(t)


def _values(tmp_path):
    f = tmp_path / "s.csv"
    f.write_text(HEAD +
                 "X-100;Мультиметр;Постоянный ток, верхний предел;10;А;;doc1;2;10A\n"
                 "X-100;Мультиметр;Постоянное напряжение, верхний предел;600;В;;doc1;2;600V\n"
                 "X-200;Осциллограф;Число каналов;2;;;doc1;6;2 канала\n"
                 "X-200;Осциллограф;Полоса пропускания;100;МГц;;doc1;6;100 МГц\n", encoding="utf-8")
    return S.parse(f, known_docs={"doc1"})[0]


def test_how_to_questions_go_to_documents_not_to_reference(tmp_path):
    values = _values(tmp_path)
    # вопрос о значении или требовании — справочник
    for q in ["Можно ли мультиметром X-100 измерить ток 15 А?", "Какая полоса пропускания у осциллографа X-200?",
              "Нужен осциллограф от 50 МГц и 2 каналами"]:
        assert S.understand(q, values) is not None, q
    # порядок работы, неисправность, меры безопасности, питание — ответ в руководстве, а не таблица параметров
    for q in ["Как измерить постоянное напряжение мультиметром X-100?",
              "Как заменить предохранитель на 10 А в мультиметре X-100?",
              "Что делать, если осциллограф X-200 не показывает сигнал на втором канале?",
              "Как подключить пробник к каналу осциллографа?",
              "Почему мультиметр показывает 0 при измерении тока?",
              "Осциллограф X-200 — какие меры безопасности при измерении высокого напряжения?",
              "Какое напряжение питания у мультиметра X-100?"]:
        assert S.understand(q, values) is None, q


def test_expert_queue_masks_personal_data(tmp_settings):
    from kbagent import experts
    it = experts.add("передал сотрудник", "Иванов Иван Иванович спрашивает про допуск", "—",
                     "Иванов Иван Иванович, тел. +7 910 123-45-67: где допуск?", comment="звонить Ивановой Анне Петровне")
    stored = experts.load()[0]
    for field in ("subject", "question", "comment"):
        assert "Иван" not in stored[field] and "123-45-67" not in stored[field], (field, stored[field])
    assert it["id"] == stored["id"]


def test_lowercase_units_in_client_letters_are_not_lost(tmp_path):
    values = _values(tmp_path)
    q = S.understand("Нужен осциллограф 150 мгц 2 канала", values)
    assert {(r.param, round(r.value)) for r in q.reqs} == {("Полоса пропускания", 150_000_000), ("Число каналов", 2)}
    assert [r.param for r in S.understand("Нужен мультиметр X-100 для тока 20 а?", values).reqs] == \
        ["Постоянный ток, верхний предел"]
    # «а» и «в» внутри фразы — союз и предлог, а не единицы
    q = S.understand("Нужен мультиметр X-100 на 20 а также для сопротивления", values)
    assert q is None or not q.reqs


def test_letter_positions_keep_requirements_next_to_service_words(tmp_path, monkeypatch):
    from kbagent import experts, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    # в письме «подключение», «обслуживание», «калибровка», «питание» — не вопрос о порядке работы
    letter = ("Просим КП: 2 осциллографа от 150 МГц на 4 канала с подключением к ПК по USB и гарантийным "
              "обслуживанием; мультиметр с калибровкой для измерения тока до 15 А с питанием от батареи — 3 шт.")
    osc, mm = tkp.parse(letter, {"all"}).positions
    assert (osc.status, osc.model, len(osc.reqs)) == ("offer", "АКИП-4143/1", 2)
    assert (mm.status, mm.model, mm.qty) == ("offer", "DT9205A", 3)
    values = _values(tmp_path)
    assert S.understand("Мультиметр X-100 с питанием от батареи — какое постоянное напряжение измеряет?", values)


def _positions(text, tmp_path, monkeypatch):
    from kbagent import experts, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    t = tkp.parse(text, {"all"})
    assert tkp.is_request(text, t), text
    return [(p.kind, p.qty, p.status, p.model) for p in t.positions], t


def test_letter_range_and_latin_units(tmp_path, monkeypatch):
    # «от 100 до 450 °C» — обе границы: Lukey и SBK (от 200 °C) не подходят
    pos, t = _positions("Нужна паяльная станция с температурой жала от 100 до 450 °C, просим КП", tmp_path, monkeypatch)
    assert pos == [("Паяльная станция", 1, "offer", "CXG DES H95")] and t.positions[0].fit_count == 1
    assert t.positions[0].reqs == ["температура жала паяльника от 100 °C", "температура жала паяльника до 450 °C"]
    # единицы латиницей: GHz, A, W, «≥»
    pos, t = _positions("Просим КП: осциллограф 4 канала, bandwidth >= 1 GHz; мультиметр DC 15 A — 2 шт.; "
                        "паяльная станция с паяльником 90 W", tmp_path, monkeypatch)
    assert pos == [("Осциллограф", 1, "offer", "АКИП-4143/1"), ("Мультиметр", 2, "offer", "DT9205A"),
                   ("Паяльная станция", 1, "offer", "CXG DES H95")]
    assert "мощность паяльника не менее 90 Вт" in t.positions[2].reqs
    # обозначение модели с латинской буквой — не единица: «DT9205A» — не «9205 А»
    assert S._requirements("мультиметр DT9205A на ток 20 A", "Мультиметр")[0].shown == "20 А"


def test_letter_with_typos_and_request_words(tmp_path, monkeypatch):
    pos, _ = _positions("Здраствуйте! Просим комерческое предложение: осцилограф с полосой от 150 МГц на 4 канала, "
                        "мультимер для измерения тока до 15 А, паяльная станцыя до 450 С.", tmp_path, monkeypatch)
    assert [p[0] for p in pos] == ["Осциллограф", "Мультиметр", "Паяльная станция"] and pos[2][3] == "CXG DES H95"
    # «просим предложение», «цену и срок поставки» — запрос клиента, а вопрос о сроке без просьбы — нет
    _positions("Просим предложение на мультиметр: переменное напряжение до 750 В", tmp_path, monkeypatch)
    pos, _ = _positions("Прошу направить цену и срок поставки: осциллограф АКИП-4122/12(V) — 1 шт., "
                        "мультиметр DT9208A — 2 шт.", tmp_path, monkeypatch)
    assert pos == [("Осциллограф", 1, "named", "АКИП-4122/12(V)"), ("Мультиметр", 2, "named", "DT9208A")]
    from kbagent import tkp
    q = "Какой срок поставки у осциллографа АКИП-4115/2В?"
    assert not tkp.is_request(q, tkp.parse(q, {"all"}))


def test_letter_quantities_tables_and_hot_air(tmp_path, monkeypatch):
    pos, t = _positions("Коммерческое предложение на: осциллограф с полосой от 300 МГц в количестве 2 (двух) штук; "
                        "мультиметр с сопротивлением до 200 МОм — 10 ед.; паяльная станция с феном — 3 комплекта.",
                        tmp_path, monkeypatch)
    assert [p[1] for p in pos] == [2, 10, 3] and t.positions[2].reqs == ["наличие фена"]
    assert t.positions[2].model in ("CXG DES H95", "Lukey 852D+")          # только станции с феном
    # таблица из Word: «|» после разбора файла, табуляция при вставке; «с феном» — та же позиция, а не новая
    for sep in (" | ", "\t"):
        rows = [["№", "Наименование", "Характеристики", "Кол-во"],
                ["1", "Осциллограф", "полоса не менее 100 МГц, 2 канала", "2"],
                ["2", "Мультиметр", "постоянное напряжение до 1000 В", "5"],
                ["3", "Паяльная станция", "температура до 480 °C, с феном", "1"]]
        text = "Просим ТКП по позициям:\n" + "\n".join(sep.join(r) for r in rows)
        pos, t = _positions(text, tmp_path, monkeypatch)
        assert [(p[0], p[1]) for p in pos] == [("Осциллограф", 2), ("Мультиметр", 5), ("Паяльная станция", 1)], sep
        assert t.positions[2].reqs == ["температура жала паяльника до 480 °C", "наличие фена"]
    # «температура фена» после станции — параметр той же станции
    pos, t = _positions("Просим КП: паяльная станция с паяльником и феном, температура фена до 450 °C — 2 шт.",
                        tmp_path, monkeypatch)
    assert pos[0][:2] == ("Паяльная станция", 2) and len(pos) == 1
    assert t.positions[0].reqs == ["температура фена до 450 °C"]


def test_letter_analog_and_subject_line(tmp_path, monkeypatch):
    from kbagent import tkp
    # тема письма .eml называет ту же продукцию — не отдельные позиции; «или аналог» — подбор и среди других моделей
    eml = ("Subject: =?utf-8?b?" + __import__("base64").b64encode("Запрос КП: осциллограф и мультиметры".encode())
           .decode() + "?=\nContent-Type: text/plain; charset=utf-8\nContent-Transfer-Encoding: 8bit\n\n"
           "Просим КП: осциллограф АКИП-4122/12(V) или аналог с полосой не менее 200 МГц — 2 шт.; мультиметр "
           "MAS830L или аналог для постоянного тока не менее 15 А — 3 шт.").encode("utf-8")
    text = tkp.letter_text("m.eml", eml)
    pos, t = _positions(text, tmp_path, monkeypatch)
    assert pos == [("Осциллограф", 2, "offer", "АКИП-4122/12(V)"), ("Мультиметр", 3, "offer", "DT9205A")]
    assert t.positions[0].alternatives                       # названная подходит — она первая, аналоги — ниже
    answer = tkp.answer_text(t)
    assert "MAS830L не подходит: постоянный ток, верхний предел 10 А при требовании не менее 15 А" in answer


def _price_list(tmp_path, monkeypatch, text, demo=True):
    import json

    from kbagent import prais
    (tmp_path / "prais.csv").write_text(text, encoding="utf-8")
    (tmp_path / "prais.json").write_text(json.dumps({"source": "p.csv", "date": "13.10.2026", "demo": demo}),
                                         encoding="utf-8")
    monkeypatch.setattr(prais, "DIR", tmp_path)
    monkeypatch.setattr(prais, "PATH", tmp_path / "prais.csv")
    monkeypatch.setattr(prais, "META", tmp_path / "prais.json")
    prais._CACHE.clear()
    return prais


def test_price_list_is_checked_like_the_reference(tmp_path, monkeypatch):
    prais = _price_list(tmp_path, monkeypatch, "Модель;Цена;Срок поставки\n"
                                               "DT9205A;1 590 ₽;в наличии\n"
                                               "MAS830L;0;в наличии\n"             # не положительная цена
                                               "Мультиметр-X;990;\n"               # нет в справочнике
                                               "DT9205A;1590;в наличии\n"          # повтор с той же ценой
                                               "DT9205A;1690;2 недели\n"           # повтор с другой ценой
                                               "DT9208A;;\n")                      # нет цены
    prices, rep = prais.load()
    assert set(prices) == {"DT9205A"} and prices["DT9205A"].price == 1590 and rep.duplicates == 1
    assert [e["line"] for e in rep.errors] == [3, 4, 6, 7]
    assert prais.label(rep) == "по прайсу от 13.10.2026 (ДЕМО — цены условные)"
    assert prais.money(1590) == "1 590 ₽" and prais.money(1234.5) == "1 234,50 ₽"
    # новый прайс от администратора: ошибки во всех строках — прежний остаётся; иначе дата загрузки и без «ДЕМО»
    import pytest
    with pytest.raises(ValueError, match="ни одна строка"):
        prais.replace_file("p.csv", "Модель;Цена\nНет такой;100\n".encode("utf-8"))
    assert set(prais.load()[0]) == {"DT9205A"}
    rep = prais.replace_file("p.csv", "Модель;Цена;Срок поставки\nMAS830L;1290;в наличии\n".encode("utf-8"))
    assert rep.prices == 1 and not rep.demo and "ДЕМО" not in prais.label(rep)


def test_tkp_takes_prices_from_price_list(tmp_path, monkeypatch):
    import io

    from docx import Document

    from kbagent import experts, tkp
    monkeypatch.setattr(experts, "PATH", tmp_path / "experts.json")
    _price_list(tmp_path, monkeypatch, "Модель;Цена;Срок поставки\nАКИП-4143/1;890000;6–8 недель\n"
                                       "DT9205A;1590;в наличии\n")
    t = tkp.parse("ТКП: 2 осциллографа от 150 МГц на 4 канала, мультиметр на ток до 15 А — 3 шт., паяльная станция "
                  "до 450 °C", {"all"})
    osc, mm, st = t.positions
    assert (osc.price, osc.lead, mm.price, st.price) == (890000, "6–8 недель", 1590, None)
    answer = tkp.answer_text(t)
    assert "Цена по прайсу: 890 000 ₽ × 2 шт. = 1 780 000 ₽, срок поставки — 6–8 недель" in answer
    assert "по прайсу от 13.10.2026 (ДЕМО" in answer and "без позиций без цены (1)" in answer
    d = t.to_dict()
    text = tkp.draft_text(d, "Сотрудник")
    assert "Цена: 1 590 ₽ за шт., сумма 4 770 ₽, срок поставки — в наличии." in text
    assert "Цена: [заполняет менеджер]" in text                       # паяльной станции в прайсе нет
    table = Document(io.BytesIO(tkp.tkp_docx(text, d, "Сотрудник", "14.10.2026 12:00", "Агент"))).tables[0]
    assert [c.text for c in table.rows[1].cells][4:] == ["890 000", "1 780 000"]
    assert table.rows[-1].cells[0].text.startswith("Итого") and table.rows[-1].cells[5].text == "1 784 770"
