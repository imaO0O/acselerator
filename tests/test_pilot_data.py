"""Шаблон данных партнёра: собирается, заполняется и загружается с проверкой каждой строки."""
import importlib.util
from pathlib import Path

from openpyxl import load_workbook

from kbagent import pilot_data

ROOT = Path(__file__).resolve().parents[1]


def _builder():
    spec = importlib.util.spec_from_file_location("make_partner_package", ROOT / "scripts" / "make_partner_package.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _fill(path: Path, sheet: str, rows: list[list]) -> None:
    wb = load_workbook(path)
    ws = wb[sheet]
    for i, row in enumerate(rows, start=3):          # строка 2 — пример
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    wb.save(path)


def test_template_is_filled_and_loaded_with_row_errors(tmp_path):
    book = tmp_path / "shablon.xlsx"
    _builder().build_xlsx(book)
    docs = tmp_path / "partner"
    docs.mkdir()
    for name in ("Regl_2025.pdf", "Regl_2024.pdf", "Manual.pdf", "Prais.xlsx"):
        (docs / name).write_bytes(b"%PDF-1.4")
    _fill(book, "Реестр документов", [
        [1, "Regl_2025.pdf", "СТП 12-2025 «Порядок ТКП»", "регламент", "2025", "01.09.2025", "", "Нач. отдела",
         "все сотрудники", 1, "действует", "Regl_2024.pdf"],
        [2, "Regl_2024.pdf", "СТП 12-2024 «Порядок ТКП»", "регламент", "2024", "01.02.2024", "", "", "все сотрудники"],
        [3, "Manual.pdf", "Руководство по эксплуатации", "", "", "32.13.2025", "", "", "все сотрудники"],
        [4, "Net_takogo.pdf", "Нет файла", "", "", "", "", "", "все сотрудники"],
        [5, "Manual.pdf", "Руководство", "", "", "", "", "", "отдел продаж"],
        [6, "Prais.xlsx", "Прайс-лист", "", "", "", "", "", "все сотрудники"],
    ])
    _fill(book, "Вопросы", [
        [1, "Срок ответа на запрос ТКП?", "Техпродажи и сервис", "в документе", "2 рабочих дня", "2 рабочих дн",
         "", "Regl_2025.pdf", "п. 4.3"],
        [2, "Сколько стоит доставка в Мурманск?", "", "нет в документах"],
        [3, "Срок поверки?", "", "в документе", "", "", "", "Regl_2025.pdf"],              # нет ключевых фактов
        [4, "Кто владелец?", "Директор", "в документе", "", "Иванов", "", "Regl_2025.pdf"],  # нет такой роли
    ])
    res = pilot_data.load(book, docs)
    by_id = {d["id"]: d for d in res.documents}
    assert set(by_id) == {"Regl_2025", "Regl_2024", "Prais"}         # пример не загружен, ошибочные строки — тоже
    assert len(res.notes) == 1 and "Prais.xlsx" in res.notes[0]       # Excel — через администрирование
    assert by_id["Regl_2025"]["effective_from"] == "2025-09-01" and by_id["Regl_2025"]["priority"] == 1
    assert by_id["Regl_2024"]["status"] == "отменён" and by_id["Regl_2024"]["replaced_by"] == "Regl_2025"
    q = res.questions
    assert [x["type"] for x in q] == ["doc", "none"]
    assert q[0]["source"] == {"doc": "Regl_2025", "clause": "4.3"} and q[0]["must_include"] == ["2 рабочих дн"]
    assert q[0]["role"] == "Техпродажи и сервис" and q[0]["id"] == "P01"
    rows = {(e["sheet"], e["row"]) for e in res.errors}
    assert rows == {("Реестр документов", 5), ("Реестр документов", 6), ("Реестр документов", 7),
                    ("Вопросы", 5), ("Вопросы", 6)}


def test_reference_and_price_sheets_are_checked_like_admin_upload(tmp_path):
    book = tmp_path / "shablon.xlsx"
    _builder().build_xlsx(book)
    _fill(book, "Справочник параметров", [
        [1, "АКИП-4115/2В", "Осциллограф", "Полоса пропускания", "200", "МГц", "", "osc_akip_4115v.pdf", 6, "200МГц"],
        [2, "АКИП-4115/2В", "Осциллограф", "Число каналов", "2", "", "", "osc_akip_4115v.pdf", "семь", ""],
    ])
    _fill(book, "Прайс", [[1, "АКИП-4115/2В", "46 500", "в наличии", "01.10.2026"], [2, "Чужая модель", "100"]])
    res = pilot_data.load(book)
    assert res.spravochnik_csv.splitlines()[1].startswith("АКИП-4115/2В;Осциллограф;Полоса пропускания;200;МГц;;osc_akip_4115v;6")
    assert res.spravochnik_models == {"АКИП-4115/2В"}
    errors = [(e["sheet"], e["row"]) for e in res.errors]
    assert ("Справочник параметров", 4) in errors                      # страница «семь» — строка Excel 4
    assert any(e["sheet"] == "Прайс" and "Чужая модель" in e["error"] for e in res.errors)
    assert "АКИП-4115/2В;46 500;в наличии" in res.prajs_csv


def test_partner_package_files_are_built(tmp_path, monkeypatch):
    mod = _builder()
    monkeypatch.setattr(mod, "OUT", tmp_path)
    mod.main()
    wb = load_workbook(tmp_path / "Shablon_dannyh_dlya_pilota.xlsx")
    assert wb.sheetnames == ["Как заполнить", *pilot_data.SHEETS]
    assert [c.value for c in wb["Вопросы"][1]] == [c.name for c in pilot_data.SHEETS["Вопросы"]]
    assert (tmp_path / "Pamyatka_partneru.docx").stat().st_size > 10_000
