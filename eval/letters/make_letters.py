"""Письма клиентов файлами для набора eval/tkp.yaml: .eml, PDF и Word с таблицей позиций.

Письма вымышленные; имена и телефоны в подписях — для проверки маскирования.
    python eval/letters/make_letters.py
"""
import io
from email.message import EmailMessage
from pathlib import Path

HERE = Path(__file__).resolve().parent

EML_BODY = """Добрый день!

Просим направить коммерческое предложение:
1) осциллограф АКИП-4122/12(V) или аналог с полосой пропускания не менее 200 МГц — 2 шт.;
2) мультиметр MAS830L или аналог для измерения постоянного тока не менее 15 А — 3 шт.

С уважением,
Сидоров Пётр Ильич, инженер-метролог
тел. +7 900 111-22-33
"""

PDF_TEXT = """ООО «Приборсервис»
Исх. № 118 от 13.10.2026

Просим предоставить предложение на поставку:
- паяльная станция с феном — 2 комплекта;
- мультиметр для измерения сопротивления ≥ 100 МОм — 2 (двух) штук.

Главный инженер Кузнецов В. А., тел. 8 (4912) 20-30-40"""

DOCX_ROWS = [("№", "Наименование", "Требования", "Кол-во"),
             ("1", "Осциллограф цифровой", "полоса пропускания не менее 100 МГц, частота дискретизации не ниже 1 Гвыб/с", "2"),
             ("2", "Мультиметр цифровой", "переменное напряжение до 750 В, масса не более 200 г", "4"),
             ("3", "Паяльная станция", "температура жала от 100 до 450 °C", "1")]


def eml() -> bytes:
    msg = EmailMessage()
    msg["Subject"] = "Запрос КП: осциллограф и мультиметры"
    msg["From"] = "p.sidorov@example.ru"
    msg["To"] = "sales@example.ru"
    msg.set_content(EML_BODY)
    return bytes(msg)


def pdf() -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    html = "".join(f"<p>{line or '&nbsp;'}</p>" for line in PDF_TEXT.splitlines())
    page.insert_htmlbox(pymupdf.Rect(60, 60, 540, 780), html, css="p {font-size: 12pt; margin: 0 0 4pt 0;}")
    return doc.tobytes()


def docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_paragraph("Добрый день! Просим подготовить ТКП по позициям:")
    table = d.add_table(rows=0, cols=4)
    for row in DOCX_ROWS:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    d.add_paragraph("Начальник отдела снабжения Орлова Е. Н., тел. +7 910 765-43-21")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    for name, make in [("K15_zapros.eml", eml), ("K16_zapros.pdf", pdf), ("K17_zapros.docx", docx)]:
        (HERE / name).write_bytes(make())
        print(HERE / name)
