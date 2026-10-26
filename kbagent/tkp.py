"""Запрос клиента → черновик технико-коммерческого предложения (ТКП).

Сотрудник вставляет письмо клиента целиком («нужны 2 осциллографа от 150 МГц на 4 канала, мультиметр на ток
до 15 А — 3 шт., паяльная станция до 450 °C»). Агент без языковой модели:
  1) делит письмо на позиции — по видам продукции и моделям из справочника параметров;
  2) для каждой позиции выделяет требования и количество;
  3) подбирает модели по справочнику: подходит / не подходит / нет данных или источники расходятся;
  4) предлагает минимально достаточную модель — с наименьшим запасом по требованиям (окончательный выбор —
     за менеджером), показывает альтернативы;
  5) составляет вопросы клиенту по позициям без требований и передаёт эксперту противоречия источников.
Цены и сроки агент не придумывает: подставляет их из прайса (kbagent/prajs.py) с пометкой «по прайсу от …» и
считает сумму «цена × количество»; чего нет в прайсе, вписывает менеджер. Утверждает ТКП сотрудник, автоотправки нет.
"""
import io
import math
import re
from dataclasses import asdict, dataclass, field

from . import prajs
from . import spravochnik as spr

TKP_WORDS_RE = re.compile(r"\bТ?КП\b|коммерческ\w*\s+предложени\w*|запрос\w*\s+клиент\w*|письм\w*\s+клиент\w*|"
                          r"заявк\w*\s+клиент\w*|подготов\w*\s+предложени\w*|подбер\w*\s+(?:и\s+)?предлож\w*|"
                          r"(?:просим|прошу|просьба)\s+(?:вас\s+)?(?:предложить|подобрать|выставить|направить\s+предложени)",
                          re.I)
# продукция, которой нет в справочнике: позиция есть, но подбирать не по чему — к техническому отделу
OTHER_RE = re.compile(r"источник\w*\s+питани\w*|блок\w*\s+питани\w*|генератор\w*|частотомер\w*|анализатор\w*|"
                      r"токов\w*\s+клещ\w*|клещ\w*|мегаомметр\w*|пирометр\w*|термометр\w*|тепловизор\w*|"
                      r"вольтметр\w*|амперметр\w*|калибратор\w*|электронн\w*\s+нагрузк\w*|паяльн\w*\s+ванн\w*|"
                      r"микроскоп\w*|дымоуловител\w*|вытяжк\w*|щуп\w*|пробник\w*", re.I)
SEP_RE = re.compile(r"[.;:!?\n]+\s*|,\s+|\s(?:и|а\s+также|также|плюс|ещё|еще|кроме\s+того)\s", re.I)
WORD_NUM = {"один": 1, "одна": 1, "одну": 1, "два": 2, "две": 2, "три": 3, "четыре": 4, "пять": 5, "шесть": 6,
            "семь": 7, "восемь": 8, "девять": 9, "десять": 10}
QTY_BEFORE_RE = re.compile(rf"(?:^|\s)(\d{{1,3}}|{'|'.join(WORD_NUM)})\s+(?:[а-яё]+\s+)?$", re.I)
QTY_RE = re.compile(r"(\d{1,3})\s*(?:шт\b\.?|штук\w*|ед\.|единиц\w*|комплект\w*)", re.I)
# строка таблицы из Word или Excel: «Мультиметр | переменное напряжение до 750 В | 5» — количество в последней ячейке
TABLE_QTY_RE = re.compile(r"[|\t]\s*(\d{1,3})\s*(?:шт\.?)?\s*\|?\s*$")
PRICE_WORDS_RE = re.compile(r"\b(?:с\s+)?(?:цен\w*|стоимост\w*|прайс\w*)\b", re.I)


@dataclass
class Check:
    param: str
    need: str            # «не менее 150 МГц» или пусто (характеристика без требования)
    value: str           # «2 ГГц»
    ref: int | None      # номер источника
    ok: bool | None      # соответствует / нет / не проверить
    note: str = ""       # «подтверждено экспертом 13.10.2026»


@dataclass
class Position:
    n: int
    text: str
    kind: str
    qty: int
    reqs: list = field(default_factory=list)          # «полоса пропускания не менее 150 МГц»
    status: str = "offer"   # offer | named | no_fit | expert | no_reqs | not_in_catalog
    model: str | None = None
    fit_count: int = 0
    total: int = 0
    alternatives: list = field(default_factory=list)
    checks: list = field(default_factory=list)
    note: str = ""
    ask_client: str = ""
    price: float | None = None     # ₽ за единицу — из прайса, только у предложенной или названной модели
    lead: str = ""                 # срок поставки из прайса
    price_since: str = ""          # дата прайса, ГГГГ-ММ-ДД
    price_note: str = ""           # «ДЕМО», «прайс старше 90 дней — проверьте цену»
    alt_prices: dict = field(default_factory=dict)
    fit_models: list = field(default_factory=list)   # все подходящие приборы: минимально достаточный — первым
    minimal: str = ""              # минимально достаточный прибор (предлагается, пока менеджер не выбрал другой)
    cheapest: str = ""             # дешевле всего из подходящих по прайсу — если это не предложенный прибор
    cheapest_price: float | None = None


@dataclass
class Tkp:
    positions: list
    sources: list
    expert_ids: list = field(default_factory=list)
    pending: list = field(default_factory=list)     # противоречия для эксперта — заводятся notify_experts
    price_list: bool = False                        # прайс загружен: «цены нет» — значит, нет в прайсе
    letter: str = ""                                # текст письма — для повторного разбора с выбором менеджера

    def to_dict(self) -> dict:
        d = asdict(self)
        d.pop("pending")
        return d


def notify_experts(t: Tkp, question: str, user: str | None = None, role: str | None = None) -> None:
    """Противоречия источников по позициям — в очередь эксперта (только для настоящего запроса клиента)."""
    from .experts import add
    for subj, details in t.pending:
        t.expert_ids.append(add("противоречие в справочнике", subj, details, question, user, role)["id"])
    t.expert_ids = list(dict.fromkeys(t.expert_ids))
    t.pending = []


def _anchors(text: str, values: list) -> list[tuple[int, int, str]]:
    """Упоминания продукции в письме: (начало, конец, вид); вид «» — продукции нет в справочнике."""
    found = []
    for kind, rx in spr.KINDS.items():
        found += [(m.start(), m.end(), kind) for m in re.finditer(rx, text, re.I)]
    found += [(m.start(), m.end(), "") for m in OTHER_RE.finditer(text)]
    for m in spr.MODEL_TOKEN_RE.finditer(text):
        models = spr._find_models(m.group(0), values)
        if models:
            kind = next((v.kind for v in values if v.model == models[0]), "")
            found.append((m.start(), m.end(), kind))
    found.sort()
    out = []
    for a in found:
        if out and a[0] < out[-1][1]:
            continue                                    # пересекается с предыдущим
        # «осциллограф АКИП-4143/1», «паяльная станция с феном» — одна позиция, если между ними нет разделителя
        if out and not SEP_RE.search(text[out[-1][1]:a[0]]) and (out[-1][2] == a[2] or not out[-1][2] or not a[2]):
            out[-1] = (out[-1][0], a[1], out[-1][2] or a[2])
            continue
        out.append(a)
    return out


def _segments(text: str, anchors: list) -> list[tuple[str, str, int]]:
    """Фрагмент письма на каждую позицию: от последнего разделителя перед упоминанием до следующей позиции."""
    starts = []
    prev_end = 0
    for start, end, _ in anchors:
        seps = list(SEP_RE.finditer(text, prev_end, start))
        starts.append(seps[-1].end() if seps else (prev_end if starts else 0))
        prev_end = end
    out = []
    for i, (start, end, kind) in enumerate(anchors):
        seg = text[starts[i]:starts[i + 1] if i + 1 < len(starts) else len(text)]
        qty = 1
        before = QTY_BEFORE_RE.search(text[starts[i]:start])
        after = QTY_RE.search(seg)
        cell = TABLE_QTY_RE.search(seg.split("\n")[0].rstrip())
        if after:
            qty = int(after.group(1))
        elif cell:
            qty = int(cell.group(1))
        elif before:
            word = before.group(1).lower()
            qty = int(word) if word.isdigit() else WORD_NUM[word]
        out.append((seg.strip(" ,;.—-\n"), kind, qty))
    return out


def _excess(v, r) -> float:
    """Запас значения над требованием (0 — точно по требованию); у диапазонов и «внутри» — 0."""
    lo, hi = v.base(v.lo), v.base(v.hi)
    try:
        if r.mode == "min" and r.value > 0:
            return math.log(max(hi, 1e-30) / r.value)
        if r.mode == "max" and lo > 0:
            return math.log(r.value / lo)
    except ValueError:
        pass
    return 0.0


def parse(text: str, allowed: set[str], chosen: dict | None = None) -> Tkp | None:
    """Разбор письма клиента. None — в тексте нет продукции из справочника (это не запрос на подбор).
    chosen — выбор менеджера {номер позиции: прибор}: прибор из подходящих вместо минимально достаточного.
    Ничего не записывает: противоречия для эксперта копятся в pending."""
    letter = text
    from .experts import decisions
    from .ingest import access_of, load_catalog, registry_suffix
    values, rep = spr.load()
    prices, _ = prajs.load()
    catalog = load_catalog(spr.settings.raw_dir)
    values = [v for v in values if access_of(catalog.get(v.doc, {})) in allowed]
    text = spr.normalize(text)                  # опечатки и единицы латиницей — к написанию справочника
    anchors = _anchors(text, values)
    if not anchors:
        return None
    contradicted = {(c["model"], c["param"]) for c in rep.contradictions}
    decided = decisions()
    by: dict = {}
    for v in values:
        by.setdefault((v.model, v.param), []).append(v)
    sources: list[dict] = []
    index: dict = {}

    def ref(v) -> int:
        k = (v.doc, v.page, v.quote)
        if k not in index:
            index[k] = len(sources) + 1
            info = catalog.get(v.doc, {})
            registry = registry_suffix(v.doc)
            sources.append({"n": index[k], "label": f"{info.get('title', v.doc)}, с. {v.page}"
                            + (f" · {registry}" if registry else ""), "text": v.quote or v.text, "doc_id": v.doc,
                            "doc_kind": info.get("kind", ""), "clause": "", "page": v.page, "page_end": v.page,
                            "cited": True, "score": 1.0})
        return index[k]

    positions, pending = [], []
    for n, (seg, kind, qty) in enumerate(_segments(text, anchors), 1):
        pos = Position(n, seg, kind or _other_name(seg), qty)
        positions.append(pos)
        if not kind:
            pos.status = "not_in_catalog"
            pos.note = "этой продукции нет в справочнике параметров — подбор вручную, вопрос техническому отделу"
            continue
        q = spr.understand(PRICE_WORDS_RE.sub(" ", seg), values)
        named = spr._find_models(seg, values)
        reqs = q.reqs if q else []
        pos.reqs = _req_texts(reqs)
        candidates = named or sorted({v.model for v in values if v.kind == kind})
        pos.total = len(candidates)
        if not reqs and not named:
            pos.status = "no_reqs"
            pos.note = "в запросе нет требований к параметрам"
            pos.ask_client = (f"{kind}: уточните, пожалуйста, требуемые характеристики "
                              f"({_typical(kind)}) — подберём модель")
            continue
        if not reqs:                                     # модель названа, требований нет — её характеристики
            pos.status, pos.model = "named", named[0]
            params = list(dict.fromkeys(v.param for v in values if v.model == named[0]))[:4]
            pos.checks = [_check(named[0], p, None, by, contradicted, decided, ref) for p in params]
            pos.checks = [c for c in pos.checks if c.value]
            continue
        state, checks = {}, {}
        for m in candidates:
            cs = [_check(m, r.param, r, by, contradicted, decided, ref, lazy=True) for r in reqs]
            checks[m] = cs
            state[m] = "fit" if all(c.ok for c in cs) else "unfit" if any(c.ok is False for c in cs) else "unknown"
        fit = [m for m in candidates if state[m] == "fit"]
        pos.fit_count = len(fit)
        if fit:
            minimal = min(fit, key=lambda m: (sum(_excess(_first(by, m, r.param, contradicted, decided), r)
                                                  for r in reqs), m))
            best = (chosen or {}).get(n) if (chosen or {}).get(n) in fit else minimal
            by_price = [m for _, m in sorted((prices[m].price, m) for m in fit if m in prices)]
            # для выбора: минимально достаточный, затем от дешёвых к дорогим, затем без цены
            pos.fit_models = list(dict.fromkeys([minimal, *by_price, *fit]))
            pos.minimal, pos.model = minimal, best
            pos.alternatives = [m for m in pos.fit_models if m != best][:3]
            if by_price and by_price[0] != best and (best not in prices or prices[by_price[0]].price < prices[best].price):
                pos.cheapest, pos.cheapest_price = by_price[0], prices[by_price[0]].price
            pos.checks = [_check(best, r.param, r, by, contradicted, decided, ref) for r in reqs]
            pos.status = "offer"
            pos.note = ("предложена минимально достаточная модель — с наименьшим запасом по требованиям; "
                        "окончательный выбор — за менеджером") if best == minimal else (
                        f"модель выбрана менеджером; минимально достаточная — {minimal}")
            continue
        unknown = [m for m in candidates if state[m] == "unknown"]
        if unknown:
            pos.status, pos.model = "expert", unknown[0]
            pos.checks = [_check(unknown[0], r.param, r, by, contradicted, decided, ref) for r in reqs]
            for c in pos.checks:
                if c.ok is None and (unknown[0], c.param) in contradicted:
                    items = by[(unknown[0], c.param)]
                    pending.append((spr.subject(unknown[0], c.param),
                                    [{"value": v.shown(), "doc": v.doc, "page": v.page} for v in items]))
            pos.note = "по части требований нет данных или источники расходятся — нужна проверка эксперта"
            unfit = [m for m in candidates if state[m] == "unfit"]
            if unfit:          # модели с известными параметрами не подходят — показать ближайшую
                closest = min(unfit, key=lambda m: (sum(1 for c in checks[m] if c.ok is False), m))
                cs = [_check(closest, r.param, r, by, contradicted, decided, ref) for r in reqs]
                missed = "; ".join(f"{c.param.lower()} {c.value} [{c.ref}] при требовании {c.need}"
                                   for c in cs if c.ok is False)
                pos.note = (f"ни одна модель с известными параметрами не подходит (ближайшая — {closest}: {missed}); "
                            f"у {', '.join(unknown[:3])} нужных данных в справочнике нет или источники расходятся — "
                            "нужна проверка эксперта")
            continue
        # ни одна модель не подходит — показываем ближайшую и спрашиваем клиента
        closest = min(candidates, key=lambda m: sum(1 for c in checks[m] if c.ok is False))
        pos.status, pos.model = "no_fit", closest
        pos.checks = [_check(closest, r.param, r, by, contradicted, decided, ref) for r in reqs]
        missed = "; ".join(f"{c.param.lower()} {c.value} при требовании {c.need}" for c in pos.checks if c.ok is False)
        pos.note = f"ни одна модель в справочнике не подходит; ближайшая — {closest}: {missed}"
        pos.ask_client = f"{kind}: в нашем ассортименте нет модели под все требования ({missed}) — допустимо ли?"
    for pos in positions:
        pos.checks = _merge_checks(pos.checks)
    for pos in positions:
        if pos.status not in ("offer", "named"):
            continue
        pos.alt_prices = {m: prices[m].price for m in pos.alternatives if m in prices}
        if pos.model in prices:
            pr = prices[pos.model]
            pos.price, pos.lead, pos.price_since = pr.price, pr.lead, pr.since
            pos.price_note = "; ".join(x for x in ("ДЕМО" if pr.demo else "",
                                                   f"прайс старше {prajs.STALE_DAYS} дней — проверьте цену"
                                                   if pr.stale else "") if x)
    return Tkp(positions, sources, [], pending, price_list=bool(prices), letter=letter)


def _req_texts(reqs: list) -> list[str]:
    """Требования для показа; нижняя и верхняя граница одного параметра — одной строкой: «температура жала
    паяльника от 200 до 480 °C»."""
    out, merged = [], set()
    for r in reqs:
        if r.param in merged:
            continue
        lo = next((x for x in reqs if x.param == r.param and x.mode == "lo"), None)
        hi = next((x for x in reqs if x.param == r.param and x.mode == "hi"), None)
        if lo and hi:
            merged.add(r.param)
            out.append(f"{r.param.lower()} от {lo.shown.rsplit(' ', 1)[0]} до {hi.shown}")
            continue
        out.append(f"{r.param.lower()} {spr.REQ_WORDS[r.mode]} {r.shown}".replace("  ", " "))
    return out


def _merge_checks(checks: list) -> list:
    """Две проверки одного значения («от 200 °C» и «до 480 °C») — одна: «нужно от 200 до 480 °C»; не подходит,
    если не выполнена любая из границ."""
    out: list[Check] = []
    for c in checks:
        prev = next((x for x in out if (x.param, x.value, x.ref) == (c.param, c.value, c.ref)), None)
        if prev is None or not (prev.need.startswith("от ") and c.need.startswith("до ")):
            out.append(c)
            continue
        prev.need = f"{prev.need.rsplit(' ', 1)[0]} {c.need}"
        prev.ok = False if False in (prev.ok, c.ok) else None if None in (prev.ok, c.ok) else True
    return out


def _first(by, m, param, contradicted, decided):
    items, _, _ = spr.pick(m, param, by, contradicted, decided)
    return items[0]


def _check(m, param, r, by, contradicted, decided, ref, lazy: bool = False) -> Check:
    """Проверка одного требования у модели. lazy — без номера источника (для моделей, которые не покажем)."""
    items, state, d = spr.pick(m, param, by, contradicted, decided)
    need = f"{spr.REQ_WORDS[r.mode]} {r.shown}".strip() if r else ""
    if state == "missing":
        return Check(param, need, "", None, None, "нет данных")
    if state in ("contradiction", "noted"):
        value = "; ".join(v.shown() for v in items)
        refs = [ref(v) for v in items][0] if not lazy else None
        return Check(param, need, value, refs, None, "источники расходятся")
    v = items[0]
    ok = spr._fits(v, r) if r else None
    return Check(param, need, v.shown(), None if lazy else ref(v), ok,
                 spr.confirmed_note(d) if state == "confirmed" else "")


def _other_name(seg: str) -> str:
    m = OTHER_RE.search(seg)
    return (m.group(0)[:1].upper() + m.group(0)[1:]) if m else "Позиция"


def _typical(kind: str) -> str:
    return {"Осциллограф": "полоса пропускания, число каналов, частота дискретизации",
            "Мультиметр": "что измерять и до каких пределов: напряжение, ток, сопротивление",
            "Паяльная станция": "температура пайки, мощность паяльника, нужен ли фен"}.get(kind, "основные параметры")


def is_request(text: str, tkp: Tkp | None) -> bool:
    """Письмо или заявка клиента, а не обычный вопрос: просят ТКП или предложение; несколько позиций с
    требованиями; длинное письмо. «Клиенту нужен осциллограф от 150 МГц — что предложить?» остаётся вопросом
    к справочнику, «сколько осциллографов и мультиметров на поверке?» — к учётной базе."""
    if tkp is None:
        return False
    if TKP_WORDS_RE.search(text):
        return True                     # просят ТКП — отвечаем разбором, даже если продукции нет в справочнике
    if not any(p.status != "not_in_catalog" for p in tkp.positions):
        return False
    specified = any(p.reqs for p in tkp.positions)
    return specified and (len(tkp.positions) >= 2 or (len(text) > 250 and chr(10) in text.strip()))


LETTER_TYPES = ["docx", "pdf", "txt", "eml"]
NORMO_RE = re.compile(r"оформлени\w*|нормоконтрол\w*|\bГОСТ\b|7\.0\.97", re.I)


def letter_text(filename: str, data: bytes) -> str:
    """Текст письма клиента из приложенного файла: Word, PDF (скан — с распознаванием), текст, письмо .eml."""
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "txt":
        for enc in ("utf-8-sig", "cp1251"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        return ""
    if ext == "eml":
        from email import policy
        from email.parser import BytesParser
        msg = BytesParser(policy=policy.default).parsebytes(data)
        body = msg.get_body(preferencelist=("plain", "html"))
        text = body.get_content() if body is not None else ""
        if body is not None and body.get_content_type() == "text/html":
            text = re.sub(r"<[^>]+>", " ", re.sub(r"(?i)<br\s*/?>|</p>", "\n", text))
        return f"{msg.get('subject', '')}\n{text}".strip()
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            lines += [" | ".join(c.text for c in row.cells) for row in table.rows]
        return "\n".join(lines)
    if ext == "pdf":
        import pymupdf

        from .ingest import _page_text
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return "\n".join(_page_text(page, {}) for page in doc)
    return ""


# ---------- текст ответа, черновик письма и ТКП в Word ----------

def _pos_title(p: Position) -> str:
    return f"{p.kind} — {p.qty} шт." if p.status != "not_in_catalog" else p.kind


def _check_text(c: Check, mark: bool = True) -> str:
    sign = {True: " ✓", False: " ✗", None: ""}[c.ok] if mark else ""
    ref = f" [{c.ref}]" if c.ref else ""
    note = f" ({c.note})" if c.note and c.note != "нет данных" else ""
    need = f", нужно {c.need}" if c.need and c.ok is not True else ""
    return f"{c.param.lower()} — {c.value or 'нет данных'}{ref}{note}{need}{sign}"


def _price_line(p: dict, price_list: bool) -> str:
    """«Цена по прайсу от 01.10.2026: 41 500 ₽ × 2 шт. = 83 000 ₽; срок поставки — в наличии.»"""
    if not p.get("price"):
        return "В прайсе цены нет — впишет менеджер." if price_list else ""
    demo = " (ДЕМО)" if "ДЕМО" in p.get("price_note", "") else ""
    line = (f"Цена по прайсу от {prajs.shown_date(p['price_since'])}{demo}: {prajs.money(p['price'])} × {p['qty']} шт. "
            f"= {prajs.money(p['price'] * p['qty'])}")
    if p.get("lead"):
        line += f"; срок поставки — {p['lead']}"
    if "старше" in p.get("price_note", ""):
        line += f"; прайс старше {prajs.STALE_DAYS} дней — проверьте цену"
    return line + "."


def _vat(total: float) -> tuple[float, float]:
    """НДС и сумма с НДС для итога по прайсу — по настройкам KB_VAT_RATE и KB_PRICES_WITH_VAT."""
    rate = spr.settings.vat_rate
    if spr.settings.prices_include_vat:
        return round(total * rate / (100 + rate), 2), total
    vat = round(total * rate / 100, 2)
    return vat, total + vat


def _rate() -> str:
    return f"{spr.settings.vat_rate:g}".replace(".", ",")


def _vat_note(total: float) -> str:
    """« (в т. ч. НДС 22 % — 70 409,84 ₽)», « + НДС 22 % … = …» или « (без НДС)»."""
    if not spr.settings.vat_rate:
        return " (без НДС)"
    vat, gross = _vat(total)
    if spr.settings.prices_include_vat:
        return f" (в т. ч. НДС {_rate()} % — {prajs.money(vat)})"
    return f" без НДС; НДС {_rate()} % — {prajs.money(vat)}, всего с НДС — {prajs.money(gross)}"


def _priced(positions: list[dict]) -> list[dict]:
    return [p for p in positions if p["status"] in ("offer", "named") and p.get("price")]


def has_offer(t: Tkp) -> bool:
    """Есть хотя бы одна позиция из справочника — иначе это честный отказ, а не ТКП."""
    return any(p.status != "not_in_catalog" for p in t.positions)


def answer_text(t: Tkp) -> str:
    if not has_offer(t):
        kinds = ", ".join(p.kind.lower() for p in t.positions)
        return (f"Разбор запроса клиента: позиций — {len(t.positions)} ({kinds}), но этой продукции нет в справочнике "
                "параметров — подобрать модели и подготовить ТКП по проверенным данным не могу. Передайте запрос "
                "техническому отделу; если продукция есть в ассортименте, её параметры нужно добавить в справочник.")
    offers = sum(1 for p in t.positions if p.status in ("offer", "named"))
    reqs = sum(len(p.reqs) for p in t.positions)
    checked = sum(p.total for p in t.positions)
    lines = [f"**Разбор запроса клиента:** позиций — {len(t.positions)}, требований — {reqs}, проверено моделей по "
             f"справочнику — {checked} → предложение готово по {offers} из {len(t.positions)}. "
             "Подбор и сравнение с требованиями — программой, без языковой модели. Предлагается минимально "
             "достаточная модель — с наименьшим запасом по требованиям; окончательный выбор, цены и сроки — "
             "за менеджером."]
    for p in t.positions:
        lines.append("")
        head = f"**{p.n}. {_pos_title(p)}**"
        if p.reqs:
            head += " Требования: " + "; ".join(p.reqs) + "."
        lines.append(head)
        if p.status == "offer":
            head = (f"Предлагаем **{p.model}**" if p.model == p.minimal or not p.minimal else
                    f"Выбрано менеджером: **{p.model}** (минимально достаточная — {p.minimal})")
            lines.append(f"- {head}: " + "; ".join(_check_text(c) for c in p.checks) + ".")
            alts = [m + (f" — {prajs.money(p.alt_prices[m])}" if m in p.alt_prices else "") for m in p.alternatives]
            more = f"Также подходят: {', '.join(alts)}" if alts else "Других подходящих нет"
            lines.append(f"- {more} (всего {p.fit_count} из {p.total}).")
        elif p.status == "named":
            lines.append(f"- **{p.model}**: " + "; ".join(_check_text(c, mark=False) for c in p.checks) + ".")
        elif p.status in ("no_fit", "expert"):
            lines.append(f"- {p.note[:1].upper() + p.note[1:]}.")
            if p.checks:
                lines.append(f"- {p.model}: " + "; ".join(_check_text(c) for c in p.checks) + ".")
        else:
            lines.append(f"- {p.note[:1].upper() + p.note[1:]}.")
        # цена и подсказка о более дешёвом — только у предложенного или названного прибора
        if p.status in ("offer", "named") and (price := _price_line(asdict(p), t.price_list)):
            lines.append(f"- {price}")
        if p.cheapest:
            lines.append(f"- Дешевле всего из подходящих: **{p.cheapest}** — {prajs.money(p.cheapest_price)} "
                         "(прибор меняется в «Черновике ТКП»).")
    priced = _priced([asdict(p) for p in t.positions])
    if priced:
        offered = sum(1 for p in t.positions if p.status in ("offer", "named"))
        total = sum(p["price"] * p["qty"] for p in priced)
        lines += ["", f"**Итого по прайсу:** {prajs.money(total)}{_vat_note(total)} — позиций с ценой {len(priced)} "
                      f"из {offered}; цены и сроки проверяет менеджер."]
    asks = [p.ask_client for p in t.positions if p.ask_client]
    if asks:
        lines += ["", "**Уточнить у клиента:**"] + [f"- {a}" for a in asks]
    if t.expert_ids or t.pending:
        lines += ["", "Расхождения источников переданы эксперту — решение появится под этим ответом."]
    lines += ["", "Это черновик: кнопка «Черновик ТКП» — письмо клиенту и ТКП в Word с таблицей позиций; "
              "утверждает сотрудник."]
    return "\n".join(lines)


def rows(t: Tkp) -> list[list]:
    out = []
    for p in t.positions:
        verdict = {"offer": "предлагаем", "named": "по запросу", "no_fit": "не подходит", "expert": "эксперту",
                   "no_reqs": "уточнить", "not_in_catalog": "нет в справочнике"}[p.status]
        if not p.checks:
            out.append([p.n, p.kind, p.qty, p.model or "—", "—", "—", verdict])
        for c in p.checks:
            out.append([p.n, p.kind, p.qty, p.model, c.param, (c.value or "нет данных")
                        + (f" (нужно {c.need})" if c.need else ""), verdict])
    return out


ROW_LABELS = ["№", "Позиция", "Кол-во", "Модель", "Параметр", "Значение", "Итог"]


def draft_text(t: dict, signer: str = "") -> str:
    """Сопроводительное письмо к ТКП. Позиции и характеристики — из справочника (в таблице Word их не правят),
    цены и сроки — пропуски для менеджера."""
    lines = ["Добрый день!", "", "Благодарим за запрос. По вашим требованиям предлагаем:"]
    for p in t["positions"]:
        if p["status"] in ("offer", "named"):
            price = (f"Цена: {prajs.money(p['price'])} за шт., сумма {prajs.money(p['price'] * p['qty'])}."
                     if p.get("price") else "Цена: [заполняет менеджер].")
            lines.append(f"{p['n']}. {p['kind']} {p['model']} — {p['qty']} шт. {price}")
        elif p["status"] == "no_fit":
            lines.append(f"{p['n']}. {p['kind']} — подходящей модели под все требования нет, предлагаем обсудить "
                         f"варианты (ближайшая — {p['model']}).")
        elif p["status"] == "expert":
            lines.append(f"{p['n']}. {p['kind']} — уточняем характеристики у технического специалиста.")
        else:
            lines.append(f"{p['n']}. {p['kind']} — уточняем возможность поставки.")
    asks = [p["ask_client"] for p in t["positions"] if p.get("ask_client")]
    if asks:
        lines += ["", "Просим уточнить:"] + [f"- {a}" for a in asks]
    priced = _priced(t["positions"])
    offered = [p for p in t["positions"] if p["status"] in ("offer", "named")]
    lines += ["", "Соответствие характеристик вашим требованиям — в таблице ТКП; значения приведены по руководствам "
              "по эксплуатации (источники — под таблицей)."]
    if priced:
        since = ", ".join(dict.fromkeys(prajs.shown_date(p["price_since"]) for p in priced))
        total = sum(p["price"] * p["qty"] for p in priced)
        lines.append(f"Цены указаны по прайс-листу от {since}. Итого: {prajs.money(total)}{_vat_note(total)}.")
    leads = {p["n"]: p.get("lead") for p in offered}
    if offered and all(leads.values()) and len(set(leads.values())) == 1:
        lines.append(f"Срок поставки: {next(iter(leads.values()))}.")
    elif any(leads.values()):
        lines.append("Срок поставки: " + "; ".join(f"поз. {n} — {lead or '[заполняет менеджер]'}"
                                                   for n, lead in leads.items()) + ".")
    else:
        lines.append("Срок поставки: [заполняет менеджер].")
    lines += ["", "С уважением,", signer or "[ФИО, должность]"]
    return "\n".join(lines)


def tkp_docx(text: str, t: dict, approved_by: str, approved_at: str, app_name: str) -> bytes:
    """ТКП в Word: заголовок, письмо, таблица позиций (цена и сумма — для менеджера), источники, отметка."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Mm, Pt, RGBColor

    from .export import _runs
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(20), Mm(10)
        section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    style.paragraph_format.space_after = Pt(0)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ТЕХНИКО-КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ")
    run.bold = True
    doc.add_paragraph(f"№ ____ от {approved_at[:10]}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    lines = text.splitlines()
    # таблица позиций — перед подписью («С уважением, …»), как в обычном ТКП
    cut = next((k for k, line in enumerate(lines) if line.strip().lower().startswith("с уважением")), len(lines))

    def paragraphs(part):
        for line in part:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(1.25) if line.strip() else None
            _runs(p, line.rstrip())

    paragraphs(lines[:cut])
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ["№", "Наименование", "Кол-во, шт.", "Соответствие требованиям",
                                                "Цена за ед., ₽", "Сумма, ₽"]):
        cell.text = head
        cell.paragraphs[0].runs[0].bold = True
    for p in t["positions"]:
        cells = table.add_row().cells
        name = f"{p['kind']} {p['model']}" if p["status"] in ("offer", "named") else p["kind"]
        checks = "\n".join(f"{c['param'][:1].upper() + c['param'][1:]}: {c['value'] or 'нет данных'}"
                           + (f" [{c['ref']}]" if c.get("ref") else "")
                           + (" — соответствует" if c["ok"] else " — не соответствует" if c["ok"] is False else "")
                           + (f" (требование: {c['need']})" if c["need"] else "") for c in p["checks"])
        status = {"no_fit": "подходящей модели нет", "expert": "уточняется у технического специалиста",
                  "no_reqs": "требования уточняются у клиента", "not_in_catalog": "уточняется возможность поставки"}
        cells[0].text, cells[1].text, cells[2].text = str(p["n"]), name, str(p["qty"])
        cells[3].text = checks or status.get(p["status"], "")
        has_price = p["status"] in ("offer", "named") and p.get("price")
        cells[4].text = prajs.money(p["price"]) if has_price else ""
        cells[5].text = prajs.money(p["price"] * p["qty"]) if has_price else ""
    priced = _priced(t["positions"])
    if priced:
        cells = table.add_row().cells
        total = sum(p["price"] * p["qty"] for p in priced)
        cells[1].text = "Итого по прайсу"
        cells[5].text = prajs.money(total)
        cells[1].paragraphs[0].runs[0].bold = cells[5].paragraphs[0].runs[0].bold = True
        vat, gross = _vat(total)
        rows = ([] if not spr.settings.vat_rate else
                [(f"в т. ч. НДС {_rate()} %", vat)] if spr.settings.prices_include_vat else
                [(f"НДС {_rate()} %", vat), ("Всего с НДС", gross)])
        for label, amount in rows or [("Без НДС", None)]:
            cells = table.add_row().cells
            cells[1].text, cells[5].text = label, prajs.money(amount) if amount is not None else ""
    for row in table.rows:
        for cell, width in zip(row.cells, (Cm(1), Cm(4.2), Cm(1.8), Cm(7), Cm(2.4), Cm(2.4))):
            cell.width = width
            for par in cell.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(10)
    if t.get("sources"):
        doc.add_paragraph()
        for line in ["Источники характеристик:"] + [f"[{s['n']}] {s['label']}" for s in t["sources"]]:
            doc.add_paragraph().add_run(line).font.size = Pt(10)
    doc.add_paragraph()
    paragraphs(lines[cut:])
    doc.add_paragraph()
    note = doc.add_paragraph()
    offered = [p for p in t["positions"] if p["status"] in ("offer", "named")]
    if priced:
        since = ", ".join(dict.fromkeys(prajs.shown_date(p["price_since"]) for p in priced))
        demo = " (ДЕМО — цены вымышленные)" if any("ДЕМО" in p.get("price_note", "") for p in priced) else ""
        stale = (f"; прайс старше {prajs.STALE_DAYS} дней — цены проверить"
                 if any("старше" in p.get("price_note", "") for p in priced) else "")
        missing = "; чего нет в прайсе, вписывает менеджер" if len(priced) < len(offered) else ""
        prices = f"Цены и сроки — по прайсу от {since}{demo}{stale}{missing}, их проверяет менеджер."
    else:
        prices = "Цены и сроки вписывает менеджер."
    run = note.add_run(f"Черновик подготовлен: {app_name}, характеристики и соответствие требованиям — по справочнику "
                       f"параметров со ссылками на руководства. {prices} "
                       f"Утвердил: {approved_by}, {approved_at}.")
    run.italic, run.font.color.rgb, run.font.size = True, RGBColor(0x5C, 0x54, 0x70), Pt(10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
