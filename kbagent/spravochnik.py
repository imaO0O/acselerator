"""Справочник параметров продукции: импорт с проверкой качества данных, доступ только на чтение,
проверка требований клиента и подбор моделей — программой, без языковой модели.

Файл справочника — CSV (разделитель «;») или XLSX с колонками: Модель, Вид, Параметр, Значение, Единица,
Условие, Документ, Страница, Цитата. Значение — число («100»), диапазон («50–350») или текст («CAT II, 600 В»).
У каждого значения есть источник: документ базы знаний, страница и цитата — она подсвечивается на странице.

При загрузке проверяются пропуски, единицы измерения, числа, страницы и документы; повторяющиеся строки
не дублируются; если для одной модели и параметра источники дают разные значения — это противоречие:
оно показывается в ответе и передаётся эксперту, а вывод «подходит / не подходит» не делается.
"""
import csv
import io
import json
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

from .config import settings

DIR = settings.raw_dir.parent / "spravochnik"
PATH = DIR / "pribory.csv"
REPORT = DIR / "import.json"
COLUMNS = ["Модель", "Вид", "Параметр", "Значение", "Единица", "Условие", "Документ", "Страница", "Цитата"]
REQUIRED = ["Модель", "Параметр", "Значение", "Документ", "Страница"]

# единица → (величина, множитель к основной единице)
UNITS = {
    "Гц": ("частота", 1), "кГц": ("частота", 1e3), "МГц": ("частота", 1e6), "ГГц": ("частота", 1e9),
    "выб/с": ("дискретизация", 1), "Мвыб/с": ("дискретизация", 1e6), "Гвыб/с": ("дискретизация", 1e9),
    "Гц/с": ("Гц/с", 1),
    "мВ": ("напряжение", 1e-3), "В": ("напряжение", 1), "кВ": ("напряжение", 1e3),
    "мкА": ("ток", 1e-6), "мА": ("ток", 1e-3), "А": ("ток", 1),
    "Ом": ("сопротивление", 1), "кОм": ("сопротивление", 1e3), "МОм": ("сопротивление", 1e6),
    "°C": ("температура", 1),
    "Вт": ("мощность", 1), "кВт": ("мощность", 1e3), "ВА": ("полная мощность", 1),
    "пс": ("время", 1e-12), "нс": ("время", 1e-9), "мкс": ("время", 1e-6), "мс": ("время", 1e-3), "с": ("время", 1),
    "г": ("масса", 1e-3), "кг": ("масса", 1),
    "": ("количество", 1),
}
UNIT_ALIASES = {"ºС": "°C", "°С": "°C", "ºC": "°C", "℃": "°C", "C": "°C", "град": "°C", "Гвыб./с": "Гвыб/с",
                "ГС/с": "Гвыб/с", "МС/с": "Мвыб/с", "V": "В", "mV": "мВ", "A": "А", "mA": "мА", "W": "Вт",
                "шт": "", "шт.": ""}
NUM = r"[-+]?\d+(?:[.,]\d+)?"


@dataclass
class Value:
    model: str
    kind: str
    param: str
    lo: float | None          # число или нижняя граница диапазона
    hi: float | None          # верхняя граница (у числа — то же значение)
    unit: str
    text: str                 # текстовое значение («CAT II, 600 В») или исходная запись
    cond: str
    doc: str
    page: int
    quote: str

    @property
    def dim(self) -> str:
        return UNITS[self.unit][0] if self.lo is not None else "текст"

    def base(self, x: float | None) -> float | None:
        return None if x is None else x * UNITS[self.unit][1]

    @property
    def key(self) -> tuple:
        return (self.dim, self.base(self.lo), self.base(self.hi), self.text if self.lo is None else "")

    def shown(self) -> str:
        if self.lo is None:
            return self.text
        num = lambda x: f"{x:g}".replace(".", ",")  # noqa: E731
        value = num(self.lo) if self.lo == self.hi else f"{num(self.lo)}–{num(self.hi)}"
        prefix = "±" if self.cond == "±" else ""
        unit = f" {self.unit}" if self.unit else ""
        note = f" ({self.cond})" if self.cond and self.cond != "±" else ""
        return f"{prefix}{value}{unit}{note}"


@dataclass
class Report:
    loaded_at: str = ""
    source: str = ""
    rows: int = 0
    values: int = 0
    models: int = 0
    errors: list = field(default_factory=list)          # строка файла и причина — строка не загружена
    duplicates: int = 0
    contradictions: list = field(default_factory=list)  # {model, param, values: [{value, doc, page}]}


def _parse_number(s: str) -> float:
    return float(s.replace(",", ".").replace(" ", ""))


def _read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).active
        data = [[("" if v is None else str(v)).strip() for v in row] for row in sheet.iter_rows(values_only=True)]
        head = data[0] if data else []
        return [dict(zip(head, r)) for r in data[1:] if any(r)]
    text = path.read_bytes().decode("utf-8-sig")
    delim = ";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") else ","
    return [{k.strip(): (v or "").strip() for k, v in r.items() if k} for r in csv.DictReader(io.StringIO(text), delimiter=delim)]


def parse(path: Path, known_docs: set[str] | None = None) -> tuple[list[Value], Report]:
    """Читает и проверяет справочник. Ошибочные строки не загружаются, повторы не дублируются."""
    rep = Report(loaded_at=time.strftime("%Y-%m-%d %H:%M:%S"), source=path.name)
    raw = _read_rows(path)
    rep.rows = len(raw)
    missing_cols = [c for c in REQUIRED if raw and c not in raw[0]]
    if missing_cols:
        rep.errors.append({"line": 1, "error": f"нет колонок: {', '.join(missing_cols)}"})
        return [], rep
    values, seen = [], set()
    for line, r in enumerate(raw, start=2):
        empty = [c for c in REQUIRED if not r.get(c)]
        if empty:
            rep.errors.append({"line": line, "error": f"не заполнено: {', '.join(empty)}"})
            continue
        unit = UNIT_ALIASES.get(r.get("Единица", ""), r.get("Единица", ""))
        if unit not in UNITS:
            rep.errors.append({"line": line, "error": f"неизвестная единица «{r.get('Единица')}»"})
            continue
        if not r["Страница"].isdigit() or int(r["Страница"]) < 1:
            rep.errors.append({"line": line, "error": f"страница «{r['Страница']}» — не номер"})
            continue
        if known_docs is not None and r["Документ"] not in known_docs:
            rep.errors.append({"line": line, "error": f"документа «{r['Документ']}» нет в базе знаний"})
            continue
        v = r["Значение"]
        rng = re.fullmatch(rf"\s*({NUM})\s*(?:–|-|—|…|\.\.\.?|~)\s*({NUM})\s*", v)
        if rng:
            lo, hi, text = _parse_number(rng.group(1)), _parse_number(rng.group(2)), v
        elif re.fullmatch(rf"\s*{NUM}\s*", v):
            lo = hi = _parse_number(v)
            text = v
        elif unit == "":
            lo = hi = None
            text = v
        else:
            rep.errors.append({"line": line, "error": f"значение «{v}» — не число, а единица «{unit}» задана"})
            continue
        item = Value(r["Модель"], r.get("Вид", ""), r["Параметр"], lo, hi, unit, text, r.get("Условие", ""),
                     r["Документ"], int(r["Страница"]), r.get("Цитата", ""))
        sig = (item.model, item.param, item.key, item.doc, item.page)
        if sig in seen:
            rep.duplicates += 1
            continue
        seen.add(sig)
        values.append(item)
    rep.values, rep.models = len(values), len({v.model for v in values})
    groups: dict[tuple, list[Value]] = {}
    for v in values:
        groups.setdefault((v.model, v.param), []).append(v)
    for (model, param), items in groups.items():
        if len({v.key for v in items}) > 1:
            rep.contradictions.append({"model": model, "param": param,
                                       "values": [{"value": v.shown(), "doc": v.doc, "page": v.page} for v in items]})
    return values, rep


_CACHE: dict = {}


def load(path: Path | None = None) -> tuple[list[Value], Report]:
    """Справочник с проверкой; перечитывается, только когда файл изменился."""
    path = path or PATH
    if not path.exists():
        return [], Report(errors=[{"line": 0, "error": "файл справочника не загружен"}])
    stamp = (str(path), path.stat().st_mtime)
    if _CACHE.get("stamp") != stamp:
        from .ingest import load_catalog
        values, rep = parse(path, set(load_catalog(settings.raw_dir)))
        _CACHE.update(stamp=stamp, values=values, report=rep)
        try:
            REPORT.write_text(json.dumps(rep.__dict__, ensure_ascii=False, indent=1), encoding="utf-8")
        except OSError:
            pass
    return _CACHE["values"], _CACHE["report"]


def replace_file(filename: str, data: bytes) -> Report:
    """Новый файл справочника от администратора: проверяется до замены; при ошибках в каждой строке не ставится."""
    from .ingest import load_catalog
    ext = Path(filename).suffix.lower()
    if ext not in (".csv", ".xlsx"):
        raise ValueError("поддерживаются CSV и XLSX")
    DIR.mkdir(parents=True, exist_ok=True)
    tmp = DIR / f"upload{ext}"
    tmp.write_bytes(data)
    values, rep = parse(tmp, set(load_catalog(settings.raw_dir)))
    if not values:
        tmp.unlink()
        raise ValueError("ни одна строка не прошла проверку: " + "; ".join(e["error"] for e in rep.errors[:3]))
    if PATH.exists():
        PATH.replace(DIR / "pribory.prev.csv")
    with open(PATH, "w", encoding="utf-8-sig", newline="") as f:     # хранится в одном формате — CSV
        w = csv.writer(f, delimiter=";")
        w.writerow(COLUMNS)
        for v in values:
            w.writerow([v.model, v.kind, v.param, v.text, v.unit, v.cond, v.doc, v.page, v.quote])
    tmp.unlink()
    _CACHE.clear()
    return load()[1]


# ---------- понимание вопроса ----------

LAT = str.maketrans("АВЕКМНОРСТХУавекмнорстху", "ABEKMHOPCTXYABEKMHOPCTXY")
KINDS = {"Осциллограф": r"осциллограф\w*", "Мультиметр": r"мультиметр\w*",
         "Паяльная станция": r"паяльн\w+\s+станц\w*|станци\w+\s+для\s+пайки|паяльник\w*|термофен\w*|\bфен\w*"}

PARAMS = {   # параметр → слова в вопросе
    "Полоса пропускания": r"полос\w*",
    "Число каналов": r"канал\w*|канальн\w*",
    "Частота дискретизации на канал": r"дискретизац\w*|выборок",
    "Время нарастания": r"нарастани\w*",
    "Максимальное входное напряжение": r"входн\w*\s+напряжени\w*",
    "Постоянное напряжение, верхний предел": r"постоянн\w*\s+напряжени\w*|\bDCV\b",
    "Переменное напряжение, верхний предел": r"переменн\w*\s+напряжени\w*|\bACV\b",
    "Постоянный ток, верхний предел": r"постоянн\w*\s+ток\w*|\bDCA\b",
    "Переменный ток, верхний предел": r"переменн\w*\s+ток\w*|\bACA\b",
    "Сопротивление, верхний предел": r"сопротивлени\w*",
    "Категория безопасности": r"категори\w*|\bCAT\b",
    "Масса": r"\bмасс\w*|\bвес\b|\bвеси\w*|\bвесом\b",
    "Стабильность температуры паяльника": r"стабильн\w*|погрешност\w*\s+(?:контрол\w*\s+)?температур\w*",
    "Потребляемая мощность": r"потребля\w*",
}
GENERIC_ASK = r"характеристик\w*|параметр\w*|технические\s+данн\w*"
# температура окружающего воздуха, хранения — условия эксплуатации, а не параметр в справочнике
NOT_PARAM_RE = re.compile(r"окружающ\w*|в\s+помещени\w*|воздух\w*\s+в\s+помещ|хранени\w*|транспортир\w*|"
                          r"влажност\w*", re.I)
PRICE_RE = re.compile(r"стоит|стоимост\w*|\bцен[аыуе]?\b|прайс", re.I)
REQ_RE = re.compile(rf"({NUM})\s*(ГГц|МГц|кГц|Гц|Гвыб/с|Мвыб/с|кВ|мВ|В|мкА|мА|А|МОм|кОм|Ом|°\s?[CС]|градус\w*|"
                    rf"кВт|Вт|ВА|пс|нс|мкс|мс|кг|г)(?![А-Яа-яA-Za-z])")
CHANNELS_RE = re.compile(r"(\d+)\s*-?\s*(?:х\s*)?канал\w*|(двух|трёх|трех|четырёх|четырех)\s*канальн\w*|"
                         # «каналов ≥ 4», «число каналов не менее 4», «каналов: 4» — число после слова
                         r"канал\w*\s*(?:≥|>=|:|не\s+менее|не\s+меньше|минимум|от)\s*(\d+)(?!\d)(?!\s*шт)|"
                         # «с четырьмя каналами», «два канала» — так говорят (голосовой вопрос) и пишут
                         r"\b(два|двумя|двух|три|тремя|трёх|трех|четыре|четырьмя|четырёх|четырех|восемь|восьмью|"
                         r"восемью|восьми)\s+канал\w*", re.I)
# Как пишут в письмах: опечатки в видах продукции, «МГЦ», единицы латиницей. Однобуквенные латинские единицы —
# только через пробел: «10 A» — ток, а «DT9205A» — модель.
TYPO_FIXES = [(re.compile(r"(?<![а-яё])[оа]с?ц[еиы]л{1,2}[оа]?граф", re.I), "осциллограф"),
              (re.compile(r"мульт[иы]мет[еи]?р", re.I), "мультиметр")]
HZ_RE = re.compile(r"(\d)\s*([кКмМгГ]?)[гГ][цЦ](?![а-яё])")
LATIN_UNITS = {"GHz": "ГГц", "GHZ": "ГГц", "MHz": "МГц", "MHZ": "МГц", "kHz": "кГц", "KHz": "кГц", "Hz": "Гц",
               "kV": "кВ", "mV": "мВ", "mA": "мА", "kW": "кВт", "MOhm": "МОм", "kOhm": "кОм", "Ohm": "Ом",
               "ns": "нс", "ps": "пс", "kg": "кг"}
LATIN_SHORT = {"A": "А", "V": "В", "W": "Вт"}
LATIN_RE = re.compile(rf"(?<![\w/.-])({NUM})\s*({'|'.join(sorted(LATIN_UNITS, key=len, reverse=True))})(?![A-Za-z])"
                      rf"|(?<![\w/.-])({NUM})\s+([AVW])(?![\w])")
WORD_NUM = {"два": 2, "двумя": 2, "двух": 2, "три": 3, "тремя": 3, "трёх": 3, "трех": 3, "четыре": 4,
            "четырьмя": 4, "четырёх": 4, "четырех": 4, "восемь": 8, "восьмью": 8, "восемью": 8, "восьми": 8}
# у этих параметров «лучше меньше»: требование «X» — значит не больше X
SMALLER_BETTER = {"Время нарастания", "Масса", "Потребляемая мощность", "Стабильность температуры паяльника"}
RANGE_PARAMS = {"Температура жала паяльника", "Температура фена"}


def normalize(text: str) -> str:
    """Письмо или вопрос — к написанию справочника: «осцилограф» → «осциллограф», «150 MHz» → «150 МГц»."""
    for rx, right in TYPO_FIXES:          # заглавная буква сохраняется: «Осцилограф» → «Осциллограф»
        text = rx.sub(lambda m: (right[0].upper() if m.group(0)[0].isupper() else right[0]) + right[1:], text)
    text = HZ_RE.sub(lambda m: m.group(1) + " " + {"к": "к", "м": "М", "г": "Г"}.get(m.group(2).lower(), "") + "Гц",
                     text)
    return LATIN_RE.sub(lambda m: f"{m.group(1)} {LATIN_UNITS[m.group(2)]}" if m.group(1)
                        else f"{m.group(3)} {LATIN_SHORT[m.group(4)]}", text)


@dataclass
class Req:
    param: str
    value: float              # в основных единицах
    shown: str                # как в вопросе: «150 МГц»
    mode: str                 # min — значение не меньше; max — не больше; inside — входит в диапазон;
                              # hi — верх диапазона не ниже; lo — низ диапазона не выше


@dataclass
class Query:
    models: list[str]
    kind: str | None
    params: list[str]
    reqs: list[Req]
    all_params: bool = False


def _norm(s: str) -> str:
    """Для сравнения обозначений: латиница вместо похожих русских букв, без дефисов, скобок и кавычек;
    пробелы между словами сохраняются — по ним видны границы обозначения в вопросе."""
    s = re.sub(r"[\-–—«»\"'()]", "", s.translate(LAT).upper())
    return re.sub(r"\s+", " ", s).strip()


def _aliases(models: set[str]) -> dict[str, set[str]]:
    """Как модель могут назвать в вопросе: полностью, без марки («4115/1В»), без «(V)»; серия — по номеру до «/»."""
    out: dict[str, set[str]] = {}
    for m in models:
        full = _norm(m)
        keys = {full, full.replace("V", "") if full.endswith("V") else full}
        no_brand = re.sub(r"^[A-ZА-ЯЁ]+(?=\d)", "", full)
        keys |= {no_brand, no_brand[:-1] if no_brand.endswith("V") else no_brand}
        for k in keys:
            if len(k) >= 3 and re.search(r"\d", k):
                out.setdefault(k, set()).add(m)
        for tok in re.split(r"\s+", m):                         # «H95», «852D+», «830B»
            if re.search(r"\d", tok) and len(_norm(tok)) >= 3:
                out.setdefault(_norm(tok), set()).add(m)
        word = re.match(r"[A-Za-zА-Яа-яЁё]{3,}", m)
        if word and not re.search(r"\d", word.group(0)) and word.group(0).upper() != "АКИП":  # «Lukey», «Магистр»
            out.setdefault(_norm(word.group(0)), set()).add(m)
        series = re.match(r"(?:АКИП-)?(\d{4})/", m)
        if series:                                              # «АКИП-4143» — вся серия
            out.setdefault(series.group(1), set()).add(m)
            out.setdefault(_norm("АКИП-" + series.group(1)), set()).add(m)
    return out


def _find_models(question: str, values: list[Value]) -> list[str]:
    q = " " + re.sub(r"[?!,.;:]", " ", _norm(question)) + " "
    found: list[str] = []
    for key, models in sorted(_aliases({v.model for v in values}).items(), key=lambda kv: -len(kv[0])):
        if re.search(rf"(?<![0-9A-ZА-ЯЁ/]){re.escape(key)}(?![0-9A-ZА-ЯЁ])", q):
            if len(models) > 1 and set(found) & models:
                continue                                        # названа модель — серия целиком не нужна
            found += [m for m in sorted(models) if m not in found]
    # «TDM» и другие марки из названия документа
    for doc_word, models in _brand_words(values).items():
        if found and set(found) & models:
            continue
        if re.search(rf"(?<![0-9A-ZА-ЯЁ]){re.escape(doc_word)}(?![0-9A-ZА-ЯЁ])", q):
            found += [m for m in sorted(models) if m not in found]
    return found


def _brand_words(values: list[Value]) -> dict[str, set[str]]:
    from .ingest import load_catalog
    catalog = load_catalog(settings.raw_dir)
    out: dict[str, set[str]] = {}
    for v in values:
        title = catalog.get(v.doc, {}).get("title", "")
        for w in re.findall(r"\b[A-Z]{3,}\b", title):          # латинские марки в названии: TDM, MASTECH
            out.setdefault(_norm(w), set()).add(v.model)
    return out


def _unit_param(unit: str, kind: str | None, question: str) -> str | None:
    dim = UNITS[unit][0]
    q = question.lower()
    if dim == "частота":
        return "Полоса пропускания"
    if dim == "дискретизация":
        return "Частота дискретизации на канал"
    if dim == "напряжение":
        if kind == "Осциллограф":
            return "Максимальное входное напряжение"
        return "Переменное напряжение, верхний предел" if "переменн" in q else "Постоянное напряжение, верхний предел"
    if dim == "ток":
        return "Переменный ток, верхний предел" if "переменн" in q else "Постоянный ток, верхний предел"
    if dim == "сопротивление":
        return "Сопротивление, верхний предел"
    if dim == "температура":
        return "Температура фена" if re.search(r"\bфен\w*|термофен", q) else "Температура жала паяльника"
    if dim == "мощность":
        if re.search(r"потребля", q):
            return "Потребляемая мощность"
        return "Мощность фена" if re.search(r"\bфен\w*|термофен", q) else "Мощность паяльника"
    if dim == "полная мощность":
        return "Потребляемая мощность"
    if dim == "время":
        return "Время нарастания"
    if dim == "масса":
        return "Масса"
    return None


MODEL_TOKEN_RE = re.compile(r"[A-Za-zА-Яа-яЁё]+-?\d+[\w/().+-]*|\d{3,}/\d+[\wА-Яа-я()+-]*")


def _mode(before: str, param: str) -> str:
    """Как понимать требование по словам перед числом. «не менее» проверяется раньше «менее»."""
    if re.search(r"не\s+(?:менее|меньше|ниже)\s*$", before):
        return "lo" if param in RANGE_PARAMS else "min"
    if re.search(r"не\s+(?:более|больше|выше)\s*$", before):
        return "hi" if param in RANGE_PARAMS else "max"
    if re.search(r"(?:минимум|больше|более|свыше|выше|≥|>=)\s*$", before):
        return "hi" if param in RANGE_PARAMS else "min"
    if re.search(r"(?:максимум|меньше|менее|ниже|≤|<=)\s*$", before):
        return "lo" if param in RANGE_PARAMS else "max"
    if re.search(r"\bдо\s*$", before):          # «измерять ток до 15 А» — нужен предел не ниже 15 А
        return "hi" if param in RANGE_PARAMS else "max" if param in SMALLER_BETTER else "min"
    if re.search(r"\bот\s*$", before):
        return "lo" if param in RANGE_PARAMS else "min"
    if param in RANGE_PARAMS:
        return "inside"                           # «пайка при 450 °C» — значение внутри диапазона
    return "max" if param in SMALLER_BETTER else "min"


def _requirements(question: str, kind: str | None) -> list[Req]:
    """Требования клиента из вопроса. Обозначения моделей вырезаются: «4115/1В» — не «1 В»."""
    question = MODEL_TOKEN_RE.sub(" ", question)
    reqs = []
    # «от 200 до 480 °C» — у нижней границы единица не повторена: у диапазонных параметров это два требования
    for m in re.finditer(rf"\bот\s+({NUM})\s+до\s+{NUM}\s*(°\s?[CС]|градус\w*)", question, re.I):
        param = _unit_param("°C", kind, question)
        if param in RANGE_PARAMS:
            reqs.append(Req(param, _parse_number(m.group(1)), f"{m.group(1)} °C", "lo"))
    for m in REQ_RE.finditer(question):
        raw_unit = m.group(2)
        unit = "°C" if raw_unit.startswith(("°", "градус")) else raw_unit
        param = _unit_param(unit, kind, question)
        if param is None or (param == "Масса" and not re.search(PARAMS["Масса"], question, re.I)):
            continue                                            # «в 2026 г.» — год, а не масса
        value = _parse_number(m.group(1)) * UNITS[unit][1]
        mode = _mode(question[max(0, m.start() - 25):m.start()].lower(), param)
        reqs.append(Req(param, value, f"{m.group(1)} {raw_unit}", mode))
    channels = []
    for m in CHANNELS_RE.finditer(question):
        n = int(m.group(1) or m.group(3)) if (m.group(1) or m.group(3)) else WORD_NUM[
            (m.group(2) or m.group(4)).lower()]
        channels.append(n)
    if channels:                                  # «4 канала … каналов не менее 4» — одно требование
        n = max(channels)
        reqs.append(Req("Число каналов", n, f"{n} каналов", "min"))
    return reqs


def understand(question: str, values: list[Value] | None = None) -> Query | None:
    """Вопрос к справочнику: названа модель или вид продукции и параметр или требование. Иначе — None
    («как заменить батарею в MAS830L» — вопрос к документам, а не к справочнику)."""
    if values is None:
        values = load()[0]
    question = normalize(question)
    if not values or PRICE_RE.search(question):
        return None
    models = _find_models(question, values)
    kind = next((k for k, rx in KINDS.items() if re.search(rx, question, re.I)), None)
    if not models and not kind:
        return None
    if kind is None:
        kind = next((v.kind for v in values if v.model == models[0]), None)
    q = question
    params = [p for p, rx in PARAMS.items() if re.search(rx, q, re.I)]
    if re.search(r"температур\w*|градус\w*", q, re.I) and not NOT_PARAM_RE.search(q):
        params.append("Температура фена" if re.search(r"\bфен\w*|термофен", q, re.I) else "Температура жала паяльника")
    if re.search(r"мощност\w*", q, re.I) and "Потребляемая мощность" not in params:
        params.append("Мощность фена" if re.search(r"\bфен\w*|термофен", q, re.I) else "Мощность паяльника")
    if re.search(r"\bток\w*", q, re.I) and not any("ток" in p for p in params):
        params.append("Переменный ток, верхний предел" if re.search(r"переменн", q, re.I)
                      else "Постоянный ток, верхний предел")
    if re.search(r"напряжени\w*", q, re.I) and not any("напряжени" in p.lower() for p in params):
        params.append("Максимальное входное напряжение" if kind == "Осциллограф" else
                      "Переменное напряжение, верхний предел" if re.search(r"переменн", q, re.I)
                      else "Постоянное напряжение, верхний предел")
    reqs = [] if NOT_PARAM_RE.search(q) else _requirements(q, kind)
    params += [r.param for r in reqs if r.param not in params]
    # «Чем отличаются A и B?» — сравнение по всем параметрам
    all_params = bool(re.search(GENERIC_ASK, q, re.I)) or (
        len(models) >= 2 and bool(re.search(r"отлича\w*|сравни\w*|разниц\w*", q, re.I)))
    if not params and not all_params:
        return None
    return Query(models, kind, params, reqs, all_params)


# ---------- ответ ----------

@dataclass
class SprAnswer:
    status: str
    text: str
    sources: list = field(default_factory=list)
    columns: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    labels: list = field(default_factory=list)
    stats: dict = field(default_factory=dict)


def _fits(v: Value, r: Req) -> bool:
    lo, hi = v.base(v.lo), v.base(v.hi)
    if r.mode == "min":
        return hi >= r.value
    if r.mode == "max":
        return lo <= r.value
    if r.mode == "hi":
        return hi >= r.value
    if r.mode == "lo":
        return lo <= r.value
    return lo <= r.value <= hi


REQ_WORDS = {"min": "не менее", "max": "не более", "hi": "до", "lo": "от", "inside": ""}


def subject(model: str, param: str) -> str:
    return f"{model} — {param}"


def pick(model: str, param: str, by: dict, contradicted: set, decided: dict) -> tuple[list[Value], str, dict | None]:
    """Значения параметра модели и их состояние: ok; missing — нет данных; contradiction — источники расходятся,
    решения нет; confirmed — эксперт выбрал верное значение; noted — эксперт решил текстом, оба значения остаются."""
    items = by.get((model, param), [])
    if not items:
        return [], "missing", None
    if (model, param) not in contradicted:
        return items, "ok", None
    d = decided.get(subject(model, param))
    if d is None:
        return items, "contradiction", None
    ch = d.get("chosen") or {}
    chosen = [v for v in items if v.shown() == ch.get("value") and v.doc == ch.get("doc") and v.page == ch.get("page")]
    return (chosen, "confirmed", d) if chosen else (items, "noted", d)


def confirmed_note(d: dict) -> str:
    day = d.get("resolved_at", "")[:10]
    return f"подтверждено экспертом {day[8:10]}.{day[5:7]}.{day[:4]}" if day else "подтверждено экспертом"


def answer(q: Query, allowed: set[str], question: str = "", user: str | None = None,
           role: str | None = None) -> SprAnswer | None:
    from .experts import decisions
    from .ingest import access_of, load_catalog, registry_suffix
    values, rep = load()
    decided = decisions()
    catalog = load_catalog(settings.raw_dir)
    # роль проверяется до ответа: значения из документов, закрытых для роли, не показываются
    values = [v for v in values if access_of(catalog.get(v.doc, {})) in allowed]
    contradicted = {(c["model"], c["param"]) for c in rep.contradictions}
    sources: list[dict] = []
    index: dict[tuple, int] = {}

    def ref(v: Value) -> str:
        k = (v.doc, v.page, v.quote)
        if k not in index:
            index[k] = len(sources) + 1
            info = catalog.get(v.doc, {})
            registry = registry_suffix(v.doc)
            label = f"{info.get('title', v.doc)}, с. {v.page}" + (f" · {registry}" if registry else "")
            sources.append({"n": index[k], "label": label, "text": v.quote or v.text,
                            "doc_id": v.doc, "doc_kind": info.get("kind", ""), "clause": "", "page": v.page,
                            "page_end": v.page, "cited": True, "score": 1.0})
        return f"[{index[k]}]"

    by = {}
    for v in values:
        by.setdefault((v.model, v.param), []).append(v)
    models = q.models or sorted({v.model for v in values if v.kind == q.kind})
    if not models:
        return None
    params = q.params or []
    if q.all_params and q.models:
        params = list(dict.fromkeys(params + [v.param for v in values if v.model in q.models]))
    rows, lines, experts = [], [], []

    confirmed = []
    demand = None                     # подбор или проверка по требованиям — в аналитику спроса (kbagent/demand.py)

    def cell(model: str, param: str) -> tuple[str, list[Value]]:
        items, state, d = pick(model, param, by, contradicted, decided)
        if state == "missing":
            return "нет данных", []
        if state == "contradiction":
            experts.append((model, param, items))
            return "; ".join(f"{v.shown()} {ref(v)}" for v in items), items
        if state == "noted":
            confirmed.append(f"{model}, {param.lower()}: решение эксперта — {d.get('note') or 'оба значения верны'}")
            return "; ".join(f"{v.shown()} {ref(v)}" for v in items), items
        if state == "confirmed":
            confirmed.append(f"{model}, {param.lower()}: {items[0].shown()} — {confirmed_note(d)}"
                             + (f" ({d['note']})" if d.get("note") else ""))
            return f"{items[0].shown()} {ref(items[0])} ({confirmed_note(d)})", items
        return f"{items[0].shown()} {ref(items[0])}", items

    if q.reqs:
        # сначала решение по каждой модели (без ссылок), ссылки — только у показанных значений:
        # при подборе это подходящие модели, при проверке названной модели — она сама
        state: dict[str, str] = {}
        for m in models:
            ok, known = True, True
            for r in q.reqs:
                items, st, _ = pick(m, r.param, by, contradicted, decided)
                if st not in ("ok", "confirmed"):
                    known = False
                    continue
                ok = ok and _fits(items[0], r)
            state[m] = "fit" if ok and known else "unfit" if not ok else "unknown"
        shown_models = q.models or [m for m in models if state[m] == "fit"][:12]

        def verdicts(m: str) -> list[str]:
            out = []
            for r in q.reqs:
                items, st, _ = pick(m, r.param, by, contradicted, decided)
                text, _ = cell(m, r.param)
                rows.append([m, r.param, re.sub(r"\s*\[\d+\]", "", text)])
                if st not in ("ok", "confirmed"):
                    out.append(f"{r.param.lower()}: {text}")
                else:
                    out.append(f"{r.param.lower()} — {text}, нужно {REQ_WORDS[r.mode]} {r.shown}".replace("  ", " "))
            return out

        fit = [(m, verdicts(m)) if m in shown_models else (m, []) for m in models if state[m] == "fit"]
        unfit = [(m, verdicts(m)) if m in shown_models else (m, []) for m in models if state[m] == "unfit"]
        unknown = [(m, verdicts(m)) if m in shown_models else (m, []) for m in models if state[m] == "unknown"]
        need = "; ".join(f"{r.param.lower()} {REQ_WORDS[r.mode]} {r.shown}".replace("  ", " ") for r in q.reqs)
        demand = {"kind": q.kind or next((v.kind for v in values if v.model in models), ""),
                  "named": list(q.models or []), "fit": [m for m, _ in fit], "unfit": [m for m, _ in unfit],
                  "unknown": [m for m, _ in unknown],
                  "reqs": [f"{r.param.lower()} {REQ_WORDS[r.mode]} {r.shown}".replace("  ", " ") for r in q.reqs]}
        if q.models:
            for m, verdicts in fit + unfit + unknown:
                state = ("подходит" if (m, verdicts) in fit else "не подходит" if (m, verdicts) in unfit
                         else "нужна проверка эксперта")
                lines.append(f"- **{m}** — {state}: " + "; ".join(verdicts))
            head = f"Проверка по справочнику параметров (требование: {need}):"
        else:
            head = (f"Подбор по справочнику параметров — {q.kind.lower() if q.kind else 'продукция'}, требование: {need}. "
                    f"Подходят: {len(fit)} из {len(models)}.")
            lines += [f"- **{m}**: " + "; ".join(v) for m, v in fit[:12]]
            if len(fit) > 12:
                lines.append(f"- и ещё моделей: {len(fit) - 12} — в таблице ниже")
            if not fit:
                lines.append("- Ни одна модель в справочнике не подходит.")
            if unfit:
                lines.append("\nНе подходят: " + ", ".join(m for m, _ in unfit[:20])
                             + (f" и ещё {len(unfit) - 20}" if len(unfit) > 20 else "") + ".")
            if unknown:
                lines.append("Нет данных или источники расходятся: " + ", ".join(m for m, _ in unknown) + ".")
        status = "answer"
    else:
        head = "По справочнику параметров:"
        for m in models:
            parts = []
            for p in params:
                text, items = cell(m, p)
                if items:
                    parts.append(f"{p.lower()} — {text}")
                    rows.append([m, p, re.sub(r"\s*\[\d+\]", "", text)])
            if parts:
                lines.append(f"- **{m}**: " + "; ".join(parts))
        if not lines:
            asked = ", ".join(p.lower() for p in params) or "параметры"
            return SprAnswer("no_data", f"В справочнике параметров нет сведений ({asked}) для: {', '.join(models[:6])}.",
                             stats={"spravochnik": True})
        status = "answer"
    if not q.reqs and len(models) >= 2 and q.all_params:
        # сравнение моделей: какие параметры различаются, какие совпадают
        same, diff = [], []
        for p in params:
            firsts = [(pick(m, p, by, contradicted, decided)[0] or [None])[0] for m in models]
            shown = {v.shown() if v else "нет данных" for v in firsts}
            (same if len(shown) == 1 and "нет данных" not in shown else diff).append(p.lower())
        lines.append("\nРазличаются: " + (", ".join(diff) or "—") + ". Совпадают: " + (", ".join(same) or "—") + ".")
    expert_ids = []
    if confirmed:
        lines.append("\n**Решения эксперта:** " + "; ".join(dict.fromkeys(confirmed)) + ".")
    if experts:
        lines.append("\n**Источники расходятся** — вывод не делаю, вопрос передан эксперту:")
        unique = {(m, p): items for m, p, items in experts}
        for (model, param), items in unique.items():
            lines.append(f"- {model}, {param.lower()}: " + "; ".join(f"{v.shown()} (с. {v.page}) {ref(v)}" for v in items))
            from .experts import add as to_expert
            expert_ids.append(to_expert("противоречие в справочнике", f"{model} — {param}",
                                        [{"value": v.shown(), "doc": v.doc, "page": v.page} for v in items],
                                        question, user, role)["id"])
    note = ("Значения — из справочника параметров, сравнение с требованием выполнено программой."
            if q.reqs else "Значения — из справочника параметров, у каждого — ссылка на страницу руководства.")
    text = head + "\n" + "\n".join(lines) + f"\n\n{note} Это черновик — проверьте его перед отправкой клиенту."
    return SprAnswer(status, text, sources, ["model", "param", "value"], rows, ["Модель", "Параметр", "Значение"],
                     stats={"spravochnik": True, "models": len(models), "contradictions": len(experts),
                            "expert_ids": expert_ids, "demand": demand})
