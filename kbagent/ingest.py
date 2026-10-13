"""Разбор PDF и DOCX на фрагменты с сохранением номера пункта и страницы.

Документ читается построчно. Строка, начинающаяся с номера пункта («4.2»,
«4.2.1», «12.»), открывает новый пункт; короткая строка с номером раздела
(«4 Порядок проведения») или римским номером («II. Требования…») становится
заголовком. Пункты затем собираются во фрагменты удобного для поиска размера:
мелкие соседние пункты объединяются, длинные режутся с перекрытием, но номер
пункта и страницы всегда остаются при фрагменте — они и есть ссылка в ответе.
"""
import re
import time
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

import pymupdf
import yaml
from docx import Document as DocxDocument

from .config import settings
from .text import join_hyphenated, normalize

CLAUSE_RE = re.compile(r"^(?:п\.\s*)?([1-9]\d{0,2}(?:\.\d{1,3}){0,4})(\.?)\s+(?=[А-ЯЁA-Z«\"(])")
ROMAN_RE = re.compile(r"^([IVXL]{1,6})\.\s+[А-ЯЁA-Z]")
ANNEX_RE = re.compile(r"^(Приложение\s+(?:N|№)?\s*[\dА-Я]+)", re.IGNORECASE)

TARGET, MAX_LEN, OVERLAP = 900, 1400, 150


@dataclass
class Chunk:
    id: str
    doc_id: str
    doc_title: str
    doc_kind: str
    access: str
    clause: str | None
    heading: str | None
    page_start: int | None
    page_end: int | None
    text: str

    @property
    def search_text(self) -> str:
        """Текст для индексации: заголовок документа и раздела помогают поиску."""
        head = f"{self.doc_title}. {self.heading or ''}."
        return f"{head}\n{self.text}"

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class _Unit:
    clause: str | None
    heading: str | None
    page_start: int | None
    page_end: int | None
    lines: list
    pages: list = field(default_factory=list)   # страница каждой строки

    @property
    def text(self) -> str:
        return " ".join(self.lines).strip()

    def spans(self) -> list[tuple[int, int | None]]:
        """Начало каждой строки в text и её страница."""
        out, pos = [], 0
        for line, page in zip(self.lines, self.pages):
            out.append((pos, page))
            pos += len(line) + 1
        return out


# ---------- чтение файлов ----------

def ocr_available() -> bool:
    return (settings.tessdata_dir / "rus.traineddata").exists()


def _page_text(page, stats: dict) -> str:
    """Текстовый слой страницы, а для скана — распознанный текст (Tesseract встроен в PyMuPDF,
    языки — tools/tessdata, скачиваются scripts/setup.py --ocr)."""
    # естественный порядок блоков: при sort=True колонки двухколоночной вёрстки перемешиваются
    text = page.get_text("text")
    if len(text.strip()) >= 20 or not page.get_images():
        return text
    if not ocr_available():
        stats["scanned_pages"] = stats.get("scanned_pages", 0) + 1
        return text
    tp = page.get_textpage_ocr(language="rus+eng", dpi=300, full=True, tessdata=str(settings.tessdata_dir))
    stats["ocr_pages"] = stats.get("ocr_pages", 0) + 1
    return page.get_text("text", textpage=tp)


def _pdf_pages(path: Path, stats: dict | None = None) -> list[tuple[int, list[str]]]:
    stats = {} if stats is None else stats
    doc = pymupdf.open(path)
    pages = []
    for i, page in enumerate(doc):
        raw = normalize(_page_text(page, stats))
        lines = [l.strip() for l in raw.split("\n")]
        pages.append((i + 1, [l for l in lines if l]))
    # повторяющиеся колонтитулы и номера страниц
    if len(pages) >= 4:
        freq = Counter(l for _, lines in pages for l in set(lines) if len(l) < 80)
        boiler = {l for l, c in freq.items() if c >= max(3, 0.3 * len(pages))}
    else:
        boiler = set()
    cleaned = []
    for no, lines in pages:
        # номер страницы — и «Страница 1 из 2» в коротком документе (иначе он прилипает к последнему пункту)
        lines = [l for l in lines if l not in boiler and not re.fullmatch(
            r"\d{1,4}|-\s*\d{1,4}\s*-|(?:Страница|Стр\.|Лист)\s*\d{1,4}(?:\s*(?:из|/)\s*\d{1,4})?", l, re.I)]
        cleaned.append((no, join_hyphenated(lines)))
    return cleaned


def _docx_pages(path: Path) -> list[tuple[int | None, list[str]]]:
    doc = DocxDocument(path)
    lines = []
    for p in doc.paragraphs:
        text = normalize(p.text).strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [normalize(c.text).strip() for c in row.cells]
            row_text = " | ".join(dict.fromkeys(c for c in cells if c))
            if row_text:
                lines.append(row_text)
    return [(None, lines)]


# ---------- структура: пункты и заголовки ----------

def _is_heading(line: str, number: str | None, dotted: bool) -> bool:
    if ROMAN_RE.match(line) or ANNEX_RE.match(line):
        return True
    if number and "." not in number:
        # «4 Порядок проведения» — заголовок раздела ГОСТа: номер без точки и короткая строка.
        # «99. Перед хранением…» — пункт приказа, даже если строка оборвана переносом.
        return not dotted and len(line) < 90 and not line.rstrip().endswith((".", ";", ":"))
    letters = [c for c in line if c.isalpha()]
    return 8 < len(line) < 90 and len(letters) > 5 and all(c.isupper() for c in letters)


def _split_inline_clauses(lines: list[str], last_top):
    """В PDF приказов пункт может начинаться в середине строки:
    «…ремонт инструмента. 117. Работать с …». Строка разрезается, если внутри неё
    стоит номер следующего по порядку пункта. Генератор: номер последнего пункта
    обновляется по мере разбора, поэтому проверка идёт по актуальному значению."""
    for line in lines:
        while True:
            n = last_top() + 1
            m = re.search(rf"(?<=[.;:)»\"])\s+(?:{n}|{n + 1}|{n + 2})\.\s+(?=[А-ЯЁ])", line)
            if not m or m.start() < 3:
                yield line
                break
            yield line[:m.start()].strip()
            line = line[m.start():].strip()


def _units(pages) -> list[_Unit]:
    units: list[_Unit] = []
    heading = None
    last_top = 0
    current = _Unit(None, None, None, None, [])

    def flush():
        nonlocal current
        if current.lines:
            units.append(current)
        current = _Unit(current.clause, heading, None, None, [])

    for page_no, lines in pages:
        for line in _split_inline_clauses(lines, lambda: last_top):
            m = CLAUSE_RE.match(line)
            number = m.group(1) if m else None
            dotted = bool(m and m.group(2))
            if number and "." not in number:
                # одиночные номера принимаем только по порядку (допуская пропуск
                # до двух нераспознанных), чтобы «220 В» в тексте не стал пунктом 220
                value = int(number)
                if not (value == 1 or last_top < value <= last_top + 3):
                    number = None
            if _is_heading(line, number, dotted):
                flush()
                heading = line[:120]
                if number and "." not in number:
                    last_top = int(number)
                current = _Unit(number, heading, page_no, page_no, [])
                continue
            if number:
                flush()
                if "." not in number:
                    last_top = int(number)
                current = _Unit(number, heading, page_no, page_no, [line], [page_no])
                continue
            if current.page_start is None:
                current.page_start = page_no
                current.heading = heading
            current.page_end = page_no
            current.lines.append(line)
            current.pages.append(page_no)
    flush()
    return [u for u in units if u.text]


def _split_long(text: str) -> list[tuple[int, str]]:
    """Куски длинного текста с перекрытием и начало каждого куска в тексте."""
    parts, start = [], 0
    while start < len(text):
        end = min(start + TARGET, len(text))
        if end < len(text):
            cut = text.rfind(". ", start + TARGET // 2, end)
            end = cut + 1 if cut > 0 else end
        raw = text[start:end]
        parts.append((start + len(raw) - len(raw.lstrip()), raw.strip()))
        if end >= len(text):
            break
        # перекрытие начинается с начала предложения (или хотя бы слова), а не с середины слова:
        # иначе фрагмент начинается с «ники из числа…» и поиск его не узнаёт
        back = max(end - OVERLAP, start + 1)
        sentence = text.find(". ", back, end)
        if sentence != -1:
            start = sentence + 2
        else:
            space = text.find(" ", back, end)
            start = space + 1 if space != -1 else back
    return parts


def _chunks(units: list[_Unit], meta: dict) -> list[Chunk]:
    chunks: list[Chunk] = []
    buf: list[_Unit] = []

    def emit(group: list[_Unit]):
        text = "\n".join(u.text for u in group)
        clauses = [u.clause for u in group if u.clause]
        if not clauses:
            clause = None
        elif clauses[0] == clauses[-1]:
            clause = clauses[0]
        else:
            clause = f"{clauses[0]}–{clauses[-1]}"
        pages = [p for u in group for p in (u.page_start, u.page_end) if p]
        base = dict(doc_id=meta["id"], doc_title=meta["title"], doc_kind=meta["kind"], access=meta["access"],
                    clause=clause, heading=group[0].heading,
                    page_start=min(pages) if pages else None, page_end=max(pages) if pages else None)
        # страница каждой строки по её началу в общем тексте: у куска длинного пункта —
        # свои страницы, а не весь пункт («стр. 125», а не «стр. 124–126»)
        line_pages, offset = [], 0
        for u in group:
            line_pages += [(offset + pos, page) for pos, page in u.spans() if page]
            offset += len(u.text) + 1
        pieces = _split_long(text) if len(text) > MAX_LEN else [(0, text)]
        opening = re.sub(r"^\s*\d{1,3}(?:\.\d{1,3})*\.?\s+", "", pieces[0][1])     # без номера пункта
        lead = re.split(r"(?<=[.;:])\s", opening, maxsplit=1)[0][:200]
        for k, (found, piece) in enumerate(pieces):
            # продолжение длинного пункта получает в заголовок его начало: иначе фрагмент «из числа
            # административно-технического персонала…» не говорит, о чём он, и поиск его не узнаёт
            heading = base["heading"] if k == 0 else f"{base['heading'] or ''} — продолжение: {lead}".strip(" —")
            own = {**base, "heading": heading}
            if len(pieces) > 1:
                # строка, в которой кусок начинается, и все строки, начатые внутри него
                first = max((pos for pos, _ in line_pages if pos <= found), default=0)
                span = [p for pos, p in line_pages if first <= pos < found + len(piece)]
                if span:
                    own.update(page_start=min(span), page_end=max(span))
            chunks.append(Chunk(id=f"{meta['id']}#{len(chunks)}", text=piece, **own))

    # Пункты одного раздела объединяются, пока фрагмент не превысит целевой размер:
    # короткий раздел («2 Классификация отказов» из четырёх пунктов) остаётся целым.
    for unit in units:
        same_section = buf and buf[-1].heading == unit.heading
        size = sum(len(u.text) for u in buf)
        if buf and (not same_section or size + len(unit.text) > TARGET):
            emit(buf)
            buf = []
        buf.append(unit)
    if buf:
        emit(buf)
    return chunks


# ---------- каталог документов ----------

ARCHIVE = "archive"   # группа доступа отменённых редакций: её нет ни у одной роли
DRAFT = "admin"       # проект новой редакции до утверждения видит только администратор


def is_cancelled(info: dict) -> bool:
    """Статус в каталоге: «действует» (по умолчанию) или «отменён» (с полем replaced_by)."""
    return str(info.get("status", "действует")).lower().startswith("отмен")


def load_catalog(raw_dir: Path) -> dict[str, dict]:
    """Названия, группы доступа и статусы (действует / отменён) из data/sources.yaml и data/raw/*/catalog.yaml."""
    catalog = {}
    for path in [raw_dir.parent / "sources.yaml", *raw_dir.glob("*/catalog.yaml")]:
        if path.exists():
            for d in yaml.safe_load(path.read_text(encoding="utf-8"))["documents"]:
                catalog[d["id"]] = d
    # новая редакция, загруженная с пометкой «заменяет X», отменяет X — только после утверждения:
    # черновик с ошибкой не должен подменять действующий регламент для всех сотрудников
    for d in list(catalog.values()):
        old = catalog.get(d.get("replaces"))
        if old is not None and not is_cancelled(d) and not is_draft_edition(d):
            old.update(status="отменён", replaced_by=d["id"])
    return catalog


def _date(value) -> str:
    """«2026-03-01» из каталога (строка или дата YAML); пусто — не задано."""
    return str(value or "")[:10]


def not_yet_effective(info: dict) -> str:
    """Дата вступления в силу, если она ещё не наступила («вступает в силу 01.12.2026»), иначе пусто."""
    start = _date(info.get("effective_from"))
    return start if start and start > time.strftime("%Y-%m-%d") else ""


def expired_on(info: dict) -> str:
    """Дата окончания действия, если она уже прошла («действует до 31 декабря 2025 года»), иначе пусто."""
    end = _date(info.get("effective_to"))
    return end if end and end < time.strftime("%Y-%m-%d") else ""


def _ru(day: str) -> str:
    return f"{day[8:10]}.{day[5:7]}.{day[:4]}" if len(day) == 10 else day


_REGISTRY: dict = {}


def registry_info(doc_id: str, raw_dir: Path | None = None) -> dict:
    """Запись реестра (каталога) о документе; каталог перечитывается, только если файлы изменились."""
    from .config import settings
    raw_dir = raw_dir or settings.raw_dir
    files = [raw_dir.parent / "sources.yaml", *raw_dir.glob("*/catalog.yaml")]
    stamp = tuple((str(f), f.stat().st_mtime) for f in files if f.exists())
    if _REGISTRY.get("stamp") != stamp:
        _REGISTRY.update(stamp=stamp, catalog=load_catalog(raw_dir))
    return _REGISTRY["catalog"].get(doc_id, {})


def registry_suffix(doc_id: str, raw_dir: Path | None = None) -> str:
    """Версия и срок действия документа для подписи ссылки: «ред. 2026, действует с 01.03.2026»;
    у истёкшего — «срок действия истёк 31.12.2025». Каталог перечитывается, только если файлы изменились."""
    info = registry_info(doc_id, raw_dir)
    parts = []
    if info.get("version"):
        parts.append(f"ред. {info['version']}")
    if expired_on(info):
        parts.append(f"срок действия истёк {_ru(expired_on(info))}")
    elif info.get("effective_from"):
        parts.append(f"действует с {_ru(_date(info['effective_from']))}")
    return ", ".join(parts)


def is_draft_edition(info: dict) -> bool:
    """Загруженная новая редакция, которую ещё не утвердили (у документов из sources.yaml поля approved нет)."""
    return bool(info.get("replaces")) and not info.get("approved", True)


def access_of(info: dict) -> str:
    if is_cancelled(info):
        return ARCHIVE
    # проект новой редакции и документ, который ещё не вступил в силу, видит только администратор
    return DRAFT if is_draft_edition(info) or not_yet_effective(info) else info.get("access", "all")


def parse_document(path: Path, catalog: dict, stats: dict | None = None) -> list[Chunk]:
    """stats (необязательно) получает счётчики: ocr_pages — распознано страниц-сканов,
    scanned_pages — сканы, которые нечем распознать."""
    info = catalog.get(path.stem, {})
    meta = {
        "id": path.stem,
        "title": info.get("title", path.stem),
        "kind": info.get("kind", "Документ"),
        # отменённая редакция остаётся в индексе для истории, но не доступна ни одной роли
        "access": access_of(info),
    }
    pages = _pdf_pages(path, stats) if path.suffix.lower() == ".pdf" else _docx_pages(path)
    return _chunks(_units(pages), meta)


def parse_corpus(raw_dir: Path) -> list[Chunk]:
    catalog = load_catalog(raw_dir)
    chunks = []
    for path in sorted(raw_dir.rglob("*")):
        if path.suffix.lower() in (".pdf", ".docx") and not path.name.startswith("~$"):
            chunks.extend(parse_document(path, catalog))
    return chunks
