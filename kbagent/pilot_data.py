"""Данные партнёра для пилота: шаблон Excel и его загрузка в агента.

Партнёр заполняет одну книгу Excel (её собирает scripts/make_partner_package.py), команда загружает её
scripts/import_partner.py. Каждая строка проверяется; ошибки — с листом и номером строки Excel, чтобы их можно
было отправить партнёру на исправление.

    «Реестр документов»       → catalog.yaml рядом с документами: название, версия, срок действия, владелец,
                                 доступ, статус редакции — как у демо-регламентов
    «Вопросы»                 → набор в формате эталона (eval/*.yaml) для eval/run_eval.py
    «Справочник параметров»   → CSV для вкладки «Справочник» — с той же проверкой, что при загрузке
    «Прайс»                   → CSV для «Прайс для ТКП» — тоже с проверкой

Закрытый приёмочный набор заполняется в том же формате (лист «Вопросы»), но остаётся у эксперта партнёра до
приёмки — разработчики его не видят (карточка задачи, п. 2; протокол замера docs/pilot/zamer_i_stoimost.md).
"""
import csv
import datetime as dt
import io
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from .config import settings

EXAMPLE = "пример"          # строка-образец в шаблоне: в столбце «№», не загружается
FILE_TYPES = (".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".tif", ".tiff")


@dataclass
class Col:
    name: str
    width: int
    required: bool = False
    hint: str = ""
    choices: list = field(default_factory=list)


ACCESS = {"все сотрудники": "all", "производство": "production", "метрология": "metrology",
          "только администратор": "admin"}
STATUS = {"действует": "", "отменён": "отменён", "отменен": "отменён"}
WHERE = {"в документе": "doc", "нет в документах": "none", "в учётной системе": "db",
         "документ и учётная система": "both"}

SHEETS = {
    "Реестр документов": [
        Col("№", 6),
        Col("Файл", 34, True, "имя файла, как он передан: Regl_TKP_2025.pdf"),
        Col("Название", 46, True, "как в самом документе: СТП 12-2025 «Порядок подготовки ТКП»"),
        Col("Вид", 18, False, "регламент, инструкция, ТУ, руководство по эксплуатации…"),
        Col("Редакция", 12, False, "номер или год редакции"),
        Col("Действует с", 13, False, "ДД.ММ.ГГГГ"),
        Col("Действует до", 13, False, "ДД.ММ.ГГГГ, если срок ограничен"),
        Col("Владелец", 22, False, "кто отвечает за актуальность: должность"),
        Col("Доступ", 20, True, "кто видит документ", list(ACCESS)),
        Col("Приоритет", 11, False, "1 — главный источник при расхождении, 2 — следующий…"),
        Col("Статус", 12, False, "отменённая редакция остаётся для истории, но в ответах не участвует",
            ["действует", "отменён"]),
        Col("Заменяет", 30, False, "файл прежней редакции, если документ её заменяет"),
    ],
    "Вопросы": [
        Col("№", 6),
        Col("Вопрос", 50, True, "как спросил сотрудник или клиент — своими словами"),
        Col("Кто спрашивает", 22, False, "роль в системе", list(settings.roles)),
        Col("Где ответ", 24, True, "", list(WHERE)),
        Col("Правильный ответ", 44, False, "кратко, для эксперта"),
        Col("Ключевые факты", 32, False, "что обязательно должно быть в ответе, через «;»: 5 рабочих дней; 2211"),
        Col("Чего быть не должно", 24, False, "через «;» — например, телефон из отменённой редакции"),
        Col("Файл документа", 30, False, "где ответ — для «в документе»"),
        Col("Пункт", 10, False, "номер пункта: 4.2"),
    ],
    "Справочник параметров": [
        Col("№", 6),
        Col("Модель", 22, True), Col("Вид", 18), Col("Параметр", 30, True),
        Col("Значение", 14, True, "число, диапазон 50–350 или текст"), Col("Единица", 10, False, "МГц, А, В, °C…"),
        Col("Условие", 16, False, "не более, на канал…"),
        Col("Документ", 30, True, "файл из реестра, откуда значение"),
        Col("Страница", 10, True), Col("Цитата", 40, False, "фраза со страницы — по ней значение подсвечивается"),
    ],
    "Прайс": [
        Col("№", 6),
        Col("Модель", 22, True), Col("Цена", 12, True, "₽ за единицу"),
        Col("Срок поставки", 18, False, "в наличии, 10 рабочих дней…"),
        Col("Действует с", 13, False, "дата прайса, ДД.ММ.ГГГГ"), Col("Примечание", 24),
    ],
}
EXAMPLES = {
    "Реестр документов": ["Regl_TKP_2025.pdf", "СТП 12-2025 «Порядок подготовки ТКП»", "регламент", "2025",
                          "01.09.2025", "", "Начальник отдела продаж", "все сотрудники", 1, "действует", ""],
    "Вопросы": ["Какой срок ответа клиенту на запрос ТКП?", "Техпродажи и сервис", "в документе",
                "Не позднее 2 рабочих дней (п. 4.3)", "2 рабочих дн", "", "Regl_TKP_2025.pdf", "4.3"],
    "Справочник параметров": ["АКИП-4115/2В", "Осциллограф", "Полоса пропускания", "200", "МГц", "",
                              "osc_akip_4115v.pdf", 6, "Полоса пропускания 100МГц 200МГц"],
    "Прайс": ["АКИП-4115/2В", "46500", "в наличии", "01.10.2026", ""],
}


@dataclass
class Result:
    documents: list = field(default_factory=list)      # для catalog.yaml
    questions: list = field(default_factory=list)      # для eval/*.yaml
    spravochnik_csv: str = ""
    prajs_csv: str = ""
    spravochnik_models: set = field(default_factory=set)
    errors: list = field(default_factory=list)         # {sheet, row, error}
    notes: list = field(default_factory=list)          # не ошибки, но что сделать вручную
    counts: dict = field(default_factory=dict)         # лист → загружено строк


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def _rows(wb, sheet: str, res: Result) -> list[tuple[int, dict]]:
    """Заполненные строки листа: (номер строки Excel, {столбец: значение}); образец и пустые — пропускаются."""
    if sheet not in wb.sheetnames:
        return []
    data = list(wb[sheet].iter_rows(values_only=True))
    if not data:
        return []
    head = [_cell(h) for h in data[0]]
    missing = [c.name for c in SHEETS[sheet] if c.required and c.name not in head]
    if missing:
        res.errors.append({"sheet": sheet, "row": 1, "error": f"нет столбцов: {', '.join(missing)}"})
        return []
    out = []
    for n, raw in enumerate(data[1:], start=2):
        row = {h: _cell(v) for h, v in zip(head, raw) if h}
        if row.get("№", "").lower() == EXAMPLE or not any(v for k, v in row.items() if k != "№"):
            continue
        empty = [c.name for c in SHEETS[sheet] if c.required and not row.get(c.name)]
        if empty:
            res.errors.append({"sheet": sheet, "row": n, "error": f"не заполнено: {', '.join(empty)}"})
            continue
        out.append((n, row))
    return out


def _iso(day: str) -> str | None:
    try:
        return dt.datetime.strptime(day, "%d.%m.%Y").date().isoformat()
    except ValueError:
        return None


def _doc_id(name: str) -> str:
    return Path(name.strip()).stem


def registry(wb, files_dir: Path | None, res: Result) -> None:
    """Лист «Реестр документов» → записи каталога; files_dir — папка, куда положены присланные файлы."""
    sheet = "Реестр документов"
    rows = _rows(wb, sheet, res)
    present = {p.name.lower() for p in files_dir.iterdir()} if files_dir and files_dir.is_dir() else None
    ids = {_doc_id(r["Файл"]) for _, r in rows}
    seen = set()
    for n, r in rows:
        problem = ""
        access = ACCESS.get(r["Доступ"].lower())
        status = STATUS.get(r.get("Статус", "").lower() or "действует")
        bad_dates = [c for c in ("Действует с", "Действует до") if r.get(c) and _iso(r[c]) is None]
        if Path(r["Файл"]).suffix.lower() not in FILE_TYPES:
            problem = f"«{r['Файл']}» — нужен файл PDF, Word, Excel или изображение"
        elif present is not None and r["Файл"].lower() not in present:
            problem = f"файла «{r['Файл']}» нет в папке с документами"
        elif _doc_id(r["Файл"]) in seen:
            problem = f"файл «{r['Файл']}» указан дважды"
        elif access is None:
            problem = f"доступ «{r['Доступ']}» — выберите из списка: {', '.join(ACCESS)}"
        elif status is None:
            problem = f"статус «{r['Статус']}» — «действует» или «отменён»"
        elif bad_dates:
            problem = "; ".join(f"«{c}»: «{r[c]}» — не дата ДД.ММ.ГГГГ" for c in bad_dates)
        elif r.get("Приоритет") and not r["Приоритет"].isdigit():
            problem = f"приоритет «{r['Приоритет']}» — не число"
        elif r.get("Заменяет") and _doc_id(r["Заменяет"]) not in ids:
            problem = f"«Заменяет»: файла «{r['Заменяет']}» нет в реестре"
        if problem:
            res.errors.append({"sheet": sheet, "row": n, "error": problem})
            continue
        doc = {"id": _doc_id(r["Файл"]), "title": r["Название"], "kind": r.get("Вид") or "Документ",
               "access": access}
        optional = {"version": r.get("Редакция"), "effective_from": _iso(r.get("Действует с", "")),
                    "effective_to": _iso(r.get("Действует до", "")), "owner": r.get("Владелец"),
                    "priority": int(r["Приоритет"]) if r.get("Приоритет") else None, "status": status,
                    "replaces": _doc_id(r["Заменяет"]) if r.get("Заменяет") else None}
        doc.update({k: v for k, v in optional.items() if v})
        seen.add(doc["id"])
        res.documents.append(doc)
        if Path(r["Файл"]).suffix.lower() not in (".pdf", ".docx"):
            res.notes.append(f"«{r['Файл']}» — Excel и изображения загружаются через «Администрирование» → "
                             "«Документы»: сборка индекса читает из папки только PDF и DOCX")
    # новая редакция «заменяет» прежнюю → у прежней статус «отменён» и ссылка на новую, как в демо-каталоге
    by_id = {d["id"]: d for d in res.documents}
    for d in res.documents:
        old = by_id.get(d.pop("replaces", None) or "")
        if old is not None:
            old.update(status="отменён", replaced_by=d["id"])
    res.counts[sheet] = len(res.documents)


def _clause(text: str) -> str:
    """«п. 4.2» → «4.2»; страница («с. 12») — не пункт."""
    m = re.search(r"\d+(?:\.\d+)*", text)
    return m.group(0) if m and not re.match(r"\s*(?:с\.|стр)", text, re.I) else ""


def questions(wb, known_docs: set[str], res: Result, prefix: str = "P") -> None:
    """Лист «Вопросы» → вопросы в формате эталона."""
    sheet = "Вопросы"
    for n, r in _rows(wb, sheet, res):
        kind = WHERE.get(r["Где ответ"].lower())
        role = r.get("Кто спрашивает", "")
        facts = [f.strip() for f in r.get("Ключевые факты", "").split(";") if f.strip()]
        banned = [f.strip() for f in r.get("Чего быть не должно", "").split(";") if f.strip()]
        doc = _doc_id(r["Файл документа"]) if r.get("Файл документа") else ""
        problem = ""
        if kind is None:
            problem = f"«Где ответ»: «{r['Где ответ']}» — выберите из списка: {', '.join(WHERE)}"
        elif role and role not in settings.roles:
            problem = f"роль «{role}» — выберите из списка: {', '.join(settings.roles)}"
        elif kind != "none" and not facts:
            problem = "для вопроса с ответом нужны ключевые факты — по ним проверяется ответ"
        elif kind in ("doc", "both") and not doc:
            problem = "укажите файл документа, где ответ"
        elif kind in ("doc", "both") and doc not in known_docs:
            problem = f"документа «{r['Файл документа']}» нет в реестре"
        if problem:
            res.errors.append({"sheet": sheet, "row": n, "error": problem})
            continue
        q = {"id": f"{prefix}{len(res.questions) + 1:02d}", "group": "партнёр", "type": kind, "question": r["Вопрос"]}
        if role:
            q["role"] = role
        if kind != "none":
            q["must_include"] = facts
        if banned:
            q["must_exclude"] = banned
        if kind in ("doc", "both"):
            q["source"] = {"doc": doc}
            if _clause(r.get("Пункт", "")):
                q["source"]["clause"] = _clause(r["Пункт"])
        if r.get("Правильный ответ"):
            q["expected"] = r["Правильный ответ"]       # для эксперта; прогон его не проверяет
        res.questions.append(q)
    res.counts[sheet] = len(res.questions)


def _table(wb, sheet: str, columns: list[str], res: Result, check, doc_column: str | None = None) -> str:
    """Лист → CSV «;» с колонками агента (имя файла документа → его id) и проверка тем же разбором, что при
    загрузке через администрирование; ошибки — с номером строки Excel."""
    rows = _rows(wb, sheet, res)
    res.counts[sheet] = len(rows)
    if not rows:
        return ""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(columns)
    for _, r in rows:
        if doc_column:
            r[doc_column] = _doc_id(r[doc_column])
        w.writerow([r.get(c, "") for c in columns])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "check.csv"
        path.write_text(buf.getvalue(), encoding="utf-8")
        problems = check(path)
    for e in problems:
        line = e.get("line", 0)              # строка CSV: 1 — заголовок, 2 — первая строка данных
        res.errors.append({"sheet": sheet, "row": rows[line - 2][0] if 2 <= line < len(rows) + 2 else "—",
                           "error": e["error"]})
    return buf.getvalue()


def load(path: Path, files_dir: Path | None = None, prefix: str = "P") -> Result:
    """Книга партнёра → реестр, вопросы, справочник и прайс; ошибки строк — в Result.errors."""
    from openpyxl import load_workbook

    from . import prajs
    from . import spravochnik as spr
    from .ingest import load_catalog
    wb = load_workbook(path, read_only=True, data_only=True)
    res = Result()
    registry(wb, files_dir, res)
    known = set(load_catalog(settings.raw_dir)) | {d["id"] for d in res.documents}
    questions(wb, known, res, prefix)

    def check_spr(csv_path: Path) -> list[dict]:
        values, rep = spr.parse(csv_path, known)
        res.spravochnik_models = {v.model for v in values}
        return rep.errors

    def check_prices(csv_path: Path) -> list[dict]:
        models = res.spravochnik_models or {v.model for v in spr.load()[0]}
        _, rep = prajs.parse(csv_path, models)
        return (rep.errors
                + [{"error": f"модели «{m}» нет в справочнике параметров — в ТКП она не попадёт"} for m in rep.unknown]
                + [{"error": f"разные цены у {c['model']}: {', '.join(c['prices'])}"} for c in rep.conflicts])

    res.spravochnik_csv = _table(wb, "Справочник параметров", spr.COLUMNS, res, check_spr, doc_column="Документ")
    res.prajs_csv = _table(wb, "Прайс", prajs.COLUMNS, res, check_prices)
    return res
