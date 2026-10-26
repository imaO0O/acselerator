"""Отчёт о пилотной эксплуатации одной кнопкой — по журналам агента за период, в Word.

Для перехода от лабораторного прототипа к проверке в условиях заказчика (TRL 4 → 5), отзыва партнёра и заявки
на грант: сколько сотрудников и вопросов, доля ответов со ссылкой на источник, честные отказы, оценки,
время до утверждённого ответа и ТКП против ручной подготовки (цель партнёра — минус 50 %), пробелы в
документации, работа эксперта, ТКП и спрос, качество на проверочных наборах, безопасность. В конце — место
для заключения и подписей заказчика.

Персональные данные в отчёт не попадают: только числа, роли и формулировки вопросов из журнала (в журнале
телефоны, почта, ФИО уже замаскированы). Прогоны эталона и проверки перед показом (пользователь eval) не считаются.
"""
import io
import json
import re
import statistics
import time
from collections import Counter
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt

from . import auth, demand, effect, experts, tkp_registry
from .actions import load_actions
from .config import ROOT, settings
from .gaps import SERVICE_USERS, UNANSWERED, find_gaps

EVAL = [("latest", "Эталон: вопросы по документам и учёту"), ("jury", "Вопросы «с ходу», как от жюри"),
        ("roles", "Разграничение доступа по ролям"), ("spravochnik", "Справочник параметров"),
        ("tkp", "Письма клиентов → ТКП")]
ROUTES = {"documents": "по документам", "database": "по учётной базе", "both": "по документам и учёту",
          "spravochnik": "по справочнику параметров", "tkp": "разбор письма клиента", "normo": "нормоконтроль документа"}


def _in(stamp: str, since: date, until: date) -> bool:
    return since.isoformat() <= (stamp or "")[:10] <= until.isoformat()


def _audit() -> list[dict]:
    path = settings.log_dir / "audit.jsonl"
    out = []
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except OSError:
        return []
    for line in lines:
        try:
            out.append(json.loads(line))
        except ValueError:
            continue
    return out


def eval_counts(md_path) -> tuple[int, int] | None:
    """Верно и всего по прогону: из CSV рядом с отчётом (в Markdown «|» из письма клиента ломал бы подсчёт)."""
    import csv
    path = md_path.with_suffix(".csv")
    if not path.exists():
        return None
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    return sum(r.get("correct") == "True" for r in rows), len(rows)


def _eval() -> list[tuple[str, str, str]]:
    """(набор, «верно из всего», дата прогона) — из eval/results/*.md."""
    out = []
    for name, title in EVAL:
        runs = []
        for path in (ROOT / "eval" / "results" / f"{name}.md", ROOT / "eval" / "results" / f"check_{name}.md"):
            if not path.exists():
                continue
            text = path.read_text(encoding="utf-8")
            counts = eval_counts(path)
            when = (re.search(r"\((\d{4}-\d{2}-\d{2}[ \d:]*)", text) or [None, ""])[1]
            if counts:
                runs.append((when, f"{counts[0]} из {counts[1]}"))
        if runs:                                       # официальный прогон или проверка ветки — что новее
            when, ok = max(runs)
            out.append((title, ok, when[:10]))
    return out


def collect(since: date, until: date) -> dict:
    """Все цифры отчёта за период [since, until] включительно."""
    records = _audit()
    people = [r for r in records if r.get("user") not in SERVICE_USERS and r.get("route") != "help"
              and _in(r.get("time"), since, until)]
    answered = [r for r in people if r.get("status") == "answer"]
    roles = {}
    for r in people:
        k = roles.setdefault(r.get("role") or "—", [0, 0])
        k[0] += 1
        k[1] += r.get("status") == "answer"
    feedback = [f for f in auth.load_feedback() if f.get("user") not in SERVICE_USERS and _in(f.get("time"), since, until)]
    eff = [r for r in effect.load() if r.get("user") not in SERVICE_USERS and _in(r.get("time"), since, until)]
    topics, _ = find_gaps([r for r in records if (r.get("time") or "")[:10] <= until.isoformat()])
    topics = [t for t in topics if _in(t["last"], since, until) or _in(t["closed"], since, until)]
    exp = [x for x in experts.load() if _in(x.get("time"), since, until)]
    registry = [r for r in tkp_registry.load() if _in(r.get("date"), since, until)]
    current = [r for r in registry if r.get("status") != tkp_registry.REPLACED]
    actions = [a for a in load_actions() if _in(a.get("time"), since, until)]
    acts = Counter(a.get("action") for a in actions)
    spros = demand.summary([r for r in demand.load(since) if _in(r.get("time"), since, until)], registry, top=5)
    return {
        "since": since, "until": until, "users": len({r.get("user") for r in people}),
        "questions": len(people), "answered": len(answered),
        "cited": sum(1 for r in answered if r.get("cited")),
        "refused": sum(1 for r in people if r.get("status") in UNANSWERED),
        "median_s": statistics.median([r["seconds"] for r in people if r.get("seconds") is not None] or [0]),
        "routes": Counter(r.get("route") for r in answered).most_common(),
        "roles": sorted(roles.items(), key=lambda x: -x[1][0]),
        "up": sum(1 for f in feedback if f.get("score") == "👍"), "down": sum(1 for f in feedback if f.get("score") == "👎"),
        "effect": effect.summary(eff),
        "gaps_open": [t for t in topics if not t["closed"]], "gaps_closed": [t for t in topics if t["closed"]],
        "experts": len(exp), "experts_done": sum(1 for x in exp if x.get("status") == "решён"),
        "experts_values": sum(1 for x in exp if x.get("chosen")),
        "tkp": len(current), "tkp_total": sum(r.get("total") or 0 for r in current),
        "tkp_won": sum(1 for r in current if r.get("status") == "выигран"),
        "tkp_lost": sum(1 for r in current if r.get("status") == "проигран"),
        "spros": spros, "eval": _eval(),
        "logins": acts["login"], "login_failed": acts["login_failed"] + acts["login_locked"],
        "api_denied": acts["api_denied"], "documents": acts["upload"] + acts["approve"],
    }


def _pct(a: int, b: int) -> str:
    return f"{100 * a / b:.0f} %" if b else "—"


def _num(x: float, digits: int = 1) -> str:
    return f"{x:.{digits}f}".replace(".", ",")


def _ru(d: date) -> str:
    return d.strftime("%d.%m.%Y")


def goal_lines(eff: dict) -> list[str]:
    """Цель партнёра — минус 50 % времени на подготовку проверенного ответа: по каждому виду результата."""
    out = []
    for kind, title in (("ответ", "ответ клиенту или сотруднику"), ("ТКП", "ТКП по письму клиента")):
        k = eff["kinds"].get(kind)
        if not k:
            out.append(f"{title}: утверждённых пока нет — цель не оценена")
            continue
        cut = 100 * (1 - k["share"])
        verdict = "цель достигнута" if k["share"] <= 0.5 else "цель пока не достигнута"
        out.append(f"{title}: медиана {_num(k['median_min'])} мин против {k['baseline_min']} мин вручную "
                   f"(−{cut:.0f} %), утверждено {k['count']} — {verdict}")
    return out


def report_docx(data: dict, organization: str = "", app_name: str = "AI-агент") -> bytes:
    doc = Document()
    for section in doc.sections:                     # ГОСТ Р 7.0.97: A4, поля 20/10/20/20 мм, Times New Roman 12
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(20), Mm(10)
        section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    style.paragraph_format.space_after = Pt(0)
    for name in ("List Bullet", "Table Grid"):
        doc.styles[name].font.name = "Times New Roman"

    def para(text: str, bold: bool = False, center: bool = False, size: int | None = None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.paragraph_format.first_line_indent = Cm(1.25)
        return p

    def heading(text: str):
        h = doc.add_paragraph()
        h.add_run(text).bold = True
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12), Pt(6)
        h.paragraph_format.keep_with_next = True

    def table(rows: list[list[str]], head: list[str], widths: list[float]):
        t = doc.add_table(rows=0, cols=len(head), style="Table Grid")
        t.autofit = False
        for i, row in enumerate([head] + rows):
            cells = t.add_row().cells
            for c, value, w in zip(cells, row, widths):
                c.width = Cm(w)
                c.text = str(value)
                if i == 0:
                    c.paragraphs[0].runs[0].bold = True
        doc.add_paragraph()

    d = data
    para("Отчёт о пилотной эксплуатации", bold=True, center=True, size=14)
    para(app_name, center=True)
    para(f"Заказчик: {organization or '____________________'}", center=True)
    para(f"Период: {_ru(d['since'])} — {_ru(d['until'])}", center=True)
    para(f"Сформирован автоматически по журналам агента {time.strftime('%d.%m.%Y %H:%M')}. Персональные данные не "
         "включены; проверочные прогоны не учитываются.", center=True, size=10)

    heading("1. Итоги")
    eff = d["effect"]
    table([["Сотрудников, задававших вопросы", d["users"]],
           ["Вопросов", d["questions"]],
           ["С ответом по документам, учёту или справочнику", f"{d['answered']} ({_pct(d['answered'], d['questions'])})"],
           ["Из них со ссылкой на источник (документ, пункт, страница)", f"{d['cited']} ({_pct(d['cited'], d['answered'])})"],
           ["Честный отказ «в документах ответа нет» вместо выдумки", d["refused"]],
           ["Время ответа, медиана", f"{_num(d['median_s'])} с"],
           ["Оценки сотрудников: полезно / неудачно", f"{d['up']} / {d['down']} ({_pct(d['up'], d['up'] + d['down'])} полезных)"],
           ["Утверждено ответов клиентам и ТКП", sum(k["count"] for k in eff["kinds"].values())],
           ["Сэкономлено (оценка)", f"{_num(eff['saved_hours'])} ч ≈ " + f"{eff['saved_rub']:,} ₽".replace(",", " ")]],
          ["Показатель", "Значение"], [11.5, 6.5])

    heading("2. Цель пилота: минус 50 % времени на подготовку проверенного ответа")
    para("Время от вопроса до утверждения сотрудником агент замеряет сам. Ручное время — оценка до замера "
         f"({effect.BASELINE_MIN['ответ']} мин на ответ, {effect.BASELINE_MIN['ТКП']} мин на ТКП); после замера у "
         "заказчика заменяется измеренным.")
    for line in goal_lines(eff):
        doc.add_paragraph(line, style="List Bullet")

    heading("3. Вопросы по ролям и видам ответа")
    if d["roles"]:
        table([[role, n, f"{ok} ({_pct(ok, n)})"] for role, (n, ok) in d["roles"]],
              ["Роль", "Вопросов", "С ответом"], [9, 4, 5])
    if d["routes"]:
        para("Ответы: " + "; ".join(f"{ROUTES.get(r, r or '—')} — {n}" for r, n in d["routes"]) + ".")

    heading("4. Пробелы в документации")
    para(f"Тем без ответа: {len(d['gaps_open'])}; закрыто за период (после загрузки документов): "
         f"{len(d['gaps_closed'])}. Это список того, что стоит дописать в регламенты или загрузить в базу знаний.")
    if d["gaps_open"]:
        table([[t["question"][:150], t["count"], ", ".join(t["roles"])] for t in d["gaps_open"][:10]],
              ["Вопрос", "Раз", "Роли"], [11, 1.5, 5.5])

    heading("5. Эксперт и справочник параметров")
    para(f"Вопросов эксперту: {d['experts']}, решено: {d['experts_done']}, значений подтверждено экспертом: "
         f"{d['experts_values']} — дальше агент отвечает ими сам.")

    heading("6. ТКП и спрос")
    s = d["spros"]
    para(f"ТКП утверждено: {d['tkp']} на {d['tkp_total']:,.0f} ₽".replace(",", " ")
         + f"; выиграно {d['tkp_won']}, проиграно {d['tkp_lost']}. Позиций в запросах клиентов: {s['positions']}, "
         f"закрыто моделью из справочника: {s['closed_share'] if s['closed_share'] is not None else '—'} %.")
    if s["kinds"]:
        table([[k, n, qty, _pct(closed, n)] for k, n, qty, closed in s["kinds"]],
              ["Вид продукции", "Запросов", "Штук", "Закрыто"], [8, 3.3, 3.3, 3.4])
    if s["unmet"]:
        para("Не закрыто: " + "; ".join(f"{k} — {why} ({n})" for k, why, n in s["unmet"]) + ".")
    if s["lost_reasons"]:
        para("Причины проигрышей: " + "; ".join(f"{r} — {n}" for r, n in s["lost_reasons"]) + ".")

    heading("7. Качество на проверочных наборах")
    if d["eval"]:
        table([[t, ok, _ru(date.fromisoformat(w)) if w else "—"] for t, ok, w in d["eval"]],
              ["Набор вопросов", "Верно", "Прогон"], [10, 4, 4])
    para("Ответ засчитывается, только если верные факты есть в тексте ответа и ссылка указывает на верный "
         "документ и пункт.")

    heading("8. Безопасность")
    para(f"Входов: {d['logins']}; неудачных попыток входа: {d['login_failed']}; запросов к API с неверным ключом: "
         f"{d['api_denied']}; загрузок и утверждений документов: {d['documents']}. Все действия — в журнале "
         "действий (только дописывается). Агент работает без интернета, к учётной базе и 1С — только чтение.")

    heading("9. Заключение заказчика")
    para("Цель пилота (минус 50 % времени): достигнута / не достигнута (нужное подчеркнуть).")
    para("Замечания и предложения:")
    box = doc.add_table(rows=1, cols=1, style="Table Grid")
    box.cell(0, 0).width = Cm(18)
    for _ in range(4):
        box.cell(0, 0).add_paragraph()
    doc.add_paragraph()
    for who in ("От заказчика", "От команды проекта"):
        para(f"{who}: ____________ / ______________________ / ______________", center=False)
        para("подпись            ФИО, должность                 дата", size=9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
