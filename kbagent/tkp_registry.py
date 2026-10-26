"""Реестр ТКП: номер, статус и файлы каждого утверждённого предложения; бланк и условия компании.

Номер присваивается при утверждении — «ТКП-2026-0001», с начала года заново; под блокировкой, чтобы два
менеджера не получили один номер. Утверждённые Word и Excel хранятся в data/tkp/files/ — из реестра скачивается
ровно то, что ушло клиенту. Статус меняет менеджер: утверждён → отправлен клиенту → выигран / проигран / отменён;
каждая смена — в истории записи и в журнале действий. Сводка для руководителя: сколько ТКП, на какую сумму,
доля выигранных среди решённых.

Новая редакция — клиент поменял прибор или количество: «ТКП-2026-0001/1», прежняя версия получает статус «заменён»
и в сводке не считается. Напоминания: клиент молчит дольше follow_up_days после отправки, срок действия истекает
(expiring_days) или истёк, ТКП утверждён, но не отправлен — с черновиком письма клиенту и отметкой «напомнили».

Бланк компании (data/tkp/blank.docx) — шапка, реквизиты, логотип, колонтитулы; условия (data/tkp/settings.yaml) —
название, префикс номера, срок действия, оплата, поставка. Всё это данные заказчика: в git не попадают
и входят в резервную копию.
"""
import datetime as dt
import io
import json
import re
from pathlib import Path

import yaml

from . import prajs, tkp
from .actions import log_action
from .config import settings

DIR = settings.raw_dir.parent / "tkp"
SETTINGS = DIR / "settings.yaml"
BLANK = DIR / "blank.docx"
REGISTRY = DIR / "registry.json"
FILES = DIR / "files"
DEFAULTS = {"company": "", "prefix": "ТКП-", "validity_days": 30, "payment": "", "delivery": "",
            "follow_up_days": 5, "expiring_days": 5}
STATUSES = ["утверждён", "отправлен клиенту", "выигран", "проигран", "отменён", "заменён"]
ACTIVE = ("утверждён", "отправлен клиенту")         # решения клиента ещё нет
REPLACED = "заменён"                                # прежняя версия после новой редакции — не считается
CLIENT_RE = re.compile(r"\b(?:ООО|АО|ПАО|ЗАО|ОАО|ФГУП|МУП|ГУП|НПП|НПО|НИИ|ИП)\s*(?:[«\"][^»\"\n]{2,80}[»\"]|"
                       r"[А-ЯЁ][\w-]+(?:\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.)?)")


def load_settings() -> dict:
    try:
        data = yaml.safe_load(SETTINGS.read_text(encoding="utf-8")) or {}
    except FileNotFoundError:
        data = {}
    return {**DEFAULTS, **{k: v for k, v in data.items() if k in DEFAULTS}}


def save_settings(values: dict) -> dict:
    data = {**load_settings(), **{k: v for k, v in values.items() if k in DEFAULTS}}
    for key in ("validity_days", "follow_up_days", "expiring_days"):
        data[key] = int(data[key] or 0)
    DIR.mkdir(parents=True, exist_ok=True)
    SETTINGS.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return data


def save_blank(data: bytes) -> str:
    """Бланк компании: проверяется, что это Word; возвращает подсказку — есть ли в нём место {{ТКП}}."""
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 — не Word или повреждён
        raise ValueError("нужен документ Word (.docx)") from e
    DIR.mkdir(parents=True, exist_ok=True)
    BLANK.write_bytes(data)
    if any(tkp.PLACEHOLDER in p.text for p in doc.paragraphs):
        return "ТКП встанет на место метки {{ТКП}}."
    return "Метки {{ТКП}} в бланке нет — ТКП встанет после текста бланка."


def blank() -> Path | None:
    return BLANK if BLANK.exists() else None


def load() -> list[dict]:
    try:
        return json.loads(REGISTRY.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def _save(records: list[dict]) -> None:
    DIR.mkdir(parents=True, exist_ok=True)
    tmp = REGISTRY.with_suffix(".tmp")
    tmp.write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(REGISTRY)


def client_of(letter: str) -> str:
    """Организация клиента из письма: «ООО «Пример»», «АО «Прибор»» — иначе пусто."""
    m = CLIENT_RE.search(letter or "")
    return re.sub(r"\s+", " ", m.group(0)).strip() if m else ""


def _file(number: str, ext: str) -> Path:
    safe = re.sub(r"[^\w-]+", "_", number)          # вне f-строки: в Python 3.11 «\» в выражении f-строки нельзя
    return FILES / f"{safe}.{ext}"


def files(number: str) -> tuple[Path, Path]:
    return _file(number, "docx"), _file(number, "xlsx")


def approve(t: dict, text: str, login: str | None, name: str, role: str, question: str, app_name: str,
            now: dt.datetime | None = None, revision_of: str | None = None) -> dict:
    """Утверждение ТКП: номер, Word и Excel на бланке компании с этим номером, запись в реестре.
    revision_of — номер ТКП, которое заменяет эта редакция: номер «…/1», «…/2», прежняя версия — «заменён»."""
    from .store import _lock
    now = now or dt.datetime.now()
    conf = load_settings()
    terms = {"payment": conf["payment"], "delivery": conf["delivery"]}
    if conf["validity_days"]:
        terms["valid_until"] = (now.date() + dt.timedelta(days=int(conf["validity_days"]))).strftime("%d.%m.%Y")
    when = now.strftime("%d.%m.%Y %H:%M")
    priced = tkp._priced(t["positions"])
    total = sum(p["price"] * p["qty"] for p in priced) if priced else None
    with _lock(DIR):
        records = load()
        old = next((r for r in records if r["number"] == revision_of), None) if revision_of else None
        if revision_of and old is None:
            raise ValueError(f"ТКП {revision_of} нет в реестре")
        if old is not None:
            base = old["number"].split("/")[0]
            rev = 1 + max((int(r["number"].split("/")[1]) for r in records
                           if r["number"].split("/")[0] == base and "/" in r["number"]), default=0)
            number, year, seq = f"{base}/{rev}", old["year"], old["seq"]
        else:
            year = now.year
            seq = 1 + max((r["seq"] for r in records if r["year"] == year and "/" not in r["number"]), default=0)
            number = f"{conf['prefix']}{year}-{seq:04d}"
        FILES.mkdir(parents=True, exist_ok=True)
        docx_path, xlsx_path = files(number)
        docx_path.write_bytes(tkp.tkp_docx(text, t, name, when, app_name, number=number, terms=terms, blank=blank()))
        xlsx_path.write_bytes(tkp.tkp_xlsx(t, number, when[:10], terms, conf["company"]))
        record = {
            "number": number, "year": year, "seq": seq, "date": now.strftime("%Y-%m-%d %H:%M"),
            "revision_of": revision_of or "",
            "client": client_of(t.get("letter", "") or question), "question": question[:300],
            "positions": len(t["positions"]), "offered": sum(1 for p in t["positions"]
                                                             if p["status"] in ("offer", "named")),
            "models": [p["model"] for p in t["positions"] if p["status"] in ("offer", "named")],
            "total": total, "vat": tkp._vat(total)[0] if total and settings.vat_rate else None,
            "valid_until": terms.get("valid_until", ""), "status": STATUSES[0], "login": login, "approved_by": name,
            "role": role, "history": [{"status": STATUSES[0], "by": name, "at": when, "comment": ""}]}
        records.append(record)
        if old is not None:
            old["status"] = REPLACED
            old["history"].append({"status": REPLACED, "by": name, "at": when, "comment": f"новая редакция {number}"})
        _save(records)
    log_action("tkp_registered", login, number + (f" (редакция {revision_of})" if revision_of else ""), f"позиций: {record['positions']}"
               + (f"; сумма: {prajs.money(total)}" if total else ""))
    return record


def set_status(number: str, status: str, login: str | None, name: str, comment: str = "") -> dict:
    if status not in STATUSES:
        raise ValueError(f"статус — один из: {', '.join(STATUSES)}")
    from .store import _lock
    with _lock(DIR):
        records = load()
        record = next((r for r in records if r["number"] == number), None)
        if record is None:
            raise ValueError(f"ТКП {number} нет в реестре")
        record["status"] = status
        record["history"].append({"status": status, "by": name, "at": dt.datetime.now().strftime("%d.%m.%Y %H:%M"),
                                  "comment": comment.strip()})
        _save(records)
    log_action("tkp_status", login, number, status + (f": {comment.strip()}" if comment.strip() else ""))
    return record


def summary(records: list[dict]) -> dict:
    """Для руководителя: всего, сумма, выигранные, доля выигранных среди решённых (выигран + проигран).
    Прежние версии, заменённые новой редакцией, не считаются."""
    records = [r for r in records if r["status"] != REPLACED]
    won = [r for r in records if r["status"] == "выигран"]
    lost = [r for r in records if r["status"] == "проигран"]
    return {"count": len(records), "total": sum(r["total"] or 0 for r in records),
            "won": len(won), "won_total": sum(r["total"] or 0 for r in won),
            "open": sum(1 for r in records if r["status"] in ACTIVE),
            "win_rate": round(100 * len(won) / (len(won) + len(lost))) if won or lost else None}


def export(records: list[dict]) -> bytes:
    """Реестр в Excel — для отчёта и для CRM."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр ТКП"
    ws.append(["Номер", "Дата", "Клиент", "Позиций", "Приборы", "Сумма, ₽", "НДС, ₽", "Действует до", "Статус",
               "Утвердил", "Запрос клиента", "История"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in records:
        ws.append([r["number"], r["date"], r["client"], r["positions"], ", ".join(r["models"]), r["total"], r["vat"],
                   r["valid_until"], r["status"], r["approved_by"], r["question"],
                   "; ".join(f"{h['at']} {h['status']} ({h['by']})" + (f": {h['comment']}" if h["comment"] else "")
                             for h in r["history"])])
    for col, width in zip("ABCDEFGHIJKL", (16, 16, 26, 9, 30, 14, 12, 13, 18, 20, 50, 60)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _at(value: str) -> dt.datetime:
    return dt.datetime.strptime(value, "%d.%m.%Y %H:%M")


def _ru(day: dt.date) -> str:
    return day.strftime("%d.%m.%Y")


def attention(records: list[dict], today: dt.date | None = None) -> list[dict]:
    """ТКП, которые требуют действия менеджера: {record, kind, text}; срочные — первыми.
    kind: expired — срок действия истёк; expiring — истекает; silent — клиент молчит после отправки;
    not_sent — утверждён, но не отправлен."""
    today = today or dt.date.today()
    conf = load_settings()
    out = []
    for r in records:
        if r["status"] not in ACTIVE:
            continue
        last = _at(r["history"][-1]["at"]).date()
        idle = (today - last).days
        until = dt.datetime.strptime(r["valid_until"], "%d.%m.%Y").date() if r.get("valid_until") else None
        if until and until < today:
            out.append({"record": r, "kind": "expired", "order": 0, "text": f"срок действия истёк {_ru(until)}"})
        elif until and (until - today).days <= conf["expiring_days"]:
            days = (until - today).days
            out.append({"record": r, "kind": "expiring", "order": 1,
                        "text": f"срок действия истекает {_ru(until)}" + (" — сегодня" if days == 0 else "")})
        if r["status"] == "отправлен клиенту" and conf["follow_up_days"] and idle >= conf["follow_up_days"]:
            out.append({"record": r, "kind": "silent", "order": 2, "text": f"клиент молчит {idle} дн. после отправки"
                        if r["history"][-1]["status"] == "отправлен клиенту" else f"нет новостей {idle} дн."})
        elif r["status"] == "утверждён" and idle >= 2:
            out.append({"record": r, "kind": "not_sent", "order": 3, "text": f"утверждён {idle} дн. назад, не отправлен"})
    return sorted(out, key=lambda x: (x["order"], x["record"]["date"]))


def reminder_text(record: dict, signer: str = "") -> str:
    """Черновик письма клиенту: напоминание о ТКП или предложение обновить истёкшее."""
    d = record["date"]
    date = f"{d[8:10]}.{d[5:7]}.{d[:4]}"
    what = ", ".join(record["models"]) or "продукцию по вашему запросу"
    total = f" на сумму {prajs.money(record['total'])}" if record.get("total") else ""
    lines = ["Добрый день!", "",
             f"Напоминаем о нашем технико-коммерческом предложении № {record['number']} от {date}: {what}{total}."]
    until = dt.datetime.strptime(record["valid_until"], "%d.%m.%Y").date() if record.get("valid_until") else None
    if until and until < dt.date.today():
        lines.append(f"Срок действия предложения истёк {_ru(until)}. Если вопрос актуален, подготовим обновлённое "
                     "предложение с текущими ценами и сроками поставки.")
    elif until:
        lines.append(f"Предложение действительно до {_ru(until)}.")
    lines += ["Готовы ответить на вопросы, а при необходимости — подобрать альтернативу.", "", "С уважением,",
              signer or "[ФИО, должность]"]
    return "\n".join(lines)


def note(number: str, comment: str, login: str | None, name: str) -> dict:
    """Отметка без смены статуса — например, «напомнили клиенту»: сбрасывает отсчёт молчания."""
    from .store import _lock
    with _lock(DIR):
        records = load()
        record = next((r for r in records if r["number"] == number), None)
        if record is None:
            raise ValueError(f"ТКП {number} нет в реестре")
        record["history"].append({"status": record["status"], "by": name,
                                  "at": dt.datetime.now().strftime("%d.%m.%Y %H:%M"), "comment": comment.strip()})
        _save(records)
    log_action("tkp_note", login, number, comment.strip())
    return record
