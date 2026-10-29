"""Ответ агента — в документ Word: вопрос, ответ, записи учёта, источники с пунктами и страницами.
Чтобы ответ можно было вложить в служебную записку или отправить коллеге без чата."""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt, RGBColor

from .database import display
from .text import md_unescape

INLINE_RE = re.compile(r"(\*\*.+?\*\*|~~.+?~~)")
GRAY = RGBColor(0x5C, 0x54, 0x70)


def _runs(paragraph, text: str) -> None:
    """**жирный**, ~~зачёркнутый~~ (сравнение редакций), остальное — обычный текст."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("~~") and part.endswith("~~"):
            paragraph.add_run(part[2:-2]).font.strike = True
        else:
            paragraph.add_run(part)


def _text(doc, text: str) -> None:
    for line in md_unescape(text).splitlines():
        line = line.rstrip()
        if not line.strip():
            continue
        if re.match(r"^\s*[-*•]\s+", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            _runs(doc.add_paragraph(), line)


def draft_text(item: dict, signer: str = "") -> str:
    """Черновик ответа клиенту или коллеге: обращение, ответ без разметки, источники, подпись.
    Ничего не отправляется — сотрудник правит текст и утверждает."""
    body = re.sub(r"\*\*(.+?)\*\*", r"\1", item["text"].replace("[БД]", "[учётная база]"))
    # служебные пометки для сотрудника («Это черновик — проверьте…») клиенту не нужны
    body = re.sub(r"(?m)^(?:Значения — из справочника|Числа посчитаны программой)[^\n]*\n?", "", body)
    body = re.sub(r"~~(.+?)~~", r"было: \1;", body)
    cited = [s for s in item.get("sources", []) if s.get("cited")]
    lines = ["Добрый день!", "", f"Отвечаем на ваш вопрос: «{item['question']}».", "", body.strip()]
    if cited:
        lines += ["", "Источники:"] + [f"[{s['n']}] {s['label']}" for s in cited]
    lines += ["", "С уважением,", signer or "[ФИО, должность]"]
    return "\n".join(lines)


def draft_docx(text: str, item: dict, approved_by: str, approved_at: str, app_name: str) -> bytes:
    """Утверждённый черновик — документ Word с отметкой, кто и когда утвердил."""
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(20), Mm(10)
        section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    style.paragraph_format.first_line_indent = Cm(1.25)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.space_after = Pt(0)
    for line in text.splitlines():
        if line.strip():
            _runs(doc.add_paragraph(), line.rstrip())
        else:
            doc.add_paragraph()
    note = doc.add_paragraph()
    run = note.add_run(f"Черновик подготовлен: {app_name}, по документам и данным предприятия со ссылками на "
                       f"источники. Утвердил: {approved_by}, {approved_at}.")
    run.italic, run.font.color.rgb = True, GRAY
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _gost(doc) -> None:
    """A4, поля 20/10/20/20 мм, Times New Roman 12 — по ГОСТ Р 7.0.97."""
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(20), Mm(10)
        section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)


def ack_docx(event: dict, rows: list[dict], app_name: str) -> bytes:
    """Лист ознакомления с документом: кто и когда отметил «Ознакомлен» в чате; графа для подписи — для архива СМК."""
    doc = Document()
    _gost(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ЛИСТ ОЗНАКОМЛЕНИЯ").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"с документом «{event['title']}»")
    details = [f"Изменение: {event['kind_label']}" + (f" (вместо «{event['old_title']}»)" if event.get("old_title")
                                                     else ""),
               f"Дата: {display(event['time'][:10])}"]
    if event.get("summary"):
        details.append(f"Что изменилось: {event['summary']}")
    for line in details:
        doc.add_paragraph(line)
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["№", "Ф. И. О.", "Роль", "Ознакомлен (отметка в системе)",
                                                 "Подпись"]):
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    for n, r in enumerate(rows, 1):
        for cell, value in zip(table.add_row().cells, [str(n), r["name"], r["role"],
                                                       display(r["at"])[:16] if r["at"] else "не ознакомлен", ""]):
            cell.text = value
    note = doc.add_paragraph()
    done = sum(1 for r in rows if r["at"])
    run = note.add_run(f"\nОзнакомились {done} из {len(rows)}. Отметки «Ознакомлен» сотрудники ставят в чате "
                       f"«{app_name}» после уведомления о новой редакции; журнал отметок хранится в системе.")
    run.italic, run.font.color.rgb = True, GRAY
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def answer_docx(item: dict, app_name: str, user: str | None = None) -> bytes:
    """item — элемент диалога чата: question, text, sources, columns/labels/rows, at, route, status."""
    doc = Document()
    # оформление по ГОСТ Р 7.0.97 — тому же, по которому агент проверяет документы: A4, поля 20/10/20/20 мм,
    # Times New Roman 12, абзацный отступ 1,25 см, выравнивание по ширине
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(20), Mm(10)
        section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.first_line_indent = Cm(1.25)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.space_after = Pt(0)
    for name in ("Heading 2", "List Bullet", "List Number"):
        doc.styles[name].font.name = "Times New Roman"
    heading = doc.styles["Heading 2"]
    heading.font.size, heading.font.color.rgb = Pt(13), RGBColor(0, 0, 0)
    heading.paragraph_format.first_line_indent = Cm(0)

    doc.add_heading("Вопрос", level=2)
    doc.add_paragraph(item["question"])
    doc.add_heading("Ответ", level=2)
    _text(doc, item["text"].replace("[БД]", "[учётная база]"))

    if item.get("rows"):
        labels = item.get("labels") or item.get("columns") or []
        doc.add_heading("Записи учётной базы", level=2)
        table = doc.add_table(rows=1, cols=len(labels))
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, labels):
            cell.text = str(label)[:1].upper() + str(label)[1:]
            cell.paragraphs[0].runs[0].bold = True
        for row in item["rows"][:50]:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = "" if value is None else display(value)

    cited = [s for s in item.get("sources", []) if s.get("cited")]
    if cited:
        doc.add_heading("Источники", level=2)
        for s in cited:
            doc.add_paragraph(f"[{s['n']}] {s['label']}")

    note = doc.add_paragraph()
    run = note.add_run(f"{app_name}. Сформировано {item.get('at', '')}" + (f", пользователь {user}" if user else "")
                       + ". Ответ составлен только по документам и учётной базе предприятия; перед применением "
                         "сверьте с источником.")
    run.italic, run.font.color.rgb = True, GRAY
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
