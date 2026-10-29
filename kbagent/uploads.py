"""Документы, загруженные через чат: сохранение, индексация, утверждение, замена редакции, удаление.

Загружать и удалять документы может только администратор (проверяется в интерфейсе).
Загруженный документ сразу участвует в поиске, но до утверждения помечается
«загружен пользователем, не утверждён» — и в ссылке на источник, и предупреждением под ответом.
Если документ загружен как новая редакция другого («заменяет»), до утверждения это проект: его видит
только администратор, сотрудники получают ответы по действующей редакции. После утверждения старая
редакция отменяется и перестаёт участвовать в поиске; при удалении новой — возвращается.

Форматы: PDF (и сканы — с распознаванием текста), Word; картинка (скан или фото страницы) превращается в PDF
и распознаётся; таблица Excel — в документ Word, где каждая строка — абзац «Строка N. Столбец: значение; …»,
чтобы ссылка в ответе вела к листу и строке.
"""
import datetime as dt
import hashlib
import io
import re
import time
from itertools import zip_longest
from pathlib import Path

import yaml

from . import changes
from .actions import log_action
from .config import settings
from .editions import chunk_clauses, compare, compare_clauses
from .ingest import ARCHIVE, access_of, expired_on, is_draft_edition, load_catalog, not_yet_effective, ocr_available
from .text import tokenize
from .store import Index

UPLOAD_DIR = settings.raw_dir / "uploaded"
CATALOG = UPLOAD_DIR / "catalog.yaml"

KIND_PENDING = "Загруженный документ"
KIND_APPROVED = "Загруженный документ (утверждён)"

IMAGE_EXT = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp")
TABLE_EXT = (".xlsx",)
UPLOAD_TYPES = ["pdf", "docx", "xlsx"] + [e[1:] for e in IMAGE_EXT]
MAX_TABLE_ROWS = 5000


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%d.%m.%Y") if v.time() == dt.time() else v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        return f"{v:.10g}".replace(".", ",")
    return re.sub(r"\s+", " ", str(v)).strip()


def image_to_pdf(data: bytes) -> bytes:
    """Скан или фото страницы → одностраничный PDF шириной A4: дальше распознавание, как у скана в PDF."""
    import pymupdf
    pix = pymupdf.Pixmap(data)
    if pix.alpha:
        pix = pymupdf.Pixmap(pix, 0)
    width = 595.0
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=width * pix.height / pix.width)
    page.insert_image(page.rect, pixmap=pix)
    return doc.tobytes()


def table_to_docx(data: bytes, title: str) -> tuple[bytes, int]:
    """Таблица Excel → документ Word: по листам, каждая строка — абзац «Строка N. Заголовок: значение; …».
    Первая непустая строка листа — заголовки столбцов. Возвращает файл и число строк данных."""
    from docx import Document
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    doc = Document()
    doc.add_paragraph(title)
    total = 0
    for ws in wb.worksheets:
        header = None
        for n, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [_cell(v) for v in row]
            if not any(cells):
                continue
            if header is None:
                header = cells
                doc.add_paragraph(f"Лист «{ws.title}»")
                continue
            parts = [f"{h or f'столбец {k + 1}'}: {v}" for k, (h, v) in enumerate(zip_longest(header, cells)) if v]
            doc.add_paragraph(f"Лист «{ws.title}», строка {n}. " + "; ".join(parts))
            total += 1
            if total >= MAX_TABLE_ROWS:
                break
        if total >= MAX_TABLE_ROWS:
            break
    wb.close()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), total


def convert_upload(filename: str, data: bytes) -> tuple[str, bytes, dict]:
    """Расширение и содержимое файла, который пойдёт в базу знаний, и заметки о преобразовании."""
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXT:
        if not ocr_available():
            raise ValueError("для картинок нужно распознавание текста: python scripts/setup.py --ocr")
        return ".pdf", image_to_pdf(data), {"image": True}
    if ext in TABLE_EXT:
        docx, rows = table_to_docx(data, Path(filename).stem.replace("_", " "))
        return ".docx", docx, {"table_rows": rows}
    if ext not in (".pdf", ".docx"):
        raise ValueError("поддерживаются PDF, Word, Excel (XLSX) и картинки (PNG, JPG, TIFF, BMP)")
    return ext, data, {}


def _read_catalog() -> list[dict]:
    if CATALOG.exists():
        return yaml.safe_load(CATALOG.read_text(encoding="utf-8")).get("documents", [])
    return []


def _write_catalog(docs: list[dict]) -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    CATALOG.write_text(yaml.safe_dump({"documents": docs}, allow_unicode=True, sort_keys=False), encoding="utf-8")


def doc_id_for(filename: str) -> str:
    """Тот же файл повторно — тот же id (замена); длинные имена, различающиеся в конце, — разные id."""
    stem = re.sub(r"[^\w\-]+", "_", Path(filename).stem).strip("_")
    if len(stem) > 60:
        stem = f"{stem[:52]}_{hashlib.sha1(stem.encode('utf-8')).hexdigest()[:7]}"
    return f"upl_{stem or 'document'}"


CODE_LINE_RE = re.compile(r"^(?:СТП|СТО|РИ|ТИ|ИОТ|ПИ|ГОСТ(?:\s*Р)?|Приказ|ПРИКАЗ|Распоряжение|РАСПОРЯЖЕНИЕ|Положение|"
                          r"ПОЛОЖЕНИЕ)\b[^\n]{0,60}\d", re.I)
BOILERPLATE_RE = re.compile(r"(?:федеральн\w*|национальн\w*|министерств\w*|ооо|ао|пао|зао|демо)\b", re.I)
ABBREVIATIONS = {"ЭРИ", "СИ", "ОТК", "БВК", "ЕСКД", "ДП", "ГОСТ", "РИ", "СТП", "ПОС", "ESD", "MSL", "ЭСР", "ТУ", "ОТ"}


def _sentence_case(line: str) -> str:
    words = [w if w.strip("«»(),.") in ABBREVIATIONS else w.lower() for w in line.split()]
    text = " ".join(words)
    return text[:1].upper() + text[1:]


def guess_title(path: Path) -> str | None:
    """Название из самого документа: строка с обозначением («РИ ДП-05-2026», «Приказ № 17 от …») и следующая
    за ней строка прописными («ОБСЛУЖИВАНИЕ УСТАНОВКИ…») → «РИ ДП-05-2026 «Обслуживание установки…»».
    У скана первая страница распознаётся. Не нашлось — None (тогда название по имени файла)."""
    from .ingest import _page_text      # распознавание скана — как при разборе документа
    from .text import normalize
    try:
        if path.suffix.lower() == ".pdf":
            import pymupdf
            with pymupdf.open(path) as doc:
                text = _page_text(doc[0], {}) if len(doc) else ""
        else:
            from docx import Document
            text = "\n".join(p.text for p in Document(path).paragraphs[:25])
    except Exception:  # noqa: BLE001 — название не критично, документ всё равно добавляется
        return None
    lines = [ln.strip() for ln in normalize(text).splitlines() if ln.strip()][:30]
    code = next((i for i, ln in enumerate(lines) if CODE_LINE_RE.match(ln)), None)
    start = code + 1 if code is not None else 0
    title = []
    for ln in lines[start:start + 6]:
        letters = [c for c in ln if c.isalpha()]
        if len(letters) >= 4 and sum(c.isupper() for c in letters) / len(letters) > 0.8:
            title.append(ln)
        elif title:
            break
    name = _sentence_case(" ".join(title)) if title else ""
    if code is not None and name:
        return f"{lines[code]} «{name}»"[:200]
    if code is not None:
        return lines[code]
    # без обозначения — только если это не шапка ведомства или организации
    return name if name and not BOILERPLATE_RE.match(name) else None


def _sync_access(index: Index, doc_id: str) -> None:
    """Приводит доступ документа в индексе к каталогу (действует / отменён)."""
    info = load_catalog(settings.raw_dir).get(doc_id)
    if info is not None:
        index.update_meta(doc_id, settings.index_dir, access=access_of(info))


def add_upload(index: Index, filename: str, data: bytes, access: str = "all", user: str | None = None,
               replaces: str | None = None, progress=None, version: str = "", effective_from: str = "",
               owner: str = "") -> dict:
    """Сохраняет файл, добавляет его в каталог и индекс.
    Возвращает id, title, chunks, replaced_same (перезаписан одноимённый), ocr_pages, scanned_pages."""
    ext, data, converted = convert_upload(filename, data)
    doc_id = doc_id_for(filename)
    title = Path(filename).stem.replace("_", " ")
    docs = _read_catalog()
    replaced_same = any(d["id"] == doc_id for d in docs)

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    for old in UPLOAD_DIR.glob(f"{doc_id}.*"):
        if old.suffix.lower() in (".pdf", ".docx"):
            old.unlink()
    path = UPLOAD_DIR / f"{doc_id}{ext}"
    path.write_bytes(data)
    title = guess_title(path) or title       # «РИ ДП-05-2026 «Обслуживание…»», а не имя файла латиницей

    old_clauses = chunk_clauses([c for c in index.chunks if c.doc_id == doc_id]) if replaced_same else []
    entry = {"id": doc_id, "title": title, "kind": KIND_PENDING, "access": access,
             "uploaded_by": user, "uploaded_at": time.strftime("%Y-%m-%d %H:%M"), "approved": False}
    if replaces and replaces != doc_id:
        entry["replaces"] = replaces
    # реестр: версия, дата вступления в силу (до неё документ видит только администратор), владелец
    entry.update({k: v for k, v in (("version", version), ("effective_from", effective_from), ("owner", owner)) if v})
    _write_catalog([d for d in docs if d["id"] != doc_id] + [entry])

    stats: dict = {}
    chunks = index.add_document(path, load_catalog(settings.raw_dir), settings.index_dir,
                                progress=progress, stats=stats)
    if replaces:
        _sync_access(index, replaces)
    elif replaced_same:          # тот же файл загружен заново — что изменилось, видно по прежним фрагментам
        changes.record("updated", doc_id, title, access, user,
                       cmp=compare_clauses(old_clauses, chunk_clauses(chunks)), ack=True)
    else:                        # новая редакция — сотрудникам сообщается при утверждении, новый документ — сразу
        changes.record("added", doc_id, title, access, user)
    log_action("upload", user, title, f"фрагментов: {len(chunks)}; доступ: {access}"
               + (f"; заменяет: {replaces}" if replaces else "") + (f"; распознано страниц: {stats['ocr_pages']}"
                                                                     if stats.get("ocr_pages") else ""))
    return {"id": doc_id, "title": title, "chunks": len(chunks), "replaced_same": replaced_same, **converted, **stats}


def approve_upload(index: Index, doc_id: str, user: str | None = None) -> None:
    """Утверждение: документ помечается утверждённым; проект новой редакции становится доступен всем,
    а прежняя редакция отменяется."""
    docs = _read_catalog()
    for d in docs:
        if d["id"] == doc_id:
            d.update(approved=True, kind=KIND_APPROVED, approved_by=user,
                     approved_at=time.strftime("%Y-%m-%d %H:%M"))
    _write_catalog(docs)
    index.update_meta(doc_id, settings.index_dir, doc_kind=KIND_APPROVED)
    _sync_access(index, doc_id)
    replaced = next((d.get("replaces") for d in docs if d["id"] == doc_id), None)
    if replaced:
        _sync_access(index, replaced)
        info = load_catalog(settings.raw_dir).get(doc_id, {})
        if not not_yet_effective(info):     # вступает в силу позже — сообщение появится в тот день (changes.dated)
            old = load_catalog(settings.raw_dir).get(replaced, {})
            try:
                cmp = compare(replaced, doc_id)
            except Exception:  # noqa: BLE001 — сравнение не должно мешать утверждению
                cmp = None
            changes.record("edition", doc_id, info.get("title", doc_id), info.get("access", "all"), user,
                           old_id=replaced, old_title=old.get("title", replaced), cmp=cmp)
    log_action("approve", user, next((d["title"] for d in docs if d["id"] == doc_id), doc_id),
               f"прежняя редакция отменена: {replaced}" if replaced else "")


def similar_documents(title: str, exclude: set[str] | None = None, threshold: float = 0.6) -> list[str]:
    """Действующие документы с почти тем же названием — возможно, загружена новая редакция без отметки «заменяет»:
    тогда в поиске окажутся обе редакции, и агент может сослаться на любую."""
    words = set(tokenize(title))
    out = []
    for doc_id, info in load_catalog(settings.raw_dir).items():
        if doc_id in (exclude or set()) or access_of(info) == ARCHIVE:
            continue
        other = set(tokenize(info.get("title", "")))
        if words and other and len(words & other) / len(words | other) >= threshold:
            out.append(info.get("title", doc_id))
    return out


def list_uploads() -> list[dict]:
    return _read_catalog()


def remove_upload(index: Index, doc_id: str, user: str | None = None) -> None:
    docs = _read_catalog()
    log_action("remove", user, next((d["title"] for d in docs if d["id"] == doc_id), doc_id))
    replaced = next((d.get("replaces") for d in docs if d["id"] == doc_id), None)
    entry = next((d for d in docs if d["id"] == doc_id), None)
    if entry is not None and not is_draft_edition(entry):     # неутверждённый проект редакции никто не видел
        changes.record("removed", doc_id, entry.get("title", doc_id), entry.get("access", "all"), user)
    _write_catalog([d for d in docs if d["id"] != doc_id])
    index.remove_document(doc_id, settings.index_dir)
    for p in UPLOAD_DIR.glob(f"{doc_id}.*"):
        if p.suffix.lower() in (".pdf", ".docx"):
            p.unlink()
    if replaced:
        _sync_access(index, replaced)   # старая редакция снова действует


def documents(index: Index) -> list[dict]:
    """Все документы базы знаний со статусом — для списка у администратора."""
    catalog = load_catalog(settings.raw_dir)
    present = index.doc_ids()
    out = []
    for doc_id, info in catalog.items():
        if doc_id in present:
            ru = lambda d: f"{d[8:10]}.{d[5:7]}.{d[:4]}"  # noqa: E731
            status = ("отменён" if access_of(info) == ARCHIVE else
                      "проект новой редакции" if is_draft_edition(info) else
                      f"вступает в силу {ru(not_yet_effective(info))}" if not_yet_effective(info) else
                      f"срок действия истёк {ru(expired_on(info))}" if expired_on(info) else "действует")
            out.append({"id": doc_id, "title": info.get("title", doc_id), "kind": info.get("kind", ""),
                        "access": info.get("access", "all"), "status": status,
                        "version": info.get("version", ""), "owner": info.get("owner", ""),
                        "effective": " – ".join(ru(str(info[k])[:10]) for k in ("effective_from", "effective_to")
                                                if info.get(k)),
                        "priority": info.get("priority", ""),
                        "replaced_by": info.get("replaced_by")})
    return out
