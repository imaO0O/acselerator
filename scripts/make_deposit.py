"""Материалы для регистрации программы для ЭВМ в Роспатенте: реферат и депонируемые материалы.

    python scripts/make_deposit.py                    → docs/registraciya/out/
    python scripts/make_deposit.py --pages 25         → по 25 первых и последних страниц листинга

Требования — приказ Минэкономразвития № 211 от 05.04.2016 и руководство ФИПС по регистрации программ (раздел VI):
  * депонируемые материалы — исходный текст полностью или фрагментами в объёме, достаточном для идентификации,
    в формате PDF/A; титульный лист — название, правообладатель и все авторы; у фрагментов, кроме сквозной
    нумерации, в скобках — номер страницы в полном исходном тексте; материалы относятся к той версии,
    которая указана в заявлении;
  * реферат — название как в заявлении, назначение, область применения, функциональные возможности; завершается
    языком программирования и объёмом программы в байтах; не больше 900 знаков; сокращения расшифрованы;
    не должен выглядеть как заявка на несколько объектов (программа и база данных).

Скрипт собирает два файла Word. PDF/A делается в Word: «Сохранить как» → PDF → «Параметры» → «Совместимо
с ISO 19005-1 (PDF/A)». Правообладатели и авторы — пропуски в квадратных скобках: их состав решает команда
вместе с партнёром (в карточке задачи — совместное владение результатом). Порядок — docs/registraciya/README.md.
"""
import argparse
import subprocess
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "registraciya" / "out"
NAME = "AI-агент для поиска ответов по внутренней документации и базам данных предприятия"
# что входит в программу: ядро, интерфейс, программный интерфейс, установка и запуск (без тестов, наборов
# проверок, документов и моделей — модели сторонние, открытые)
PARTS = ["kbagent", "app", "api", "scripts"]
SUFFIXES = (".py", ".ps1", ".bat")
FIRST = ["kbagent/config.py", "kbagent/llm.py", "kbagent/text.py", "kbagent/ingest.py", "kbagent/store.py",
         "kbagent/retriever.py", "kbagent/answer.py", "kbagent/agent.py"]     # ядро — в начало листинга
WIDTH, LINES = 100, 70          # знаков в строке листинга и строк на странице (Courier New 8 пт)
REFERAT = (
    "{name} (AI — искусственный интеллект). Предназначена для подготовки ответов сотрудникам предприятия по "
    "внутренним документам, справочнику параметров продукции и учётным системам. Область применения: "
    "промышленные предприятия — продажи и сервис, производство, метрология. Функциональные возможности: разбор "
    "документов (PDF, Word, сканы, Excel) на пункты; поиск по словам и по смыслу; ответ локальной языковой "
    "моделью только по найденным фрагментам со ссылкой на документ и пункт, отказ при отсутствии сведений; "
    "запросы к учётной системе только на чтение; подбор продукции по требованиям клиента и черновик "
    "технико-коммерческого предложения; сравнение редакций документов; разграничение доступа по ролям, журналы "
    "действий. Работает без подключения к интернету. Язык: Python. "
    "Объём программы: {size} байт."
)


def num(x: int) -> str:
    return f"{x:,}".replace(",", " ")


def git(*args: str) -> str:
    return subprocess.run(["git", "-C", str(ROOT), *args], capture_output=True, text=True, encoding="utf-8",
                          check=True).stdout.strip()


def program_files() -> list[str]:
    files = [f for f in git("ls-files", *PARTS).splitlines() if f.endswith(SUFFIXES)]
    return [f for f in FIRST if f in files] + sorted(f for f in files if f not in FIRST)


def listing(files: list[str]) -> list[tuple[str, list[str]]]:
    """Полный исходный текст, разбитый на страницы: (файл в начале страницы, строки страницы)."""
    lines: list[tuple[str, str]] = []
    for f in files:
        text = (ROOT / f).read_text(encoding="utf-8-sig").expandtabs(4)
        lines.append((f, f"# {'=' * 30} {f} {'=' * 30}"))
        for raw in text.splitlines():
            wrapped = textwrap.wrap(raw, WIDTH, subsequent_indent="    ", drop_whitespace=False,
                                    replace_whitespace=False) or [""]
            lines += [(f, w.rstrip()) for w in wrapped]
    return [(lines[i][0], [t for _, t in lines[i:i + LINES]]) for i in range(0, len(lines), LINES)]


def setup(doc: Document) -> None:
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(25), Mm(10), Mm(20), Mm(20)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    st.paragraph_format.space_after = Pt(0)


def para(doc: Document, text: str, bold: bool = False, center: bool = False, size: int = 12, before: int = 0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.5
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold, run.font.size = bold, Pt(size)
    return p


def deposit(files: list[str], size: int, version: str, pages_each: int) -> Path:
    pages = listing(files)
    total = len(pages)
    take = list(range(total)) if total <= 2 * pages_each else \
        list(range(pages_each)) + list(range(total - pages_each, total))
    doc = Document()
    setup(doc)
    para(doc, "ДЕПОНИРУЕМЫЕ МАТЕРИАЛЫ,", bold=True, center=True, before=120)
    para(doc, "идентифицирующие программу для ЭВМ", center=True)
    para(doc, f"«{NAME}»", bold=True, center=True, size=14, before=24)
    para(doc, version, center=True, before=12)
    para(doc, "Правообладатель(и): [полное наименование или ФИО — по заявлению]", before=48)
    para(doc, "Авторы: [ФИО всех авторов, не отказавшихся быть указанными]")
    para(doc, "Язык программирования: Python")
    para(doc, f"Объём программы: {num(size)} байт, файлов исходного текста — {len(files)}")
    para(doc, (f"Фрагменты исходного текста: первые {pages_each} и последние {pages_each} страниц из {total}; "
               "в скобках — номер страницы в полном исходном тексте.") if len(take) < total
         else f"Исходный текст полностью: {total} страниц.")
    for n, k in enumerate(take, start=1):
        first_file, lines = pages[k]
        head = doc.add_paragraph()
        head.paragraph_format.page_break_before = True
        run = head.add_run(f"Стр. {n} ({k + 1})    {first_file}")
        run.bold, run.font.size, run.font.name = True, Pt(8), "Courier New"
        for line in lines:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = Pt(9.6)
            r = p.add_run(line)
            r.font.size, r.font.name = Pt(8), "Courier New"
    path = OUT / "Deponiruemye_materialy.docx"
    doc.save(path)
    return path


def referat(size: int) -> tuple[Path, int]:
    text = REFERAT.format(name=NAME, size=num(size))
    if len(text) > 900:
        raise SystemExit(f"Реферат — {len(text)} знаков, больше 900: сократите REFERAT")
    doc = Document()
    setup(doc)
    para(doc, "РЕФЕРАТ", bold=True, center=True)
    p = para(doc, text, before=12)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(12.5)
    path = OUT / "Referat.docx"
    doc.save(path)
    (OUT / "referat.txt").write_text(text + "\n", encoding="utf-8")
    return path, len(text)


def main() -> None:
    ap = argparse.ArgumentParser(description="Реферат и депонируемые материалы для Роспатента")
    ap.add_argument("--pages", type=int, default=25, help="первых и последних страниц листинга (по умолчанию 25)")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    if git("status", "--porcelain", "--", *PARTS):
        print("Внимание: в коде есть незакоммиченные изменения — материалы должны соответствовать версии "
              "из заявления, соберите их из закоммиченного состояния.")
    files = program_files()
    size = sum((ROOT / f).stat().st_size for f in files)
    commit, day = git("log", "-1", "--format=%h"), git("log", "-1", "--format=%cd", "--date=format:%d.%m.%Y")
    version = f"Версия от {day} (коммит {commit}), год создания — 2026"
    ref, chars = referat(size)
    dep = deposit(files, size, version, args.pages)
    print(f"Реферат: {ref} — {chars} знаков из 900")
    print(f"Депонируемые материалы: {dep}")
    print(f"Программа: {len(files)} файлов, {num(size)} байт. {version}")
    print("Дальше: впишите правообладателей и авторов, сохраните оба файла в Word как PDF/A — docs/registraciya/README.md")


if __name__ == "__main__":
    main()
