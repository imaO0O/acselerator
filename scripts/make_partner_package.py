"""Пакет для партнёра перед пилотом: что прислать и в каком виде.

    python scripts/make_partner_package.py   → docs/pilot/paket/Shablon_dannyh_dlya_pilota.xlsx
                                                docs/pilot/paket/Pamyatka_partneru.docx

Книга Excel — листы «Реестр документов», «Вопросы», «Справочник параметров», «Прайс» с подсказками в заголовках,
выпадающими списками и строкой-примером; столбцы — те, что читает scripts/import_partner.py
(kbagent/pilot_data.py). Памятка — одна страница: что нужно, обезличивание, как передать, что дальше.
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from kbagent.pilot_data import EXAMPLE, EXAMPLES, SHEETS  # noqa: E402

OUT = ROOT / "docs" / "pilot" / "paket"
HEAD = PatternFill("solid", fgColor="E9E4F3")
REQUIRED = PatternFill("solid", fgColor="FFD8A8")
ROWS = 300                                  # строк с выпадающими списками

GUIDE = [
    ("Шаблон данных для пилота AI-агента", True),
    ("Заполните листы и пришлите книгу вместе с документами. Строку «пример» можно оставить — она не загружается.",
     False),
    ("Оранжевые заголовки — обязательные столбцы. Подсказка — в примечании к заголовку (наведите курсор).", False),
    ("", False),
    ("Реестр документов", True),
    ("Каждый присланный файл — одна строка: название, редакция, срок действия, владелец, кто видит. Версия и срок "
     "попадают в каждую ссылку ответа, отменённая редакция в ответах не участвует.", False),
    ("Вопросы — для настройки", True),
    ("20–30 реальных вопросов сотрудников или клиентов с правильными ответами. «Ключевые факты» — короткие "
     "точные фрагменты, без которых ответ неверен: «2 рабочих дня», «200 МГц», «2211». Варианты — через «|».",
     False),
    ("Вопросы — закрытый приёмочный набор", True),
    ("Ещё 30–50 других вопросов в этой же форме составляет эксперт отдела в отдельной копии книги и хранит у себя "
     "до приёмки: команда их не видит, поэтому проверка честная.", False),
    ("Справочник параметров", True),
    ("Если есть выгрузка справочника — пришлите её как есть. Если нет — соберём по руководствам сами; лист — для "
     "тех значений, которые вы хотите задать явно. Каждое значение — с документом, страницей и цитатой.", False),
    ("Прайс", True),
    ("По желанию — для черновиков ТКП: цена за единицу и срок поставки. Без прайса цены в ТКП вписывает менеджер.",
     False),
]


def build_xlsx(path: Path) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "Как заполнить"
    guide.column_dimensions["A"].width = 120
    for i, (text, bold) in enumerate(GUIDE, start=1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        for j, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=j, value=col.name)
            cell.font = Font(bold=True)
            cell.fill = REQUIRED if col.required else HEAD
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if col.hint:
                cell.comment = Comment(col.hint, "команда проекта")
            ws.column_dimensions[cell.column_letter].width = col.width
            if col.choices:
                dv = DataValidation(type="list", formula1='"' + ",".join(col.choices) + '"', allow_blank=True)
                dv.add(f"{cell.column_letter}2:{cell.column_letter}{ROWS}")
                ws.add_data_validation(dv)
        for j, value in enumerate([EXAMPLE] + EXAMPLES[name], start=1):
            c = ws.cell(row=2, column=j, value=value)
            c.font = Font(italic=True, color="7A7A7A")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
    wb.save(path)


def build_memo(path: Path) -> None:
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(20), Mm(15), Mm(15), Mm(15)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)
    st.paragraph_format.space_after = Pt(3)

    def para(text: str, bold: bool = False, size: int = 11, center: bool = False, before: int = 0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
        return p

    para("Пилот AI-агента: какие материалы нужны и в каком виде", bold=True, size=13, center=True)
    para("Пилот — один отдел (технические продажи или сервис) и одна группа продукции. Цель из карточки задачи — "
         "сократить время подготовки проверенного ответа не менее чем на 50 % без ухудшения качества. Данные "
         "не передаются в облачные сервисы: агент работает на компьютере предприятия без интернета.", before=6)
    para("Что нужно", bold=True, before=6)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ["Что", "Сколько и в каком виде", "Зачем"]):
        cell.text = head
        cell.paragraphs[0].runs[0].bold = True
    items = [
        ("Документы группы продукции", "20–50 файлов: руководства, ТУ, регламенты, прайс-листы — PDF, Word, Excel, "
         "сканы; лист «Реестр документов»", "по ним агент отвечает со ссылкой на пункт и страницу"),
        ("Справочник параметров", "выгрузка CSV или Excel, если есть; если нет — соберём по руководствам",
         "подбор модели под требования клиента программой, без «догадок» модели"),
        ("Прайс (по желанию)", "лист «Прайс»: цена за единицу и срок поставки", "цены и суммы в черновике ТКП"),
        ("Вопросы для настройки", "20–30 реальных вопросов с правильными ответами — лист «Вопросы»",
         "настройка агента на вашу терминологию"),
        ("Закрытый приёмочный набор", "30–50 других вопросов в той же форме и 10–20 реальных запросов клиентов "
         "со временем ручного ответа; составляет эксперт отдела и хранит у себя до приёмки",
         "честная проверка: команда эти вопросы заранее не видит"),
        ("Сведения для установки", "компьютер или сервер: Windows или Astra Linux / РЕД ОС, видеокарта NVIDIA и её "
         "память, оперативная память; вход — учётные записи системы или Active Directory; две роли и что видит каждая",
         "подготовить поставку под ваш контур"),
    ]
    for row in items:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for row in table.rows:
        for cell, width in zip(row.cells, (Mm(42), Mm(80), Mm(53))):
            cell.width = width
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    para("Обезличивание", bold=True, before=8)
    para("Уберите из документов и вопросов ФИО и контакты клиентов, реквизиты, договорные цены и условия, которые "
         "нельзя показывать. Техническое содержание оставьте — по нему агент и отвечает. Если документы нельзя "
         "выносить, работаем на вашей территории: системе не нужен интернет. По вашему желанию подпишем соглашение "
         "о конфиденциальности.")
    para("Как передать", bold=True, before=6)
    para("Архив на флешке или по защищённому каналу предприятия — не через публичные облачные хранилища. Шаблон "
         "Excel заполните и приложите к архиву; имена файлов в реестре — как в архиве.")
    para("Что дальше", bold=True, before=6)
    para("Через 5–6 недель после получения материалов — установка пилотной версии в отделе и обучение сотрудников "
         "(2–3 часа); если сервер на Astra Linux или РЕД ОС или нужен вход через Active Directory — на 3–4 недели "
         "позже. Затем 1–2 месяца пилота с замером «вручную / с агентом» на закрытом наборе. В конце — отчёт: "
         "время подготовки ответа, качество по оценке эксперта, пробелы в документации, стоимость обработки запроса.")
    para("Контакты команды: [ФИО лидера команды], [телефон], [почта]", before=8)
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx, memo = OUT / "Shablon_dannyh_dlya_pilota.xlsx", OUT / "Pamyatka_partneru.docx"
    build_xlsx(xlsx)
    build_memo(memo)
    print(xlsx, memo, sep="\n")


if __name__ == "__main__":
    main()
